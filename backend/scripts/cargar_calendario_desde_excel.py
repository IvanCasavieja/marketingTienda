"""
Carga el calendario comercial del Excel 'CALENDARIO 2025 Y 2026 TI.xlsx' en la
plataforma, por la API, sin pisar lo que la gente ya edito en la pantalla.

Por que existe: el 02/10/2026 Ivan pidio que el calendario de promociones que
vive en ese Excel (24 hojas, una por mes de 2025 y 2026) se cargue y se vaya
actualizando en la plataforma. El calendario de la plataforma guarda una fila
por barra desde el 28/09 (migracion 0057) y se escribe por /calendario/barras,
asi que lo mas seguro es ir por la MISMA API que usa la pantalla: misma
validacion, misma revision, y no hace falta DATABASE_URL en la PC.

Como lee el Excel: con el mismo lector que importar_calendario_excel.py (la
fila de dias, las celdas combinadas como barras, el color del relleno), mas un
corte que ese lector no tenia: debajo de la tabla del mes muchas hojas traen
la tabla del MISMO mes del año anterior, para comparar, y sin el corte se leia
entera como si fueran acciones de este año (139 barras en vez de 36).

Que hace con cada barra del Excel, por banda y mes:
  - si ya hay una igual (banda, nombre, fechas) en la plataforma, nada;
  - si hay una con el mismo nombre pero otras fechas o color, la actualiza;
  - si hay una sin pareja con fechas que se tocan y un nombre parecido (la
    renombraron en el Excel), la actualiza conservando su id, sus piezas y sus
    avisos;
  - si no hay ninguna, la crea con un id determinista (br-xl-<mes>-<hash>):
    correrlo dos veces no duplica nada.
Lo que esta en la plataforma y no en el Excel se LISTA y no se borra: puede
ser algo que alguien creo en la pantalla, y borrar es decision de una persona.

Las bandas se unifican solo por mayusculas y espacios ('Container' y
'CONTAINER' son la misma); 'FARMA' y 'TIENDA FARMA' se dejan distintas porque
no es seguro que sean lo mismo.

Uso:
  MKTG_EMAIL=... MKTG_PASSWORD=... python backend/scripts/cargar_calendario_desde_excel.py
  MKTG_EMAIL=... MKTG_PASSWORD=... python backend/scripts/cargar_calendario_desde_excel.py --aplicar
  (--meses 2026-11,2026-12 limita los meses; MKTG_API_URL opcional)

Sin --aplicar solo muestra el plan. Con --aplicar guarda antes un respaldo del
calendario entero en backend/backups/.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import importlib.util
import json
import os
import re
import sys
import time
import warnings
from datetime import date, datetime
from pathlib import Path

import httpx
import openpyxl
from rapidfuzz import fuzz

warnings.filterwarnings("ignore")

API = os.environ.get("MKTG_API_URL", "https://marketingtienda.onrender.com/api/v1").rstrip("/")
AQUI = Path(__file__).resolve().parent
BACKUPS = AQUI.parent / "backups"

_spec = importlib.util.spec_from_file_location("importar_calendario_excel", AQUI / "importar_calendario_excel.py")
imp = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(imp)

MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6, "JULIO": 7,
         "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12}
PARECIDO_MINIMO = 80     # token_set_ratio para dar dos nombres por el mismo
DIAS_DE_TOLERANCIA = 3   # cuanto pueden separarse las fechas de una renombrada
PAUSA = 0.35             # el limite de la API es 200 pedidos por minuto


def clave_de_hoja(ws) -> str:
    """'Octubre 2026' -> '2026-10'. Las hojas de 2025 no dicen el año en el
    nombre; algunas lo traen como fecha en B1 y si no, es 2025."""
    t = ws.title.strip().upper()
    anio = 2026 if "2026" in t else 2025
    b1 = ws.cell(1, 2).value
    if hasattr(b1, "year"):
        anio = b1.year
    mes = next(m for nombre, m in MESES.items() if t.startswith(nombre))
    return f"{anio}-{mes:02d}"


def es_fila_de_dias(ws, r) -> bool:
    n = sum(1 for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(r, c).value, (int, float)) and 1 <= ws.cell(r, c).value <= 31)
    return n >= 25


def corte_de(ws, fd) -> int:
    """Donde termina la tabla del mes: antes del bloque de comparacion."""
    for r in range(fd + 2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is not None and re.match(r"^\s*(DIA DE LA SEMANA|20\d\d)\s*$", str(a).upper()):
            return r
        if es_fila_de_dias(ws, r):
            return r - 1
    return ws.max_row + 1


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().casefold()


def d(s: str) -> date:
    return date.fromisoformat(s)


def se_tocan(a: dict, b: dict) -> bool:
    return ((d(a["desde"]) - d(b["hasta"])).days <= DIAS_DE_TOLERANCIA
            and (d(b["desde"]) - d(a["hasta"])).days <= DIAS_DE_TOLERANCIA)


def parecidos(n1: str, n2: str) -> bool:
    a, b = norm(n1), norm(n2)
    if not a or not b:
        return False
    if a.startswith(b) or b.startswith(a):
        return True
    return fuzz.token_set_ratio(a, b) >= PARECIDO_MINIMO


def leer_excel(meses_filtro: set[str] | None) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(imp.RUTAS["comercial"], data_only=True)
    por_mes: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        clave = clave_de_hoja(ws)
        if meses_filtro and clave not in meses_filtro:
            continue
        fd = imp.fila_de_dias(ws)
        col2dia, ci, cf, _dow = imp.ejes(ws, fd)
        bandas = imp.leer_bandas(ws, fd, 1, corte_de(ws, fd), ci, cf, col2dia)
        barras = []
        for b in bandas:
            if re.match(r"^20\d\d$", b["nombre"]):
                continue
            for carril, fila in enumerate(b["filas"]):
                for x in fila:
                    desde, hasta = min(x["desde"], x["hasta"]), max(x["desde"], x["hasta"])
                    barras.append({
                        "banda": b["nombre"], "carril": carril, "nombre": x["nombre"], "color": x["color"],
                        "desde": f"{clave}-{desde:02d}", "hasta": f"{clave}-{hasta:02d}",
                    })
        por_mes[clave] = barras
    return por_mes


def id_de(mes: str, b: dict) -> str:
    crudo = "|".join([norm(b["banda"]), str(b["carril"]), norm(b["nombre"]), b["desde"], b["hasta"]])
    return f"br-xl-{mes}-{hashlib.sha1(crudo.encode('utf-8')).hexdigest()[:10]}"


def carril_libre(existentes, banda, desde, hasta, preferido, excluir=None) -> int:
    """Mismo criterio que carrilLibre en frontend/lib/calendario/derivar.ts."""
    ocupados = {b["carril"] for b in existentes
                if b["seccion"] == "comercial" and norm(b["banda"]) == norm(banda) and b["id"] != excluir
                and b["desde"] <= hasta and desde <= b["hasta"]}
    if preferido not in ocupados:
        return preferido
    i = 0
    while i in ocupados:
        i += 1
    return i


def cambios_entre(c: dict, x: dict, existentes: list[dict], con_nombre: bool) -> dict:
    cambios = {}
    if con_nombre and c["nombre"] != x["nombre"]:
        cambios["nombre"] = x["nombre"]
    if c["desde"] != x["desde"]:
        cambios["desde"] = x["desde"]
    if c["hasta"] != x["hasta"]:
        cambios["hasta"] = x["hasta"]
    if x["color"] and c.get("color") != x["color"]:
        cambios["color"] = x["color"]
    if "desde" in cambios or "hasta" in cambios:
        nuevo = carril_libre(existentes, x["banda"], x["desde"], x["hasta"], c["carril"], excluir=c["id"])
        if nuevo != c["carril"]:
            cambios["carril"] = nuevo
    return cambios


def armar_plan(excel: dict[str, list[dict]], existentes: list[dict]) -> dict:
    comerciales = [b for b in existentes if b["seccion"] == "comercial"]
    bandas_plataforma = {norm(b["banda"]): b["banda"] for b in comerciales}
    plan = {"crear": [], "cambiar": [], "iguales": 0, "solo_plataforma": [], "bandas_nuevas": []}
    bandas_vistas = dict(bandas_plataforma)
    for mes in sorted(excel):
        del_mes = [b for b in comerciales if b["desde"][:7] == mes or b["hasta"][:7] == mes]
        usadas: set[str] = set()
        exactas, por_nombre = {}, {}
        for b in del_mes:
            exactas[(norm(b["banda"]), norm(b["nombre"]), b["desde"], b["hasta"])] = b
            por_nombre.setdefault((norm(b["banda"]), norm(b["nombre"])), []).append(b)

        pendientes = []
        # Primero las que ya estan iguales, asi un cambio de fecha no se roba la pareja de otra.
        orden = sorted(excel[mes], key=lambda x: (norm(x["banda"]), norm(x["nombre"]), x["desde"], x["hasta"]) not in exactas)
        for x in orden:
            x["banda"] = bandas_vistas.setdefault(norm(x["banda"]), x["banda"])
            k = (norm(x["banda"]), norm(x["nombre"]), x["desde"], x["hasta"])
            if k in exactas and exactas[k]["id"] not in usadas:
                usadas.add(exactas[k]["id"])
                plan["iguales"] += 1
                continue
            candidatos = [c for c in por_nombre.get((norm(x["banda"]), norm(x["nombre"])), []) if c["id"] not in usadas]
            if candidatos:
                c = min(candidatos, key=lambda c: abs((d(c["desde"]) - d(x["desde"])).days))
                usadas.add(c["id"])
                cambios = cambios_entre(c, x, existentes, con_nombre=False)
                if cambios:
                    plan["cambiar"].append({"id": c["id"], "mes": mes, "banda": x["banda"], "nombre": x["nombre"],
                                            "antes": {k2: c.get(k2) for k2 in ("nombre", "desde", "hasta", "color")},
                                            "cambios": cambios, "motivo": "mismo nombre, otras fechas"})
                else:
                    plan["iguales"] += 1
                continue
            pendientes.append(x)

        # Renombradas: sin pareja por nombre, pero misma banda, fechas que se tocan y nombre parecido.
        sobrantes = [b for b in del_mes if b["id"] not in usadas]
        for x in pendientes:
            pareja = next((c for c in sobrantes if c["id"] not in usadas and norm(c["banda"]) == norm(x["banda"])
                           and se_tocan(c, x) and parecidos(c["nombre"], x["nombre"])), None)
            if pareja:
                usadas.add(pareja["id"])
                plan["cambiar"].append({"id": pareja["id"], "mes": mes, "banda": x["banda"], "nombre": x["nombre"],
                                        "antes": {k2: pareja.get(k2) for k2 in ("nombre", "desde", "hasta", "color")},
                                        "cambios": cambios_entre(pareja, x, existentes, con_nombre=True),
                                        "motivo": "renombrada en el Excel"})
                continue
            nueva = {"id": id_de(mes, x), "seccion": "comercial", "banda": x["banda"],
                     "carril": carril_libre(existentes, x["banda"], x["desde"], x["hasta"], x["carril"]),
                     "nombre": x["nombre"], "color": x["color"], "desde": x["desde"], "hasta": x["hasta"]}
            existentes.append(nueva)
            plan["crear"].append({"mes": mes, **nueva})
        for b in del_mes:
            if b["id"] not in usadas:
                plan["solo_plataforma"].append({k2: b.get(k2) for k2 in ("id", "banda", "nombre", "desde", "hasta")})
    plan["bandas_nuevas"] = sorted({b["banda"] for b in plan["crear"] if norm(b["banda"]) not in bandas_plataforma})
    return plan


def mostrar(plan: dict, excel: dict) -> None:
    por_mes = collections.Counter(b["mes"] for b in plan["crear"])
    print("== PLAN ==")
    print("meses leidos:", ", ".join(f"{m}:{len(v)}" for m, v in sorted(excel.items())))
    print("iguales (nada que hacer):", plan["iguales"])
    print("a crear:", len(plan["crear"]), "->", dict(sorted(por_mes.items())))
    print("a cambiar:", len(plan["cambiar"]))
    for c in plan["cambiar"]:
        print(f"    {c['mes']} {c['banda']} | {c['antes']['nombre'][:40]!r} {c['antes']['desde']}..{c['antes']['hasta']}"
              f" -> {c['cambios']}  [{c['motivo']}]")
    print("en la plataforma pero no en el Excel (se dejan):", len(plan["solo_plataforma"]))
    for s in plan["solo_plataforma"]:
        print("   ", s)
    print("bandas que la plataforma todavia no tiene:", plan["bandas_nuevas"])


class Api:
    def __init__(self) -> None:
        self.email, self.password = os.environ.get("MKTG_EMAIL"), os.environ.get("MKTG_PASSWORD")
        if not self.email or not self.password:
            print("ERROR: faltan MKTG_EMAIL y MKTG_PASSWORD (un usuario con calendario.edit)")
            sys.exit(1)
        self.cli = httpx.Client(timeout=120)
        self.h: dict[str, str] = {}
        self.login()

    def login(self) -> None:
        r = self.cli.post(f"{API}/auth/login", json={"email": self.email, "password": self.password})
        r.raise_for_status()
        self.h = {"Authorization": f"Bearer {r.json()['access_token']}"}

    def pedir(self, metodo: str, path: str, **kw) -> httpx.Response:
        r = None
        for _ in range(6):
            r = self.cli.request(metodo, API + path, headers=self.h, **kw)
            if r.status_code == 429:
                time.sleep(20)
                continue
            if r.status_code == 401:
                self.login()
                continue
            return r
        return r


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--aplicar", action="store_true", help="sin esto solo muestra el plan")
    ap.add_argument("--meses", default="", help="YYYY-MM,YYYY-MM para limitar")
    args = ap.parse_args()
    filtro = {m.strip() for m in args.meses.split(",") if m.strip()} or None

    if not imp.RUTAS["comercial"].exists():
        print("No encuentro el Excel:", imp.RUTAS["comercial"])
        return 1

    api = Api()
    datos = api.pedir("GET", "/calendario/datos").json()
    existentes = list(datos["barras"])
    excel = leer_excel(filtro)
    plan = armar_plan(excel, existentes)
    mostrar(plan, excel)

    if not args.aplicar:
        print("\n(simulacion) sin --aplicar no se escribe nada")
        return 0

    BACKUPS.mkdir(exist_ok=True)
    marca = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    (BACKUPS / f"calendario_datos_pre_carga_{marca}.json").write_text(
        json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    (BACKUPS / f"calendario_plan_{marca}.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=1), encoding="utf-8")

    errores, hechas = [], {"creadas": 0, "ya_existian": 0, "cambiadas": 0}
    for c in plan["cambiar"]:
        r = api.pedir("PATCH", f"/calendario/barras/{c['id']}", json=c["cambios"])
        if r.status_code >= 300:
            errores.append(("PATCH", c["id"], r.status_code, r.text[:200]))
        else:
            hechas["cambiadas"] += 1
        time.sleep(PAUSA)
    for i, b in enumerate(plan["crear"], 1):
        payload = {k: b[k] for k in ("id", "seccion", "banda", "carril", "nombre", "color", "desde", "hasta")}
        r = api.pedir("POST", "/calendario/barras", json=payload)
        if r.status_code >= 300:
            errores.append(("POST", b["id"], r.status_code, r.text[:200]))
        elif r.json().get("yaExistia"):
            hechas["ya_existian"] += 1
        else:
            hechas["creadas"] += 1
        if i % 100 == 0:
            print(f"   ... {i}/{len(plan['crear'])}")
        time.sleep(PAUSA)

    print("\n== HECHO ==", hechas, "| errores:", len(errores))
    for e in errores[:20]:
        print("   ", e)
    despues = api.pedir("GET", "/calendario/datos").json()
    cnt = collections.Counter(b["desde"][:7] for b in despues["barras"] if b["seccion"] == "comercial")
    print("comercial por mes ahora:", dict(sorted(cnt.items())))
    return 0 if not errores else 1


if __name__ == "__main__":
    sys.exit(main())
