"""Banco de verdad: qué dibuja PowerPoint DE VERDAD contra lo que el repo predice.

-----------------------------------------------------------------------------
Por qué existe
-----------------------------------------------------------------------------

El achique de las cenefas depende de medir texto, y esa medición está escrita
DOS VECES, a mano:

  - el exportador (Python) mide con la tabla de anchos
    ``app/data/font_metrics.json``, vía ``font_metrics.py``;
  - el preview (TypeScript) mide con el ``ctx.measureText()`` del navegador,
    en ``frontend/lib/cenefas/dibujarTextoEnriquecido.ts``.

Ninguna de las dos le preguntó nunca a PowerPoint, que es el único que importa
porque es el que imprime. Cada vez que una regla se escribió en un lado y no en
el otro, el preview mintió: la voladita (c92b482) y Arial Black medida como
Arial (673aa56) son los dos casos ya documentados. Es el mismo error dos veces.

Este script es la tercera pata: la MEDICIÓN REAL. Toma una plantilla de
producción, la renderiza con el motor del repo, la exporta a PDF con PowerPoint
de verdad (no LibreOffice: achica la voladita a 0,580 en vez de 0,667, o sea
que como patrón mentiría), lee cada pedazo de texto dibujado con PyMuPDF y lo
compara contra lo que el repo había predicho para ESE MISMO pedazo.

Lo que devuelve no es un "pasa/no pasa": es un MAPA de divergencias por
tipografía y por cuadro. Sirve para dos cosas distintas:

  1. saber qué fuentes están mal medidas hoy y cuánto, con números;
  2. correrlo de nuevo cuando alguien agregue un cuadro o una plantilla, y que
     la divergencia aparezca sola en vez de descubrirse en góndola.

-----------------------------------------------------------------------------
QUÉ CUBRE Y QUÉ NO  (leer antes de confiar en un resultado)
-----------------------------------------------------------------------------

CUBRE la mitad de abajo del problema: PowerPoint contra el EXPORTADOR. O sea,
la tabla ``font_metrics.json`` y ``pt_efectivo`` contra la realidad impresa.

NO CUBRE la mitad de arriba: el PREVIEW del navegador contra esa misma tabla.
El preview no mide con la tabla, mide con ``ctx.measureText()`` de Chrome, y
para compararlo hace falta un browser (Playwright o similar) que hoy no está en
este entorno. Mientras esa mitad no exista, que este script dé todo verde
significa "el archivo que se baja es fiel", NO "la pantalla no miente".

La forma definitiva de cerrar el agujero es que las dos implementaciones dejen
de ser dos: que el preview pida los anchos a la misma tabla en vez de medirlos
por su cuenta. Este script es lo que permite afirmar que esa tabla es correcta;
sin él, unificar sería mudar el error de lugar.

-----------------------------------------------------------------------------
Uso
-----------------------------------------------------------------------------

    python backend/scripts/verificar_preview_contra_powerpoint.py
    python backend/scripts/verificar_preview_contra_powerpoint.py "Rompe Precios-202608-A4"
    python backend/scripts/verificar_preview_contra_powerpoint.py --csv salida.csv

Requiere: Windows con PowerPoint instalado, pywin32, PyMuPDF, y el .env del
backend apuntando a la base (SOLO SE LEE: la única consulta es un SELECT).

Abrir PowerPoint por COM es lento y deja procesos colgados si algo explota, así
que la aplicación se abre UNA vez para toda la corrida y se cierra siempre en un
finally. Y si una plantilla falla, se anota el error y se sigue con la
siguiente: una corrida de veintipico de plantillas no se puede perder entera
porque una tenga un cuadro roto.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import io
import json
import os
import pathlib
import re
import sys
import tempfile
import traceback
from dataclasses import dataclass, field

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from app.services.cenefas.font_metrics import (  # noqa: E402
    FACTOR_VOLADITA,
    ancho_texto_cm,
    fuentes_conocidas,
    pt_efectivo,
)
from app.services.cenefas.layout_engine import get_format  # noqa: E402
from app.services.cenefas.variables import DECIMAL_VARS, PRICE_VARS  # noqa: E402

EMU_POR_PT = 12700.0
PT_POR_CM = 72.0 / 2.54


# ---------------------------------------------------------------------------
# Los productos de prueba
# ---------------------------------------------------------------------------
#
# Los valores son CONOCIDOS y de largo variado a propósito. Un cartel de prueba
# con "123" en todos lados no prueba nada: el error de medición crece con el
# ancho, así que lo que rompe es el precio largo y la descripción larga, no el
# caso cómodo.
#
# En los precios se usan dígitos ANCHOS (6 y 9, los dos que más miden en
# Impact) y también el "1", que es el más angosto: si la tabla de anchos
# tuviera los dígitos mal, la diferencia entre "1.111" y "9.669" es lo que la
# delata. Ver `digito_mas_ancho` en font_metrics.py.
_VALORES_PRECIO = ("99", "669", "1.969", "16.996")
_VALORES_DECIMAL = (",90", ",99", ",50", ",95")

# Textos de largo creciente. El tercero es a propósito más largo que cualquier
# caja: si el motor no achica, PowerPoint lo va a partir en renglones y el
# script lo va a ver -- que es justamente el caso que llegó a góndola.
_VALORES_TEXTO = {
    "descripcion": (
        "YERBA CANARIAS 1 KG",
        "GALLETITAS SURTIDAS FAMILIARES CHOCOLATE 500 G",
        "DETERGENTE LIQUIDO CONCENTRADO PARA ROPA DELICADA PERFUME LAVANDA 3 LT",
        "QUESO",
    ),
    "legales": (
        "Promocion valida por tiempo limitado. No acumulable con otras promociones.",
    ),
    "vigencia": ("DEL 18 AL 24 DE SETIEMBRE",),
    "tipoOferta": ("2x1", "25% OFF", "6x4", "SOLO X 25"),
    "tipoOfertaComprando": ("Comprando 2",),
    "unidad": ("la unidad", "el kilo"),
    "mecanica": ("$75,33 la unidad.",),
    "unidadMoneda": ("$",),
    "categoria": ("ALMACEN",),
    "marca": ("CANARIAS",),
    "sku": ("1234567",),
}
# Cualquier variable de texto que no esté arriba cae acá. Tres largos, para que
# una plantilla de 3 o 6 celdas ejercite los tres en la misma hoja.
_TEXTO_GENERICO = ("CORTO", "TEXTO DE LARGO MEDIO", "TEXTO BASTANTE MAS LARGO QUE LOS OTROS DOS")


def productos_de_prueba(definicion: dict, celdas: int) -> list[dict]:
    """Un producto por celda de la hoja, con valores conocidos y de largo variado.

    Se llenan TODAS las variables que declara la plantilla, incluso las que el
    diseño podría ocultar por regla: si una regla la esconde, el cuadro no se
    dibuja y el script simplemente no ve ese pedazo. Lo que no se puede hacer es
    dejar variables vacías, porque entonces medio cartel no se mide.
    """
    nombres = [v.get("name") for v in (definicion.get("variables") or []) if v.get("name")]
    productos = []
    for i in range(max(1, celdas)):
        prod: dict[str, str] = {}
        for n in nombres:
            if n in PRICE_VARS:
                prod[n] = _VALORES_PRECIO[i % len(_VALORES_PRECIO)]
            elif n in DECIMAL_VARS:
                prod[n] = _VALORES_DECIMAL[i % len(_VALORES_DECIMAL)]
            else:
                opciones = _VALORES_TEXTO.get(n, _TEXTO_GENERICO)
                prod[n] = opciones[i % len(opciones)]
        productos.append(prod)
    return productos


# ---------------------------------------------------------------------------
# Lo que el REPO predice: se lee del PPTX ya generado, no de la definición
# ---------------------------------------------------------------------------
#
# Es a propósito. Entre la definición y el archivo hay un camino largo --reglas,
# ocultamientos, smart_bold, el tamaño manual, los overrides por formato-- y
# reproducirlo acá sería escribir la medición una TERCERA vez, que es
# exactamente la enfermedad que este script viene a diagnosticar.
#
# Así que la predicción sale de los runs del PPTX: cada run declara su cuerpo,
# su familia, su negrita y su voladita, y eso es literalmente lo que el
# exportador le dijo a PowerPoint. Comparar eso contra lo dibujado mide la
# brecha entre lo que el repo CREE que pidió y lo que PowerPoint HACE.

@dataclass
class Pedazo:
    """Un run del PPTX: un pedazo de texto con un estilo homogéneo."""
    hoja: int
    cuadro: str
    # Ordinal del cuadro dentro de la hoja. Dos cuadros pueden llamarse igual
    # (el importer repite "CuadroTexto 5" en las celdas de una 3xA4), así que el
    # nombre no sirve para agrupar los runs de UN cuadro.
    idx_cuadro: int
    texto: str
    pt_declarado: float | None
    voladita: int
    familia: str | None
    negrita: bool
    # Rectángulo del cuadro que lo contiene, en puntos de la hoja.
    caja: tuple[float, float, float, float]

    @property
    def pt_predicho(self) -> float | None:
        """El cuerpo con el que el repo cree que se va a dibujar."""
        return pt_efectivo(self.pt_declarado, self.voladita)


_NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _baseline(run) -> int:
    """La voladita del run, tal como la escribe `_apply_run_style` (atributo crudo).

    python-pptx no expone la voladita en Font, así que hay que bajar al XML —
    lo mismo que hace el exportador para escribirla.
    """
    rPr = run._r.find(_NS_A + "rPr")
    if rPr is None:
        return 0
    try:
        return int(rPr.get("baseline") or 0)
    except (TypeError, ValueError):
        return 0


def _shapes_planos(shapes, dx=0.0, dy=0.0, ex=1.0, ey=1.0):
    """Todos los shapes, entrando a los grupos y arrastrando su transformación.

    Un grupo mueve y escala a sus hijos: las coordenadas del hijo están en el
    sistema del grupo (chOff/chExt) y hay que llevarlas al de la hoja, o el
    rectángulo queda en cualquier lado y con él el emparejamiento por geometría
    de más abajo.
    """
    for sh in shapes:
        if getattr(sh, "shape_type", None) == 6 and hasattr(sh, "shapes"):  # GROUP
            try:
                gex = sh.width / sh.child_extents.cx if sh.child_extents.cx else 1.0
                gey = sh.height / sh.child_extents.cy if sh.child_extents.cy else 1.0
                ndx = dx + (sh.left - sh.child_offset.x * gex) * ex
                ndy = dy + (sh.top - sh.child_offset.y * gey) * ey
                yield from _shapes_planos(sh.shapes, ndx, ndy, ex * gex, ey * gey)
            except Exception:
                continue
        else:
            yield sh, dx, dy, ex, ey


def pedazos_del_pptx(pptx_bytes: bytes) -> tuple[list[Pedazo], float, float]:
    """Los runs de cada cuadro de texto, con su caja en puntos de la hoja."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx_bytes))
    ancho_pt = prs.slide_width / EMU_POR_PT
    alto_pt = prs.slide_height / EMU_POR_PT

    pedazos: list[Pedazo] = []
    for n, slide in enumerate(prs.slides, start=1):
        for idx, (sh, dx, dy, ex, ey) in enumerate(_shapes_planos(slide.shapes)):
            if not getattr(sh, "has_text_frame", False):
                continue
            try:
                x0 = (dx + sh.left * ex) / EMU_POR_PT
                y0 = (dy + sh.top * ey) / EMU_POR_PT
                caja = (x0, y0,
                        x0 + sh.width * ex / EMU_POR_PT,
                        y0 + sh.height * ey / EMU_POR_PT)
            except Exception:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if not (run.text or "").strip():
                        continue
                    pedazos.append(Pedazo(
                        hoja=n,
                        cuadro=(sh.name or "")[:40],
                        idx_cuadro=idx,
                        texto=run.text,
                        pt_declarado=run.font.size.pt if run.font.size else None,
                        voladita=_baseline(run),
                        familia=run.font.name,
                        negrita=bool(run.font.bold),
                        caja=caja,
                    ))
    return pedazos, ancho_pt, alto_pt


# ---------------------------------------------------------------------------
# Lo que PowerPoint DIBUJA
# ---------------------------------------------------------------------------

class PowerPoint:
    """PowerPoint real, abierto una sola vez para toda la corrida.

    Se usa como context manager porque lo único inaceptable es dejar el proceso
    colgado: si el script muere a mitad de camino, POWERPNT.EXE se queda en
    memoria y la próxima corrida abre otro encima.
    """

    FORMATO_PDF = 32  # ppSaveAsPDF

    def __enter__(self):
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        self.app = win32com.client.Dispatch("PowerPoint.Application")
        return self

    def __exit__(self, *exc):
        try:
            self.app.Quit()
        except Exception:
            pass
        try:
            import pythoncom
            pythoncom.CoUninitialize()
        except Exception:
            pass
        return False

    def a_pdf(self, pptx: pathlib.Path, pdf: pathlib.Path) -> None:
        pres = self.app.Presentations.Open(str(pptx), ReadOnly=-1, Untitled=0, WithWindow=0)
        try:
            pres.SaveAs(str(pdf), self.FORMATO_PDF)
        finally:
            try:
                pres.Close()
            except Exception:
                pass


@dataclass
class Dibujado:
    """Un span del PDF: un pedazo que PowerPoint dibujó de verdad."""
    hoja: int
    texto: str
    pt: float
    fuente: str
    bbox: tuple[float, float, float, float]

    @property
    def ancho_cm(self) -> float:
        return (self.bbox[2] - self.bbox[0]) / PT_POR_CM


def spans_del_pdf(pdf: pathlib.Path, ancho_hoja_pt: float) -> list[Dibujado]:
    """Cada pedazo de texto del PDF, normalizado a puntos de la HOJA del PPTX.

    PowerPoint a veces exporta la página con un tamaño distinto al de la
    diapositiva (redondeos de la hoja A4, sobre todo). Sin normalizar, todos los
    anchos salen corridos por el mismo porcentaje y el informe acusa a la tabla
    de un error que en realidad es de la exportación.
    """
    import fitz

    salida: list[Dibujado] = []
    doc = fitz.open(pdf)
    try:
        for n, page in enumerate(doc, start=1):
            escala = (ancho_hoja_pt / page.rect.width) if page.rect.width else 1.0
            for bloque in page.get_text("dict").get("blocks", []):
                for linea in bloque.get("lines", []):
                    for span in linea.get("spans", []):
                        if not (span.get("text") or "").strip():
                            continue
                        b = span["bbox"]
                        salida.append(Dibujado(
                            hoja=n,
                            texto=span["text"],
                            pt=span["size"] * escala,
                            # El PDF le pone un prefijo de seis letras al nombre
                            # de toda fuente incrustada ("ABCDEF+Impact").
                            fuente=re.sub(r"^[A-Z]{6}\+", "", span.get("font", "")),
                            bbox=(b[0] * escala, b[1] * escala,
                                  b[2] * escala, b[3] * escala),
                        ))
    finally:
        doc.close()
    return salida


# ---------------------------------------------------------------------------
# Emparejar un span del PDF con el run que lo originó
# ---------------------------------------------------------------------------
#
# No hay identificador que los una: hay que deducirlo. Dos señales, y hacen
# falta las dos:
#
#   TEXTO     -- el texto del span tiene que aparecer en el del cuadro.
#   GEOMETRÍA -- con "2x1" en la cocarda y "2x1" tapando el precio, el texto
#                solo no alcanza; gana el cuadro cuyo rectángulo esté más cerca.
#
# Se compara sin espacios: PowerPoint se come el espacio donde cortó el
# renglón, así que el texto del span no es literalmente una subcadena del run.
#
# El emparejamiento es CONTRA EL CUADRO, no contra el run, porque el corte no
# coincide en ninguna de las dos direcciones:
#
#   - un run largo se parte en varios spans, uno por renglón;
#   - y al revés: tres runs seguidos con el MISMO estilo ("$" + "99" + ",90" en
#     el precio regular) los junta en UN span. Emparejando run por run ese span
#     quedaba huérfano y el precio regular no se medía nunca -- justo uno de los
#     cuadros donde el desborde se ve.
#
# Así que se arma una tira por cuadro, con el run dueño de cada caracter. Un
# span que cae sobre varios runs de estilo distinto no se puede juzgar (no hay
# un cuerpo único que comparar) y se cuenta aparte en vez de inventarle uno.

@dataclass
class _Cuadro:
    pedazos: list[Pedazo] = field(default_factory=list)
    cadena: str = ""
    duenio: list[int] = field(default_factory=list)
    cursor: int = 0

    def agregar(self, p: Pedazo) -> None:
        i = len(self.pedazos)
        self.pedazos.append(p)
        limpio = _sin_espacios(p.texto)
        self.cadena += limpio
        self.duenio.extend([i] * len(limpio))


def _sin_espacios(s: str) -> str:
    return re.sub(r"\s+", "", s)


def _distancia(caja, bbox) -> float:
    """Cuánto se aleja el span de su caja. 0 si el centro del span cae adentro."""
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    dx = max(caja[0] - cx, 0.0, cx - caja[2])
    dy = max(caja[1] - cy, 0.0, cy - caja[3])
    return (dx * dx + dy * dy) ** 0.5


@dataclass
class Comparacion:
    plantilla: str
    hoja: int
    cuadro: str
    texto: str
    familia_declarada: str
    familia_dibujada: str
    negrita: bool
    pt_declarado: float | None
    pt_predicho: float | None
    pt_dibujado: float
    ancho_predicho_cm: float
    ancho_dibujado_cm: float

    @property
    def medible(self) -> bool:
        """Si el repo pudo siquiera predecir un ancho para este pedazo.

        Un run sin cuerpo declarado (el exportador no le escribió `sz` porque el
        componente no traía font_size) hereda el tamaño del layout: PowerPoint
        sabe cuál es, el repo no. Ahí no hay divergencia que medir -- hay un
        agujero, y es peor, porque el motor de achique tampoco puede medirlo.
        """
        return self.pt_predicho is not None

    @property
    def div_cuerpo_pct(self) -> float | None:
        if not self.pt_predicho or not self.pt_dibujado:
            return None
        return (self.pt_predicho - self.pt_dibujado) / self.pt_dibujado * 100

    @property
    def div_ancho_pct(self) -> float | None:
        if not self.ancho_dibujado_cm or not self.medible:
            return None
        return (self.ancho_predicho_cm - self.ancho_dibujado_cm) / self.ancho_dibujado_cm * 100

    @property
    def div_ancho_cm(self) -> float:
        if not self.medible:
            return 0.0
        return self.ancho_predicho_cm - self.ancho_dibujado_cm

    @property
    def dudoso(self) -> bool:
        """El cuerpo dibujado no es el declarado: o el par está mal, o falta una regla.

        PowerPoint respeta el cuerpo declarado con un error del orden del 0,1 %.
        Cuando no lo respeta pasa una de dos cosas, y desde acá no se distinguen:

          - el script emparejó mal (un span de un caracter suelto, "9" o "$",
            que existe en varios cuadros de la misma hoja);
          - o PowerPoint aplica una regla de dibujo que el repo NO modela, que
            es exactamente cómo apareció la voladita en su momento.

        Por eso no se descartan en silencio: salen listados aparte para que
        alguien los mire, y quedan fuera de las estadísticas de ancho, donde
        serían ruido.
        """
        d = self.div_cuerpo_pct
        return d is not None and abs(d) > 2.0


def comparar(nombre: str, pedazos: list[Pedazo],
             spans: list[Dibujado]) -> tuple[list[Comparacion], int, int]:
    """Empareja y compara.

    Devuelve (comparaciones, spans sin pareja, spans de estilo mezclado).
    """
    cuadros: dict[tuple[int, int], _Cuadro] = {}
    for p in pedazos:
        cuadros.setdefault((p.hoja, p.idx_cuadro), _Cuadro()).agregar(p)

    salida: list[Comparacion] = []
    huerfanos = 0
    mezclados = 0
    for span in spans:
        clave = _sin_espacios(span.texto)
        if not clave:
            continue
        candidatos = []
        for (hoja, _), c in cuadros.items():
            if hoja != span.hoja:
                continue
            desde = c.cadena.find(clave, c.cursor)
            if desde < 0:
                continue
            cubiertos = {c.duenio[i] for i in range(desde, desde + len(clave))}
            p = c.pedazos[min(cubiertos)]
            # Desempate por CUERPO antes que por distancia.
            #
            # Un span de un caracter ("9", "$", "G") aparece en media docena de
            # cuadros y la distancia sola elegía mal: quedaban emparejamientos
            # con el cuerpo al doble o a la mitad, que después el informe leía
            # como si PowerPoint hubiera dibujado otro tamaño. PowerPoint
            # respeta el cuerpo declarado con un error del orden del 0,1 %, así
            # que "el cuerpo coincide" es una señal de identidad casi perfecta.
            #
            # No es circular: si NINGÚN candidato coincide en cuerpo, se elige
            # igual por distancia y la divergencia se informa. Esa es la puerta
            # por la que en su momento apareció la voladita, y tiene que seguir
            # abierta: una regla de dibujo que el repo todavía no modela se ve
            # exactamente así.
            coincide = (p.pt_predicho is not None and span.pt
                        and abs(p.pt_predicho - span.pt) / span.pt <= 0.02)
            candidatos.append((0 if coincide else 1,
                               _distancia(p.caja, span.bbox), c, desde, cubiertos, p))
        if not candidatos:
            # Texto que está en el PDF pero no en ningún run de la diapositiva:
            # casi siempre vive en el LAYOUT o el MASTER del PPTX original
            # (rótulos del arte, "OFERTA", "PRECIO REGULAR:"), que python-pptx no
            # devuelve en slide.shapes. Es diseño, no dato: no lo llena ninguna
            # variable y el motor de achique nunca lo toca. Se cuenta y se sigue.
            huerfanos += 1
            continue
        _, _, cuadro, desde, cubiertos, p = min(candidatos, key=lambda t: (t[0], t[1]))
        cuadro.cursor = desde + len(clave)

        if len({(cuadro.pedazos[i].pt_predicho, cuadro.pedazos[i].familia,
                 cuadro.pedazos[i].negrita) for i in cubiertos}) > 1:
            mezclados += 1
            continue

        pt_pred = p.pt_predicho
        # El ancho se predice con el cuerpo que el repo CREE (pt_predicho): así
        # la divergencia de ancho arrastra el error de la voladita, que es
        # justamente como lo sufre el motor de achique cuando decide si algo
        # entra en su caja.
        ancho_pred = (
            ancho_texto_cm(span.texto, pt_pred, p.familia, p.negrita) if pt_pred else 0.0
        )
        salida.append(Comparacion(
            plantilla=nombre,
            hoja=span.hoja,
            cuadro=p.cuadro,
            texto=span.texto,
            familia_declarada=p.familia or "(sin declarar)",
            familia_dibujada=span.fuente,
            negrita=p.negrita,
            pt_declarado=p.pt_declarado,
            pt_predicho=pt_pred,
            pt_dibujado=span.pt,
            ancho_predicho_cm=ancho_pred,
            ancho_dibujado_cm=span.ancho_cm,
        ))
    return salida, huerfanos, mezclados


# ---------------------------------------------------------------------------
# Traer las plantillas de la base — SOLO LECTURA
# ---------------------------------------------------------------------------

async def _traer(nombre: str | None) -> list[dict]:
    import asyncpg
    from dotenv import load_dotenv

    # override=True a propósito: el entorno de desarrollo exporta un
    # DATABASE_URL de mentira (localhost/test_dummy) para poder importar la app
    # sin base, y load_dotenv por defecto NO pisa lo que ya está exportado --
    # así que sin esto el script se conecta a la nada. Acá la base que vale es
    # la del .env, siempre, y se usa nada más que para leer.
    load_dotenv(pathlib.Path(__file__).resolve().parent.parent / ".env", override=True)
    dsn = re.sub(r"^postgresql\+\w+://", "postgresql://",
                 os.getenv("DATABASE_URL", "")).split("?")[0]
    if not dsn:
        raise SystemExit("ERROR: falta DATABASE_URL en backend/.env")

    # statement_cache_size=0: la base está detrás del pooler de Supabase en modo
    # transacción, que no soporta prepared statements con nombre.
    conn = await asyncpg.connect(dsn, timeout=40, statement_cache_size=0)
    try:
        sql = ("select name, category, formats, definition, source_pptx "
               "from cenefa_templates_v2 ")
        if nombre:
            filas = await conn.fetch(sql + "where name = $1 order by created_at", nombre)
            if not filas:
                filas = await conn.fetch(sql + "where name ilike $1 order by created_at",
                                         f"%{nombre}%")
        else:
            # El destino `pruebas` queda afuera: no es producción y lo que se
            # esté probando ahí puede tener el autoajuste prendido, que cambia
            # el cuerpo dibujado y ensuciaría el mapa (ver pruebas.py).
            filas = await conn.fetch(
                sql + "where coalesce(category,'') <> 'pruebas' order by category, name")
        return [dict(f) for f in filas]
    finally:
        await conn.close()


# ---------------------------------------------------------------------------
# Una plantilla, de punta a punta
# ---------------------------------------------------------------------------

def revisar(fila: dict, pp: PowerPoint,
            carpeta: pathlib.Path) -> tuple[list[Comparacion], str]:
    from app.services.cenefas.component_renderer import render_template_to_pptx

    nombre = fila["name"]
    definicion = fila["definition"]
    if isinstance(definicion, str):
        definicion = json.loads(definicion)
    # `category` vive en la columna y adentro del JSON falta en varias
    # plantillas: el motor la necesita estampada para decidir aislamientos
    # (ver pruebas.es_de_pruebas). Se replica lo que hace _resolve_template_v2.
    definicion.setdefault("category", fila.get("category"))

    formato = (list(fila.get("formats") or []) or ["a4"])[0]
    celdas = get_format(formato)["slots"]
    productos = productos_de_prueba(definicion, celdas)

    pptx_bytes, faltantes = render_template_to_pptx(
        definicion, productos, formato, None, fila.get("source_pptx")
    )
    seguro = re.sub(r"[^A-Za-z0-9_-]+", "_", nombre)[:60]
    ruta_pptx = carpeta / f"{seguro}.pptx"
    ruta_pdf = carpeta / f"{seguro}.pdf"
    ruta_pptx.write_bytes(pptx_bytes)

    pedazos, ancho_pt, _ = pedazos_del_pptx(pptx_bytes)
    pp.a_pdf(ruta_pptx, ruta_pdf)
    spans = spans_del_pdf(ruta_pdf, ancho_pt)
    comps, huerfanos, mezclados = comparar(nombre, pedazos, spans)

    nota = (f"{formato:<6} {len(pedazos):>3} runs | {len(spans):>3} dibujados | "
            f"{len(comps):>3} emparejados")
    if huerfanos:
        nota += f" | {huerfanos} sin pareja"
    if mezclados:
        nota += f" | {mezclados} de estilo mezclado"
    if faltantes:
        nota += f" | faltan vars: {','.join(sorted(faltantes))[:40]}"
    return comps, nota


# ---------------------------------------------------------------------------
# El informe
# ---------------------------------------------------------------------------

def _pct(v):
    return "     -" if v is None else f"{v:+6.1f}"


def _med(vs):
    return f"{sum(vs) / len(vs):+10.1f}" if vs else "         -"


def _peor(vs):
    return f"{max(vs, key=abs):+10.1f}" if vs else "         -"


def _promedio_abs(cs: list[Comparacion]) -> float:
    vs = [abs(c.div_ancho_pct) for c in cs if c.div_ancho_pct is not None]
    return sum(vs) / len(vs) if vs else 0.0


# Un nombre de fuente degenerado en el PDF ("7", "23", "2,Bold") es la firma de
# que PowerPoint NO tenía instalada la que el diseño pedía y la sustituyó por
# otra sin poder nombrarla. Ahí lo dibujado no es la fuente declarada, así que
# ni la tabla ni el preview pueden acertarle: es un problema anterior a medir.
_RE_FUENTE_DEGENERADA = re.compile(r"^\d+(,\w+)?$")


def informar(todas: list[Comparacion]) -> None:
    if not todas:
        print("\nNo se pudo comparar nada.")
        return

    sin_cuerpo = [c for c in todas if not c.medible]
    dudosos = [c for c in todas if c.medible and c.dudoso]
    medibles = [c for c in todas if c.medible and not c.dudoso]

    print("\n" + "=" * 118)
    print("PEORES DIVERGENCIAS DE ANCHO  (lo que el repo predice vs lo que PowerPoint dibujó)")
    print("=" * 118)
    print(f"{'plantilla':<30} {'cuadro':<20} {'texto':<22} {'fuente real':<18} "
          f"{'pred':>7} {'real':>7} {'dif%':>6} {'dif cm':>7}")
    for c in sorted(medibles, key=lambda c: -abs(c.div_ancho_cm))[:35]:
        print(f"{c.plantilla[:29]:<30} {c.cuadro[:19]:<20} {c.texto[:21]:<22} "
              f"{c.familia_dibujada[:17]:<18} {c.ancho_predicho_cm:>7.2f} "
              f"{c.ancho_dibujado_cm:>7.2f} {_pct(c.div_ancho_pct)} {c.div_ancho_cm:>+7.2f}")

    # --- Por tipografía: el mapa que de verdad sirve --------------------------
    #
    # Un cuadro puntual mal medido se arregla a mano. Una FUENTE mal medida
    # rompe todos los cuadros que la usan, incluidos los que todavía no
    # existen -- que es el problema que hay que cerrar.
    conocidas = set(fuentes_conocidas())
    por_fuente: dict[tuple[str, str], list[Comparacion]] = {}
    for c in medibles:
        por_fuente.setdefault((c.familia_dibujada, c.familia_declarada), []).append(c)

    print("\n" + "=" * 118)
    print("MAPA POR TIPOGRAFÍA  (la que PowerPoint dibujó / la que el PPTX declara)")
    print("=" * 118)
    print(f"{'fuente dibujada':<24} {'declarada':<26} {'en tabla':<9} {'n':>4} "
          f"{'ancho med%':>11} {'ancho peor%':>12} {'cuerpo med%':>12}")
    for (dibujada, declarada), cs in sorted(por_fuente.items(),
                                            key=lambda kv: -_promedio_abs(kv[1])):
        anchos = [c.div_ancho_pct for c in cs if c.div_ancho_pct is not None]
        cuerpos = [c.div_cuerpo_pct for c in cs if c.div_cuerpo_pct is not None]
        en_tabla = "si" if (declarada or "").strip().lower() in conocidas else "NO"
        if _RE_FUENTE_DEGENERADA.match(dibujada):
            en_tabla = "SUSTIT."
        print(f"{dibujada[:23]:<24} {declarada[:25]:<26} {en_tabla:<9} {len(cs):>4} "
              f"{_med(anchos):>11} {_peor(anchos):>12} {_med(cuerpos):>12}")

    # --- Fuentes que PowerPoint no tenía -------------------------------------
    sustituidas = sorted({c.familia_declarada for c in todas
                          if _RE_FUENTE_DEGENERADA.match(c.familia_dibujada)})
    if sustituidas:
        print("\n" + "=" * 118)
        print("FUENTES QUE POWERPOINT NO TENÍA Y SUSTITUYÓ")
        print("=" * 118)
        print("El PDF les puso un nombre degenerado ('7', '23', '2,Bold'): la que se")
        print("dibujó NO es la que el diseño pide. Medirlas bien es imposible mientras")
        print("la máquina que exporta no las tenga instaladas.")
        for f in sustituidas:
            n = sum(1 for c in todas if c.familia_declarada == f
                    and _RE_FUENTE_DEGENERADA.match(c.familia_dibujada))
            print(f"  {f:<32} {n:>4} pedazos")

    # --- Runs sin cuerpo declarado -------------------------------------------
    if sin_cuerpo:
        print("\n" + "=" * 118)
        print("PEDAZOS QUE EL REPO NO PUEDE MEDIR (el run no declara cuerpo)")
        print("=" * 118)
        print("El componente no trae font_size, así que el exportador no le escribe `sz`")
        print("y el cuerpo lo pone el layout. PowerPoint lo dibuja igual; el motor de")
        print("achique se queda sin número y da el cuadro por bueno sin haberlo medido.")
        for c in sin_cuerpo:
            print(f"  {c.plantilla[:30]:<31} {c.cuadro[:20]:<21} {c.texto[:24]:<25} "
                  f"dibujado a {c.pt_dibujado:>6.1f} pt en {c.familia_dibujada}")

    # --- La negrita, que también es UN número global --------------------------
    #
    # font_metrics ensancha un 8 % (_FACTOR_BOLD) cualquier texto en negrita, sea
    # la fuente que sea. Es una estimación, nunca se midió, y se aplica a TODAS.
    # Partir las cuentas por familia y por negrita la deja a la vista: si una
    # familia da bien en redonda y mal en negrita por el mismo porcentaje en
    # todas, el que está mal es el factor, no la tabla.
    por_peso: dict[tuple[str, bool], list[float]] = {}
    for c in medibles:
        if c.div_ancho_pct is not None:
            por_peso.setdefault((c.familia_declarada, c.negrita), []).append(c.div_ancho_pct)
    print("\n" + "=" * 118)
    print(f"NEGRITA  (_FACTOR_BOLD = 1.08 en font_metrics, el mismo para todas las fuentes)")
    print("=" * 118)
    print(f"{'familia declarada':<30} {'redonda: n':>12} {'error %':>10}   "
          f"{'negrita: n':>12} {'error %':>10}   {'lo que agrega la negrita':>24}")
    familias = sorted({f for f, _ in por_peso},
                      key=lambda f: -len(por_peso.get((f, False), []) + por_peso.get((f, True), [])))
    for f in familias:
        r = por_peso.get((f, False), [])
        b = por_peso.get((f, True), [])
        mr = sum(r) / len(r) if r else None
        mb = sum(b) / len(b) if b else None
        # Cuánto ensancha la negrita DE VERDAD en esa familia: el 1,08 corregido
        # por lo que cada lado se desvía. Solo tiene sentido con las dos medidas.
        real = (1.08 * (1 + mr / 100) / (1 + mb / 100)) if (mr is not None and mb is not None) else None
        txt_r = "-" if mr is None else f"{mr:+.1f}"
        txt_b = "-" if mb is None else f"{mb:+.1f}"
        print(f"{f[:29]:<30} {len(r):>12} {txt_r:>10}   {len(b):>12} {txt_b:>10}   "
              f"{'-' if real is None else f'x{real:.3f}':>24}")

    # --- Emparejamientos que no cierran por cuerpo ----------------------------
    if dudosos:
        print("\n" + "=" * 118)
        print("PARES DUDOSOS: POWERPOINT DIBUJÓ OTRO CUERPO DEL DECLARADO")
        print("=" * 118)
        print("Cada uno es, o un pedazo que el script emparejó mal (texto de un caracter")
        print("repetido en varios cuadros de la hoja), o una regla de dibujo que el repo")
        print("no modela. Quedan fuera de las cuentas de ancho; hay que mirarlos a ojo.")
        for c in sorted(dudosos, key=lambda c: -abs(c.div_cuerpo_pct or 0)):
            print(f"  {c.plantilla[:28]:<29} {c.cuadro[:18]:<19} {c.texto[:16]:<17} "
                  f"declara {c.pt_declarado:>6.1f} pt -> predice {c.pt_predicho:>6.1f} | "
                  f"dibuja {c.pt_dibujado:>6.1f} pt en {c.familia_dibujada[:18]}")

    # --- La voladita, que es UN número y se puede calibrar --------------------
    volados = [c for c in medibles
               if c.pt_declarado and c.pt_predicho
               and abs(c.pt_predicho - c.pt_declarado) > 0.01]
    print("\n" + "=" * 118)
    print(f"VOLADITA  (FACTOR_VOLADITA = {FACTOR_VOLADITA} en los dos motores)")
    print("=" * 118)
    if not volados:
        print("Ningún pedazo volado en las plantillas revisadas.")
    else:
        factores = [c.pt_dibujado / c.pt_declarado for c in volados if c.pt_declarado]
        medio = sum(factores) / len(factores)
        print(f"{len(volados)} pedazos volados. Factor REAL medido: "
              f"min {min(factores):.4f} | promedio {medio:.4f} | max {max(factores):.4f}")
        print(f"El repo usa {FACTOR_VOLADITA}: {(FACTOR_VOLADITA - medio) / medio * 100:+.1f}% de error.")

    print("\n" + "=" * 118)
    print("RESUMEN")
    print("=" * 118)
    anchos = [c.div_ancho_pct for c in medibles if c.div_ancho_pct is not None]
    cuerpos = [c.div_cuerpo_pct for c in medibles if c.div_cuerpo_pct is not None]
    print(f"pedazos comparados            : {len(todas)}")
    print(f"  sin cuerpo declarado        : {len(sin_cuerpo)}  (no se pueden medir)")
    print(f"  pares dudosos               : {len(dudosos)}  (fuera de las cuentas)")
    if anchos:
        print(f"error de ancho  |medio|       : {sum(abs(a) for a in anchos) / len(anchos):.1f}%")
        print(f"error de ancho  |peor|        : {max(abs(a) for a in anchos):.1f}%")
        malos = sum(1 for a in anchos if abs(a) > 5)
        print(f"pedazos con más de 5% de error: {malos} de {len(anchos)} "
              f"({malos / len(anchos) * 100:.0f}%)")
    if cuerpos:
        print(f"error de cuerpo |medio|       : {sum(abs(a) for a in cuerpos) / len(cuerpos):.1f}%")
    if medibles:
        peor = max(medibles, key=lambda c: abs(c.div_ancho_cm))
        print(f"peor divergencia absoluta     : {peor.div_ancho_cm:+.2f} cm "
              f"({peor.plantilla} / {peor.cuadro} / {peor.texto[:20]!r})")


def _num(v):
    return "" if v is None else round(v, 2)


# ---------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description="Banco de verdad: PowerPoint real vs lo que el repo predice.")
    ap.add_argument("plantilla", nargs="?", help="nombre (exacto o parcial); sin esto, todas")
    ap.add_argument("--csv", help="además, volcar cada pedazo comparado a un CSV")
    ap.add_argument("--carpeta", help="dónde dejar los PPTX y PDF intermedios")
    args = ap.parse_args()

    filas = asyncio.run(_traer(args.plantilla))
    if not filas:
        print("No hay plantillas que revisar.")
        return 1

    carpeta = pathlib.Path(args.carpeta or tempfile.mkdtemp(prefix="banco_verdad_"))
    carpeta.mkdir(parents=True, exist_ok=True)
    print(f"{len(filas)} plantillas | intermedios en {carpeta}\n")

    todas: list[Comparacion] = []
    fallidas: list[tuple[str, str]] = []
    with PowerPoint() as pp:
        for i, fila in enumerate(filas, start=1):
            etiqueta = f"[{i}/{len(filas)}] {fila['name'][:45]:<45}"
            try:
                comps, nota = revisar(fila, pp, carpeta)
                todas.extend(comps)
                print(f"{etiqueta} {nota}", flush=True)
            except Exception as exc:
                # Una plantilla rota no puede voltear la corrida entera: el mapa
                # de las otras veintipico sigue siendo el entregable.
                fallidas.append((fila["name"], f"{type(exc).__name__}: {exc}"))
                print(f"{etiqueta} FALLO -- {type(exc).__name__}: {exc}", flush=True)
                traceback.print_exc(limit=3, file=sys.stderr)

    informar(todas)
    if fallidas:
        print("\nPlantillas que no se pudieron revisar:")
        for n, e in fallidas:
            print(f"  {n}: {e}")

    if args.csv and todas:
        with open(args.csv, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(["plantilla", "hoja", "cuadro", "texto", "familia_declarada",
                        "familia_dibujada", "negrita", "pt_declarado", "pt_predicho", "pt_dibujado",
                        "div_cuerpo_pct", "ancho_predicho_cm", "ancho_dibujado_cm",
                        "div_ancho_pct", "div_ancho_cm"])
            for c in todas:
                w.writerow([c.plantilla, c.hoja, c.cuadro, c.texto, c.familia_declarada,
                            c.familia_dibujada, int(c.negrita), c.pt_declarado, c.pt_predicho,
                            round(c.pt_dibujado, 2), _num(c.div_cuerpo_pct),
                            round(c.ancho_predicho_cm, 3), round(c.ancho_dibujado_cm, 3),
                            _num(c.div_ancho_pct), round(c.div_ancho_cm, 3)])
        print(f"\nCSV: {args.csv}  ({len(todas)} filas)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
