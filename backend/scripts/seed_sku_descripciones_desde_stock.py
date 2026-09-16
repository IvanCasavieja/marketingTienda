"""
Script standalone -- usa asyncpg directo, sin dependencias de la app.
Carga en sku_descripciones los pares (id_producto, descripcion) de un export de
STOCK POR SUCURSAL, el mismo formato que lee el Convertidor desde el 15/09/2026:
una fila por producto Y por sucursal, con el codigo en `id_producto`.
Toca la base de PRODUCCION -- por eso pide confirmacion salvo que se pase --yes.

Por que existe, teniendo ya seed_sku_descripciones_standalone.py: aquel lee
Gestion/Diccionario.xlsx, que trae una fila por SKU y la descripcion de cenefa ya
escrita. Este lee el export de gestion, que repite cada producto una vez por
sucursal y trae la descripcion del ERP. Cambia la forma del archivo; la semantica
de escritura esta copiada a proposito.

SOLO AGREGA, NO PISA. Es lo que pidio Ivan el 16/09/2026, textual: "estan ok para
AGREGAR al diccionario asi tal cual estan". Un SKU que ya tiene descripcion se
saltea, aunque nadie la haya tocado a mano. El motivo es concreto: las filas que
dejaron el seed original y fix_sku_descripciones_casing.py tienen updated_by_id
NULL --ese script hace UPDATE sin escribirlo-- asi que el candado de
"correccion humana" NO las protege, y son descripciones de cenefa BIEN escritas.
Pisarlas con el crudo del ERP seria cambiar texto bueno por texto en MAYUSCULA
con abreviaturas, que es justo lo que el otro seed se niega a persistir a
proposito ("es la descripcion mala", seed_sku_descripciones_standalone.py:52).
Para pisar igual esta --pisar-existentes, que ademas deja backup obligatorio.

SOBRE EL ESTILO: las descripciones del ERP vienen en MAYUSCULA COMPLETA y con
abreviaturas ("AURICULAR PHILIPS INALAM. UPBEAT NEGRO"), y la regla vigente del
catalogo (_STYLE_RULES en convertidor_ai.py) pide marca en MAYUSCULA y el resto
en minuscula tipo oracion. Se cargan tal cual igual, por decision de Ivan: tener
algo es mejor que la fila roja de "falta descripcion". Ojo con una consecuencia:
una vez cargadas NO vuelven a aparecer en el modal de "Generar con IA" (dejan de
tener el warning missing_description), asi que nadie las revisa de nuevo. Para
re-castearlas despues hay que usar fix_sku_descripciones_casing_AI.py y no el
heuristico: el heuristico deja afuera por diseno las filas 100% en mayuscula,
porque no tiene como saber cual era la marca.

Y UNA EXPECTATIVA, PARA QUE NO SORPRENDA: esto NO cambia nada del archivo que lo
motivo. Un listado que trae su propia columna Descripcion gana sobre el catalogo
(decision del 24/08/2026, ver el alias `descripcion` en convertidor.py). El seed
sirve para la PROXIMA oferta de esos productos que venga SIN esa columna.

Uso:
  DATABASE_URL=postgresql+asyncpg://... python backend/scripts/seed_sku_descripciones_desde_stock.py "ruta/al/export.xlsx" --dry-run
  DATABASE_URL=postgresql+asyncpg://... python backend/scripts/seed_sku_descripciones_desde_stock.py "ruta/al/export.xlsx" --yes
  (con --dry-run y sin DATABASE_URL valida solo el archivo, sin tocar la base)
"""
import argparse
import asyncio
import json
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    print("ERROR: pip install openpyxl")
    sys.exit(1)

try:
    import asyncpg
except ImportError:
    print("ERROR: pip install asyncpg")
    sys.exit(1)


_MAX_DESCRIPCION_LEN = 300   # mismo limite que SkuDescripcion.descripcion (String(300))
_MAX_SKU_LEN = 600           # mismo limite que SkuDescripcion.sku (migracion 0049)
_HEADER_SCAN_ROWS = 10       # mismo barrido que detectar_fila_headers en convertidor.py

# Copia textual de _RE_CLAVE_PLURAL (convertidor.py): "123-456" es la clave de un
# GRUPO unificado y vive en otra tabla. upsert_sku_descripcion la rechaza con un
# 400; el seed la tiene que saltear por el mismo motivo, o mete en el catalogo
# singular filas que despues nadie puede editar desde la pantalla del Diccionario.
_RE_CLAVE_PLURAL = re.compile(r"^\d+(?:\s*[-/]\s*\d+)+$")

# Solo `descripcion`. NO se acepta `descripcion web` como respaldo: en la app esa
# columna mapea a descripcionWeb, que es contexto para el prompt de Tinin y no el
# texto del cartel -- y viene larga, asi que cargarla dejaria esos SKU con el
# warning descripcion_larga en todos los listados futuros.
_COL_CODIGO = ("idproducto", "codigo")
_COL_DESCRIPCION = ("descripcion",)


def _norm(name) -> str:
    """Igual que _norm en convertidor.py: sin acentos, sin espacios/guiones, minusculas."""
    s = unicodedata.normalize("NFD", str(name)).encode("ascii", "ignore").decode()
    return re.sub(r"[\s_\-]+", "", s).lower()


def _limpiar(texto: str) -> str:
    """Colapsa espacios raros del ERP. No toca el contenido.

    `.strip()` no alcanza: el export real trae espacios duros (U+00A0) en el medio
    --'MOULINEX\\xa0BY2M0810'-- y dobles espacios. Guardados asi, el PPT no corta
    de linea ahi nunca y buscar el texto con un espacio normal no lo encuentra.
    """
    return " ".join(str(texto).split())


def normalize_sku(raw) -> str | None:
    """int/float/str crudo de la celda -> string canonico ('503996.0' -> '503996')."""
    if raw is None:
        return None
    if isinstance(raw, float):
        raw = int(raw) if raw.is_integer() else raw
    s = str(raw).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s or None


def _abrir_hoja(path: str):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"ERROR: no pude abrir el archivo como .xlsx ({type(e).__name__}: {e})")
        sys.exit(1)
    return wb[wb.sheetnames[0]]


def parse_export(path: str) -> tuple[list[tuple[str, str]], dict]:
    """(pares [(sku, descripcion)], resumen).

    Un export de stock repite cada producto una vez por sucursal: se deduplica por
    SKU y gana la primera aparicion, igual que el otro seed.
    """
    ws = _abrir_hoja(path)
    filas = ws.iter_rows(values_only=True)

    # No se asume la fila 1: el export real puede traer una fila de titulo antes,
    # igual que contempla detectar_fila_headers en convertidor.py.
    col_sku = col_desc = None
    cabecera = None
    for _ in range(_HEADER_SCAN_ROWS):
        try:
            fila = next(filas)
        except StopIteration:
            break
        if not fila:
            continue
        idx: dict[str, int] = {}
        for i, c in enumerate(fila):
            if c is None:
                continue
            k = _norm(c)
            if k and k not in idx:      # gana la PRIMERA: con dos columnas que
                idx[k] = i              # normalizan igual, la de atras no pisa
        s = next((idx[c] for c in _COL_CODIGO if c in idx), None)
        d = next((idx[c] for c in _COL_DESCRIPCION if c in idx), None)
        if s is not None and d is not None:
            col_sku, col_desc, cabecera = s, d, fila
            break

    if col_sku is None or col_desc is None:
        print(f"ERROR: no encontre las columnas en las primeras {_HEADER_SCAN_ROWS} filas.")
        print(f"       busco el codigo en {_COL_CODIGO} y la descripcion en {_COL_DESCRIPCION}")
        sys.exit(1)
    print(f"Encabezados: {[str(c) for c in cabecera if c is not None]}")

    pares: list[tuple[str, str]] = []
    vistos: set[str] = set()
    resumen = {"filas": 0, "sin_codigo": 0, "sin_descripcion": 0, "repetidos": 0,
               "sku_largo": 0, "clave_plural": 0, "truncadas": 0}

    for row in filas:
        resumen["filas"] += 1
        sku = normalize_sku(row[col_sku]) if col_sku < len(row) else None
        desc = _limpiar(row[col_desc]) if col_desc < len(row) and row[col_desc] is not None else ""
        if not sku:
            resumen["sin_codigo"] += 1
            continue
        # Las filas de pie del export ('Total', 'Filtros aplicados:') traen texto en
        # la columna del codigo y NADA en la descripcion: caen aca, igual que
        # cualquier fila sin descripcion. No hace falta reconocerlas por su texto.
        if not desc:
            resumen["sin_descripcion"] += 1
            continue
        if sku in vistos:
            resumen["repetidos"] += 1
            continue
        if _RE_CLAVE_PLURAL.match(sku):
            resumen["clave_plural"] += 1
            continue
        if len(sku) > _MAX_SKU_LEN:
            # Un codigo NO se trunca: cortado apunta a otro producto.
            resumen["sku_largo"] += 1
            continue
        vistos.add(sku)
        if len(desc) > _MAX_DESCRIPCION_LEN:
            resumen["truncadas"] += 1
        pares.append((sku, desc[:_MAX_DESCRIPCION_LEN]))

    return pares, resumen


def _imprimir_resumen_archivo(pares, resumen, path) -> None:
    print(f"\nLeidas {resumen['filas']} filas de {path}")
    print(f"  -> {len(pares)} productos distintos con descripcion")
    print(f"     {resumen['repetidos']} filas repetidas del mismo producto (gana la primera)")
    print(f"     {resumen['sin_descripcion']} sin descripcion, {resumen['sin_codigo']} sin codigo")
    for clave, texto in (("truncadas", f"descripciones truncadas a {_MAX_DESCRIPCION_LEN} caracteres"),
                         ("clave_plural", "salteadas por ser clave de grupo (van en Plurales)"),
                         ("sku_largo", f"salteadas por codigo mas largo que {_MAX_SKU_LEN}")):
        if resumen[clave]:
            print(f"     {resumen[clave]} {texto}")


def _guardar_backup(detalle_update: list[dict]) -> str:
    """Backup JSON de lo que se va a pisar. Obligatorio: la tabla no tiene
    historial (ninguna migracion agrego columnas de auditoria) y updated_at se
    reemplaza con now(), asi que sin esto el texto viejo no queda en ningun lado.
    Mismo lugar y forma que el resto de los backups de datos del repo."""
    carpeta = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backups"))
    os.makedirs(carpeta, exist_ok=True)
    sello = datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d_%H%M%S")
    ruta = os.path.join(carpeta, f"sku_descripciones_pre_stock_{sello}.json")
    with open(ruta, "w", encoding="utf-8") as fh:
        json.dump(detalle_update, fh, ensure_ascii=False, indent=1, default=str)
    return ruta


async def run_seed(excel_path: str, dry_run: bool, pisar: bool) -> None:
    pares, resumen = parse_export(excel_path)
    _imprimir_resumen_archivo(pares, resumen, excel_path)

    raw_url = os.environ.get("DATABASE_URL", "")
    if not raw_url:
        if dry_run:
            print("\n[DRY RUN] Sin DATABASE_URL: valide solo el archivo, no compare contra la base.")
            return
        print("\nERROR: DATABASE_URL no definida")
        sys.exit(1)
    if not pares:
        print("No hay nada para cargar.")
        return

    dsn = re.sub(r"^postgresql\+asyncpg://", "postgresql://", raw_url)
    dsn = re.sub(r"^postgres://", "postgresql://", dsn)

    try:
        # statement_cache_size=0 es obligatorio contra el pooler de Supabase
        # (puerto 6543, pgBouncer en transaction mode): cada transaccion puede
        # caer en un backend distinto y el cache de prepared statements de
        # asyncpg reusa nombres que ese backend no conoce. Mismo motivo que el
        # comentario de app/core/database.py.
        conn = await asyncpg.connect(dsn, statement_cache_size=0)
    except Exception as e:
        print(f"\nERROR: no pude conectar a la base ({type(e).__name__}: {e})")
        sys.exit(1)

    try:
        # Un solo round-trip: un SELECT por fila serian cientos de idas y vueltas
        # secuenciales contra Supabase para nada.
        skus = [sku for sku, _ in pares]
        existing_rows = await conn.fetch(
            "SELECT sku, id, updated_by_id, descripcion, created_at, updated_at "
            "FROM sku_descripciones WHERE sku = ANY($1::text[])",
            skus,
        )
        existing = {r["sku"]: r for r in existing_rows}

        to_insert: list[tuple[str, str]] = []
        to_update: list[tuple[str, int]] = []
        detalle_update: list[dict] = []
        ya_iguales = 0
        con_correccion_humana: list[str] = []
        ya_tenian: list[str] = []

        for sku, descripcion in pares:
            row = existing.get(sku)
            if row is None:
                to_insert.append((sku, descripcion))
            elif row["descripcion"] == descripcion:
                ya_iguales += 1
            elif row["updated_by_id"] is not None:
                con_correccion_humana.append(sku)
            elif not pisar:
                ya_tenian.append(sku)
            else:
                to_update.append((descripcion, row["id"]))
                detalle_update.append({
                    "sku": sku, "id": row["id"],
                    "descripcion_anterior": row["descripcion"],
                    "descripcion_nueva": descripcion,
                    "updated_by_id": row["updated_by_id"],
                    "created_at": row["created_at"], "updated_at": row["updated_at"],
                })

        print(f"\nContra la base:")
        print(f"   {len(to_insert)} SKU nuevos (se agregan)")
        print(f"   {ya_iguales} ya estaban con el mismo texto")
        print(f"   {len(ya_tenian)} ya tienen otra descripcion -> SE SALTEAN"
              f"{' (usa --pisar-existentes para reemplazarlas)' if ya_tenian else ''}")
        print(f"   {len(con_correccion_humana)} corregidos a mano por alguien -> nunca se tocan")
        if pisar:
            print(f"   {len(to_update)} SE VAN A PISAR (--pisar-existentes)")

        # La lista completa, no una muestra: es el unico lugar donde se ve el texto
        # concreto que entra a produccion, y con un preview de 5 una degradacion
        # masiva pasa desapercibida detras de un resumen que suena a exito.
        for sku, desc in to_insert:
            print(f"   NUEVO  {sku}  {desc}")
        for d in detalle_update:
            print(f"   PISA   {d['sku']}  {d['descripcion_anterior']!r} -> {d['descripcion_nueva']!r}")
        if ya_tenian:
            print(f"   SALTEA (ya tenian): {', '.join(ya_tenian)}")
        if con_correccion_humana:
            print(f"   SALTEA (correccion humana): {', '.join(con_correccion_humana)}")

        if dry_run:
            print("\n[DRY RUN] No escribi nada.")
            return

        if detalle_update:
            print(f"\nBackup de lo que se pisa -> {_guardar_backup(detalle_update)}")

        async with conn.transaction():
            if to_insert:
                # ON CONFLICT DO NOTHING: entre el SELECT y esto alguien puede crear
                # el SKU desde el Convertidor (que usa el mismo upsert). Sin esto el
                # executemany es atomico y se cae el lote ENTERO por una sola fila.
                await conn.executemany(
                    "INSERT INTO sku_descripciones (sku, descripcion) VALUES ($1, $2) "
                    "ON CONFLICT (sku) DO NOTHING",
                    to_insert,
                )
            if to_update:
                # El AND revalida el candado del lado del servidor: entre el SELECT y
                # el UPDATE alguien pudo corregir ese SKU a mano, y la regla "nunca
                # pisar una correccion humana" tiene que valer tambien en esa ventana.
                await conn.executemany(
                    "UPDATE sku_descripciones SET descripcion = $1, updated_at = now() "
                    "WHERE id = $2 AND updated_by_id IS NULL",
                    to_update,
                )
    finally:
        await conn.close()

    print(f"\nListo -- {len(to_insert)} agregados, {len(to_update)} pisados, "
          f"{len(ya_tenian) + len(con_correccion_humana)} salteados")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel_path", help="Ruta al export de stock por sucursal (.xlsx)")
    parser.add_argument("--dry-run", action="store_true", help="Solo simula, no escribe nada en la base")
    parser.add_argument("--yes", action="store_true", help="Saltea la confirmacion interactiva")
    parser.add_argument("--pisar-existentes", action="store_true",
                        help="Ademas de agregar, reemplaza descripciones que ya estan "
                             "(salvo las corregidas a mano). Deja backup JSON obligatorio.")
    args = parser.parse_args()

    if not os.path.exists(args.excel_path):
        print(f"ERROR: no encontre el Excel en {args.excel_path}")
        sys.exit(1)

    if not args.dry_run and not args.yes:
        print("Esto escribe en sku_descripciones de la base de PRODUCCION.")
        if args.pisar_existentes:
            print("Y con --pisar-existentes ADEMAS reemplaza descripciones que ya estan.")
        try:
            confirm = input("Escribi CONFIRMAR para continuar: ")
        except EOFError:
            print("Cancelado: no hay terminal interactiva (usa --yes si estas seguro).")
            sys.exit(1)
        if confirm.strip() != "CONFIRMAR":
            print("Cancelado.")
            sys.exit(0)

    print(f"Importando desde: {args.excel_path}")
    asyncio.run(run_seed(args.excel_path, args.dry_run, args.pisar_existentes))


if __name__ == "__main__":
    main()
