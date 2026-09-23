#!/usr/bin/env python
"""
Regenera frontend/lib/calendario/seed.ts leyendo los dos Excel del calendario.

    python backend/scripts/importar_calendario_excel.py

Los dos archivos se esperan en Descargas. Si los movés, cambiá RUTAS.

Esto es REFERENCIA, no parte del producto: los datos de seed.ts son un muestreo
(setiembre y octubre 2026) para que el calendario arranque con algo real. Cuando
el calendario guarde contra el backend los datos salen de ahí, y este script
queda como la documentación de cómo se parsean esos Excel.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

DESCARGAS = Path.home() / "Downloads"
RUTAS = {
    "retail": DESCARGAS / "Calendario espacios digitales y datos de campaña.xlsx",
    "comercial": DESCARGAS / "CALENDARIO 2025 Y 2026 TI.xlsx",
}

# hoja del Excel -> clave del mes en la app
HOJAS_RETAIL = {"SETIEMBRE 26": "2026-09", "OCTUBRE 26": "2026-10"}
HOJAS_COMERCIAL = {"Septiembre 2026": "2026-09", "Octubre 2026": "2026-10"}

DESTINO = (Path(__file__).resolve().parents[2]
           / "frontend" / "lib" / "calendario" / "seed.ts")


def limpiar(v) -> str:
    return re.sub(r"\s+", " ", str(v).strip())


def reparar(s: str) -> str:
    """Algunas hojas vienen con acentos rotos por doble codificación."""
    try:
        r = s.encode("latin1").decode("utf-8")
        if sum(c in "áéíóúñÁÉÍÓÚÑüÜ" for c in r) > sum(c in "áéíóúñ" for c in s):
            return r
    except (UnicodeDecodeError, UnicodeEncodeError):
        pass
    return s


def color(celda) -> str | None:
    try:
        relleno = celda.fill
        if relleno is None or relleno.patternType is None:
            return None
        rgb = relleno.start_color.rgb
        if isinstance(rgb, str) and len(rgb) == 8 and rgb[2:] not in ("000000", "FFFFFF"):
            return "#" + rgb[2:]
    except Exception:
        pass
    return None


def fila_de_dias(ws) -> int:
    """La fila con los números 1..31. Las hojas no la tienen siempre igual."""
    for r in range(1, 12):
        n = sum(
            1 for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(r, c).value, (int, float)) and 1 <= ws.cell(r, c).value <= 31
        )
        if n >= 25:
            return r
    raise ValueError(f"no encontré la fila de días en {ws.title!r}")


def ejes(ws, fila_dias: int):
    col2dia = {
        c: int(ws.cell(fila_dias, c).value)
        for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(fila_dias, c).value, (int, float))
        and 1 <= ws.cell(fila_dias, c).value <= 31
    }
    dow = {col2dia[c]: limpiar(ws.cell(fila_dias - 1, c).value or "") for c in col2dia}
    return col2dia, min(col2dia), max(col2dia), dow


def leer_bandas(ws, fila_dias, col_etiqueta, fila_corte, col_ini, col_fin, col2dia):
    """
    Una banda es una fila del Excel (un tipo de acción o un formato). Ocupa
    tantas filas como carriles tenga. Las campañas son celdas combinadas que
    abarcan un rango de días.
    """
    inicios = {(m.min_row, m.min_col): (m.max_row, m.max_col) for m in ws.merged_cells.ranges}
    tapadas = {
        (r, c)
        for m in ws.merged_cells.ranges
        for r in range(m.min_row, m.max_row + 1)
        for c in range(m.min_col, m.max_col + 1)
        if (r, c) != (m.min_row, m.min_col)
    }

    bandas, actual = [], None
    for r in range(fila_dias + 1, fila_corte):
        etiqueta = ws.cell(r, col_etiqueta).value
        if etiqueta and limpiar(etiqueta):
            actual = {"nombre": reparar(limpiar(etiqueta)), "filas": []}
            bandas.append(actual)
        if actual is None:
            continue

        carril = []
        for c in range(col_ini, col_fin + 1):
            v = ws.cell(r, c).value
            if v in (None, "") or (r, c) in tapadas:
                continue
            _, col_max = inicios.get((r, c), (r, c))
            desde = col2dia.get(c)
            if desde is None:
                continue
            carril.append({
                "nombre": reparar(limpiar(v)),
                "desde": desde,
                "hasta": col2dia.get(min(col_max, col_fin), desde),
                "color": color(ws.cell(r, c)),
            })
        actual["filas"].append(carril)

    for b in bandas:
        while len(b["filas"]) > 1 and not b["filas"][-1]:
            b["filas"].pop()
    return bandas


def extraer_retail():
    wb = openpyxl.load_workbook(RUTAS["retail"], data_only=True)
    out = {}
    for hoja, clave in HOJAS_RETAIL.items():
        ws = wb[hoja]
        fd = fila_de_dias(ws)
        col2dia, ci, cf, dow = ejes(ws, fd)
        # estas hojas traen el calendario comercial pegado abajo: cortar ahí
        corte = ws.max_row + 1
        for r in range(fd + 1, min(ws.max_row + 1, 200)):
            b = ws.cell(r, 2).value
            if b and re.match(r"^\s*20\d\d", str(b)):
                corte = r
                break
        bandas = [b for b in leer_bandas(ws, fd, 3, corte, ci, cf, col2dia) if any(b["filas"])][:20]

        # el Excel repite "Banner check out" en eComm y en Express
        vistos = set()
        for b in bandas:
            if b["nombre"] in vistos:
                b["nombre"] += " (Express)"
                b["grupo"] = "Express"
            else:
                vistos.add(b["nombre"])
                b["grupo"] = "eComm"
        out[clave] = {"dias": sorted(col2dia.values()), "dow": dow, "bandas": bandas}
    return out


def extraer_comercial():
    wb = openpyxl.load_workbook(RUTAS["comercial"], data_only=True)
    out = {}
    for hoja, clave in HOJAS_COMERCIAL.items():
        ws = wb[hoja]
        fd = fila_de_dias(ws)
        col2dia, ci, cf, dow = ejes(ws, fd)
        # cada hoja trae el año actual y abajo el anterior: quedarse con el primero
        corte = ws.max_row + 1
        for r in range(fd + 2, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a and "DIA DE LA SEMANA" in str(a).upper():
                corte = r - 1
                break
        out[clave] = {
            "dias": sorted(col2dia.values()),
            "dow": dow,
            "bandas": leer_bandas(ws, fd, 1, corte, ci, cf, col2dia),
        }
    return out


def firma(retail, comercial) -> str:
    """Huella de los datos extraidos. Si cambia, los meses que nadie edito
    se rehacen solos la proxima vez que se abre el calendario (ver el
    `merge` de frontend/lib/calendario/store.ts)."""
    crudo = json.dumps({"retail": retail, "comercial": comercial},
                       sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(crudo.encode("utf-8")).hexdigest()[:12]


def a_typescript(retail, comercial) -> str:
    buf = io.StringIO()
    w = buf.write
    w("// Generado por backend/scripts/importar_calendario_excel.py — no editar a mano.\n")
    w("//   retail    <- 'Calendario espacios digitales y datos de campaña.xlsx'\n")
    w("//   comercial <- 'CALENDARIO 2025 Y 2026 TI.xlsx'\n\n")
    w("import type { SeedMes } from './tipos'\n\n")
    w("/** Cambia cada vez que se reimportan los Excel. */\n")
    w("export const SEED_VERSION = %r\n\n" % firma(retail, comercial))

    def carriles(filas, sangria):
        s = "[\n"
        for fila in filas:
            s += sangria + "  [" + ", ".join(
                "{{ nombre: {}, desde: {}, hasta: {}, color: {} }}".format(
                    json.dumps(b["nombre"], ensure_ascii=False), b["desde"], b["hasta"],
                    json.dumps(b["color"]) if b["color"] else "null",
                ) for b in fila
            ) + "],\n"
        return s + sangria + "]"

    w("export const SEED: Record<string, SeedMes> = {\n")
    for clave in sorted(set(retail) | set(comercial)):
        r = retail.get(clave, {"dias": [], "dow": {}, "bandas": []})
        c = comercial.get(clave, {"dias": [], "dow": {}, "bandas": []})
        dias = c["dias"] or r["dias"]
        dow = c["dow"] or r["dow"]
        w(f"  '{clave}': {{\n")
        w(f"    dias: {len(dias)},\n")
        w("    dow: %s,\n" % json.dumps({str(k): v for k, v in sorted(dow.items())}, ensure_ascii=False))
        for campo, bandas, con_grupo in (("comercial", c["bandas"], False), ("retail", r["bandas"], True)):
            w(f"    {campo}: [\n")
            for b in bandas:
                grupo = ", grupo: %s" % json.dumps(b.get("grupo", "eComm")) if con_grupo else ""
                w("      { nombre: %s%s, filas: %s },\n" % (
                    json.dumps(b["nombre"], ensure_ascii=False), grupo, carriles(b["filas"], "      ")))
            w("    ],\n")
        w("  },\n")
    w("}\n")
    return buf.getvalue()


def main() -> int:
    faltan = [str(p) for p in RUTAS.values() if not p.exists()]
    if faltan:
        print("No encuentro estos archivos:", *faltan, sep="\n  ")
        return 1

    retail, comercial = extraer_retail(), extraer_comercial()
    DESTINO.write_text(a_typescript(retail, comercial), encoding="utf-8")

    print(f"Escrito {DESTINO}")
    for clave in sorted(set(retail) | set(comercial)):
        nr = sum(len(f) for b in retail.get(clave, {}).get("bandas", []) for f in b["filas"])
        nc = sum(len(f) for b in comercial.get(clave, {}).get("bandas", []) for f in b["filas"])
        print(f"  {clave}: {nc} acciones comerciales, {nr} campañas de Retail Media")
    return 0


if __name__ == "__main__":
    sys.exit(main())
