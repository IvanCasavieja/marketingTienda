"""Validación de placas contra una PLANILLA, de punta a punta.

El camino que pidió Ivan el 18/09/2026 para las campañas que no tienen mailing
físico: se sube un .xlsx o un .csv, se lee como DATOS (sin IA) y las placas se
comparan contra eso.

El archivo de prueba es `tests/data/planilla_rrss.xlsx` y está armado como sale
un export de gestión de verdad: una fila de título, una en blanco, recién ahí
los encabezados, y una fila de pie que no es un producto.

Lo que estos tests fijan, que son las decisiones del diseño:
  - la descripción es la que empareja, y tolera estar mal escrita;
  - los precios se comparan por IMPORTE, porque la planilla trae un número y no
    un texto impreso;
  - lo que la planilla NO dice no se valida ni se inventa;
  - los tres casos de emparejamiento (ninguna fila, dos filas, filas que nadie
    reclamó) son visibles, no se resuelven en silencio.
"""
import io
import pathlib
import re
import time

import pytest
from openpyxl import Workbook, load_workbook

from app.services.rrss import archivos, comparador as c, excel, planilla

DATOS = pathlib.Path(__file__).parent / "data" / "planilla_rrss.xlsx"


@pytest.fixture(scope="module")
def mailing() -> dict:
    return planilla.leer(DATOS.read_bytes(), "planilla_rrss.xlsx").mailing


def placa_de(**kw) -> dict:
    """Una lectura de placa como la que devuelve CatTi, con lo mínimo."""
    producto = dict(
        descripcion="", precio_anterior="", precio_anterior_tachado=True, mecanica="",
        oferta_encabezado="Oferta", oferta_precio="", oferta_pie="", es_alcohol=False, cajas={},
    )
    producto.update(kw)
    return {
        "producto": producto, "fecha": "", "logo_campana_presente": True, "isotipo_presente": True,
        "legal_bases": "Bases y condiciones en tiendainglesa.com.uy", "legal_alcohol": "", "cta": "",
        "otros_textos": [], "cajas": {},
        "imagen_producto": dict(presente=True, que_se_ve="", coincide_con_descripcion=True, motivo=""),
    }


CONFIG = {"legal_bases": "Bases y condiciones en tiendainglesa.com.uy", "legal_alcohol": "", "fecha": ""}


def por_campo(filas):
    return {f["campo"]: f for f in filas}


# ---------------------------------------------------------------- leer

def test_encuentra_los_encabezados_aunque_no_esten_en_la_primera_fila(mailing):
    """El export real trae una fila de título y una en blanco antes. Si se
    asumiera la fila 1, el archivo entero se rechazaría."""
    assert mailing["planilla"]["fila_encabezado"] == 3
    assert mailing["planilla"]["hoja"] == "Alemania 2026"
    assert len(mailing["productos"]) == 7


def test_el_numero_de_fila_es_el_real_del_archivo(mailing):
    """El que muestra Excel al abrirlo. Si el diseñador va a la fila 4 y ahí hay
    otra cosa, se perdió toda la credibilidad de una."""
    assert mailing["productos"][0]["fila"]["numero"] == 4
    assert mailing["productos"][-1]["fila"]["numero"] == 10


def test_la_fila_de_pie_no_es_un_producto(mailing):
    """En ESTE archivo "Total: 7 artículos" está en la columna CODIGO, así que
    se cae sola por no tener descripción. Cuando gestión lo escribe en la
    columna DESCRIPCION no se caía nada: ver
    `test_el_pie_del_reporte_no_es_un_producto`, más abajo."""
    assert all(p["descripcion"] for p in mailing["productos"])
    assert mailing["planilla"]["filas_ignoradas"] == 1


def test_la_descripcion_que_manda_es_DESCRIPCION_y_no_NOMBREARTICULO(mailing):
    """NOMBREARTICULO es el nombre corto de gestión ("CERVEZA PATRICIA LATA
    473ML"); lo que va impreso en la placa es DESCRIPCION."""
    assert mailing["productos"][0]["descripcion"] == "Cerveza PATRICIA lata. 473 ml"


def test_el_precio_canonico_lleva_la_moneda_de_la_planilla(mailing):
    assert mailing["productos"][0]["oferta_precio"] == "$74,50"
    assert mailing["productos"][2]["oferta_precio"] == "U$S39"
    assert mailing["productos"][3]["precio_anterior"] == "$18.990"


def test_no_gasta_ia(mailing):
    """Es media razón de ser de este camino: el otro son dos llamadas al modelo
    por página del mailing."""
    assert mailing["origen"] == "planilla"
    assert mailing["fecha_pagina"] is None
    assert all(p["pagina"] is None and p["cajas"] == {} for p in mailing["productos"])


def test_un_csv_entra_por_el_mismo_lector():
    crudo = (
        "CODIGO;DESCRIPCION;MONEDA;PRECIOANT;PRECIO\n"
        "501233;Cerveza PATRICIA lata. 473 ml;$;99;74,5\n"
        "501235;Mostaza Dijon MAILLE. 215 g;$;245;199\n"
    ).encode("utf-8")
    m = planilla.leer(crudo, "listado.csv").mailing
    assert [p["descripcion"] for p in m["productos"]] == [
        "Cerveza PATRICIA lata. 473 ml", "Mostaza Dijon MAILLE. 215 g",
    ]
    assert m["productos"][0]["oferta_precio"] == "$74,50"


def test_nombre_articulo_es_la_descripcion_cuando_no_hay_otra():
    """El listado real de los rompeprecios (22/09/2026) trae la descripción
    impresa en NOMBRE ARTÍCULO y ninguna otra columna de texto. Hasta ese día
    NOMBREARTICULO no contaba nunca y el archivo se rechazaba entero."""
    crudo = (
        "CODIGO;NOMBRE ARTÍCULO;MONEDA;PRECIOANT;PRECIO\n"
        "478352;Arvejas TIENDA INGLESA. 300 g;$;48;37,5\n"
    ).encode("utf-8")
    m = planilla.leer(crudo, "listado.csv").mailing
    assert m["productos"][0]["descripcion"] == "Arvejas TIENDA INGLESA. 300 g"


def test_sin_ninguna_columna_de_descripcion_no_se_valida_nada():
    crudo = b"CODIGO;PRECIO\n501233;74,5\n"
    with pytest.raises(planilla.PlanillaInvalida) as exc:
        planilla.leer(crudo, "listado.csv")
    assert "descripción" in str(exc.value)


# ---------------------------------------------------------------- CatTi decide las columnas

# El listado REAL que subió Ivan el 22/09/2026, con su forma: la descripción en
# NOMBRE ARTÍCULO y la mecánica en una columna que se llama OFERTA.
def _listado_real() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "mbdelcerro_Mailing_13316"
    ws.append(["CODIGO", "NOMBRE ARTÍCULO", "MONEDA", "PRECIOANT", "PRECIO", "OFERTA"])
    ws.append([31026, "Agua mineral natural sin gas NATIVA. 1 L", "$", 80, 64, None])
    ws.append([478352, "Arvejas TIENDA INGLESA. 300 g", "$", 48, 37.5, "2x$75"])
    ws.append([14022, "Lomito canadiense TIENDA INGLESA. 100g", "$", 990, 830, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Lo que contestó CatTi de verdad sobre ese listado (corrida con la API del
# 22/09/2026), para no depender del modelo en los tests.
MAPEO_REAL = {
    "hoja": "mbdelcerro_Mailing_13316",
    "fila_encabezados": 1,
    "columnas": {
        "descripcion": "B", "precio_anterior": "D", "precio_oferta": "E", "moneda": "C",
        "mecanica": "F", "arriba_del_precio": "", "abajo_del_precio": "", "vigencia": "",
        "leyenda_alcohol": "", "sucursal": "", "stock": "", "codigo": "A",
    },
    "explicacion": "El nombre del producto está en B y la mecánica tipo 2x$75 en F.",
    "dudas": "",
}


def test_con_lo_que_decidio_catti_se_lee_el_listado_real():
    m = planilla.leer(_listado_real(), "LISTADO FINAL.xlsx", mapeo=MAPEO_REAL).mailing
    arvejas = m["productos"][1]
    assert arvejas["descripcion"] == "Arvejas TIENDA INGLESA. 300 g"
    assert arvejas["mecanica"] == "2x$75"          # salió de la columna OFERTA
    assert arvejas["precio_anterior"] == "$48"
    assert arvejas["oferta_precio"] == "$37,50"
    assert m["campos"] == ["descripcion", "precio_anterior", "oferta_precio", "mecanica"]
    assert arvejas["fila"]["numero"] == 3


def test_la_interpretacion_de_catti_queda_a_la_vista():
    """La persona tiene que poder ver CÓMO se leyó su archivo antes de creerle
    a una corrección: qué columna tomó para cada dato, con el título real."""
    it = planilla.leer(_listado_real(), "x.xlsx", mapeo=MAPEO_REAL).mailing["planilla"]["interpretacion"]
    assert it["explicacion"].startswith("El nombre del producto")
    assert {"dato": "Descripción", "letra": "B", "titulo": "NOMBRE ARTÍCULO"} in it["columnas"]
    assert {"dato": "Mecánica", "letra": "F", "titulo": "OFERTA"} in it["columnas"]


def test_los_valores_salen_de_las_celdas_no_de_catti():
    """CatTi decide COORDENADAS; el contenido se copia de la celda tal cual. Una
    planilla es dato exacto y tiene que seguir siéndolo."""
    m = planilla.leer(_listado_real(), "x.xlsx", mapeo=MAPEO_REAL).mailing
    assert [p["descripcion"] for p in m["productos"]] == [
        "Agua mineral natural sin gas NATIVA. 1 L", "Arvejas TIENDA INGLESA. 300 g",
        "Lomito canadiense TIENDA INGLESA. 100g",
    ]


def test_sin_fila_de_titulos_tambien_se_lee():
    """Una planilla sin títulos: CatTi contesta fila_encabezados=0 y los datos
    arrancan en la fila 1. Los títulos que se muestran son "columna B"."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Arvejas TIENDA INGLESA. 300 g", 48, 37.5])
    ws.append(["Bola de lomo. Kg", 499, 399])
    buf = io.BytesIO()
    wb.save(buf)
    mapeo = {**MAPEO_REAL, "hoja": ws.title, "fila_encabezados": 0,
             "columnas": {**{k: "" for k in MAPEO_REAL["columnas"]},
                          "descripcion": "A", "precio_anterior": "B", "precio_oferta": "C"}}
    m = planilla.leer(buf.getvalue(), "sin_titulos.xlsx", mapeo=mapeo).mailing
    assert [p["descripcion"] for p in m["productos"]] == ["Arvejas TIENDA INGLESA. 300 g", "Bola de lomo. Kg"]
    assert m["productos"][0]["fila"]["numero"] == 1
    assert m["planilla"]["titulos"]["descripcion"] == "columna A"


def test_si_catti_no_encuentra_la_descripcion_se_dice():
    mapeo = {**MAPEO_REAL, "columnas": {**MAPEO_REAL["columnas"], "descripcion": ""}, "dudas": "No hay nombres."}
    with pytest.raises(planilla.PlanillaInvalida) as exc:
        planilla.leer(_listado_real(), "x.xlsx", mapeo=mapeo)
    assert "descripción" in str(exc.value) and "No hay nombres." in str(exc.value)


def test_una_columna_que_no_existe_o_repetida_no_se_cree():
    mapeo = {**MAPEO_REAL, "columnas": {**MAPEO_REAL["columnas"], "vigencia": "Z", "codigo": "B"}}
    m = planilla.leer(_listado_real(), "x.xlsx", mapeo=mapeo).mailing
    avisos = " ".join(m["planilla"]["avisos"])
    assert "columna Z, que no existe" in avisos
    assert "misma columna (B)" in avisos
    assert m["productos"][1]["descripcion"] == "Arvejas TIENDA INGLESA. 300 g"


def test_letras_de_columna_ida_y_vuelta():
    for i in (0, 5, 25, 26, 27, 51, 52, 701, 702):
        assert planilla.indice_de_letra(planilla.letra(i)) == i
    assert planilla.letra(0) == "A" and planilla.letra(26) == "AA"
    assert planilla.indice_de_letra("") is None and planilla.indice_de_letra("3") is None


def test_lo_que_ve_catti_tiene_fila_y_letra():
    texto = planilla.muestra(_listado_real(), "x.xlsx")["texto"]
    assert "Hoja «mbdelcerro_Mailing_13316»" in texto
    assert "fila | A | B | C | D | E | F" in texto
    assert "   3 | 478352 | Arvejas TIENDA INGLESA. 300 g | $ | 48 | 37,5 | 2x$75" in texto


def test_preparar_planilla_usa_lo_que_decide_catti(monkeypatch):
    import asyncio
    from app.services.rrss import catti, validador

    async def interpretar(texto):
        assert "NOMBRE ARTÍCULO" in texto
        return MAPEO_REAL, 5600, 400

    monkeypatch.setattr(catti, "interpretar_planilla", interpretar)
    prep = asyncio.run(validador.preparar_planilla(_listado_real(), "LISTADO FINAL.xlsx"))
    assert prep.mailing["productos"][1]["mecanica"] == "2x$75"
    assert (prep.tokens_in, prep.tokens_out) == (5600, 400)


def test_si_catti_no_esta_se_lee_por_nombres_y_se_avisa(monkeypatch):
    import asyncio
    from app.services.rrss import catti, validador

    async def caido(texto):
        raise RuntimeError("ANTHROPIC_API_KEY no configurado")

    monkeypatch.setattr(catti, "interpretar_planilla", caido)
    prep = asyncio.run(validador.preparar_planilla(_listado_real(), "LISTADO FINAL.xlsx"))
    assert prep.mailing["productos"][1]["descripcion"] == "Arvejas TIENDA INGLESA. 300 g"
    assert "CatTi no pudo mirar la planilla" in prep.mailing["planilla"]["avisos"][0]
    assert prep.mailing["planilla"]["interpretacion"] is None


# ---------------------------------------------------------------- qué exige y qué no

def test_solo_exige_las_columnas_que_trae(mailing):
    assert mailing["campos"] == ["descripcion", "precio_anterior", "oferta_precio"]


def test_lo_que_la_planilla_no_dice_no_se_marca_como_error(mailing):
    """Una placa con mecánica "2x$199" contra una planilla sin columna de
    mecánica NO es un error: la planilla no dice nada de eso. Marcarlo sería
    acusar a la placa con un dato que nadie escribió."""
    p = placa_de(descripcion="Cerveza PATRICIA lata. 473 ml", precio_anterior="$99",
                 oferta_precio="$74,50", mecanica="2x$199", oferta_pie="unidad")
    idx, _, filas = c.comparar_placa(p, mailing, CONFIG)
    campos = por_campo(filas)
    assert idx == 0
    assert "mecanica" not in campos and "oferta_pie" not in campos
    assert c.estado_de_la_placa(filas, idx) == "ok"


def test_se_dice_cuales_campos_quedaron_afuera(mailing):
    """No se resuelve en silencio: la pantalla y el Excel muestran esta lista."""
    # La planilla de prueba no trae VIGENCIA ni LEYENDA ALCOHOL y la config no
    # las escribe: la fecha y la leyenda tampoco se revisaron, y se dice.
    assert planilla.campos_que_no_dicta(mailing, CONFIG) == [
        "Mecánica", "Texto arriba del precio", "Texto abajo del precio",
        "Si el precio anterior va tachado", "Fecha de la campaña", "Leyenda de alcohol",
    ]


def test_contra_un_mailing_se_sigue_exigiendo_todo():
    """La planilla no aflojó la regla del mailing: ahí sigue mandando el texto
    impreso, carácter por carácter."""
    reglas = c.reglas_de_la_fuente({"origen": "mailing"})
    assert reglas["precios"] == "texto"
    assert set(reglas["campos"]) == {campo for campo, _ in c.CAMPOS_PRODUCTO}


# ---------------------------------------------------------------- precios

def test_el_precio_se_compara_por_importe_no_por_como_esta_escrito(mailing):
    """La planilla trae 74,5: cómo se escribe en la placa ("$74,5", "$74,50")
    no lo dicta la planilla. Exigir un formato sería inventar una regla."""
    p = placa_de(descripcion="Cerveza PATRICIA lata. 473 ml", precio_anterior="$99", oferta_precio="$74,5")
    _, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert por_campo(filas)["oferta_precio"]["estado"] == "ok"


def test_un_precio_distinto_de_verdad_sigue_siendo_un_error(mailing):
    p = placa_de(descripcion="Cerveza PATRICIA lata. 473 ml", precio_anterior="$99", oferta_precio="$75")
    _, _, filas = c.comparar_placa(p, mailing, CONFIG)
    fila = por_campo(filas)["oferta_precio"]
    assert fila["severidad"] == "error" and fila["mailing"] == "$74,50"


def test_lo_que_acompana_al_precio_no_lo_dicta_la_planilla(mailing):
    """"$99 unidad" contra un 99: el importe es el mismo y la palabra "unidad"
    no está en ninguna celda."""
    p = placa_de(descripcion="Cerveza PATRICIA lata. 473 ml", precio_anterior="$99 unidad", oferta_precio="$74,50")
    _, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert por_campo(filas)["precio_anterior"]["estado"] == "ok"


@pytest.mark.parametrize("texto, esperado", [
    ("$37,50", 37.5), ("U$S149", 149.0), ("$1.090", 1090.0), ("$1.090,50", 1090.5),
    ("1.7 L", 1.7), ("$74,5", 74.5), ("", None), ("Oferta", None),
])
def test_como_se_lee_un_precio(texto, esperado):
    """"U$S149" empieza con una letra: por eso no sirve el parse de precios del
    Convertidor, cuya regex exige empezar con un dígito."""
    assert c.importe(texto) == esperado


# ---------------------------------------------------------------- emparejar

def test_la_descripcion_empareja_aunque_este_mal_escrita(mailing):
    """Es la respuesta de Ivan: "la descripción que extraés de las fotos es la
    misma descripción que debe tener el Excel". Si estuviera bien escrita no
    haría falta validar nada."""
    p = placa_de(descripcion="Cerveza PATRICIA lata. 473ml.", oferta_precio="$74,50")
    idx, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert idx == 0
    assert por_campo(filas)["descripcion"]["severidad"] == "error"


def test_dos_filas_que_se_parecen_igual_no_se_adivinan(mailing):
    """WARSTEINER botella y WARSTEINER lata. Un emparejamiento equivocado hace
    que se acuse de error a una placa que está bien, y eso es peor que no
    validar."""
    p = placa_de(descripcion="Cerveza WARSTEINER. 500 ml", oferta_precio="$99,50")
    idx, _ = c.emparejar(p["producto"], mailing["productos"])
    assert idx is None
    cands = c.candidatos(p["producto"], mailing["productos"])
    assert c.motivo_sin_pareja(cands) == "ambiguo"
    assert {k["descripcion"] for k in cands[:2]} == {
        "Cerveza WARSTEINER botella. 500 ml", "Cerveza WARSTEINER lata. 500 ml",
    }


def test_cuando_no_se_parece_ninguna_se_dice_y_se_muestra_contra_que(mailing):
    p = placa_de(descripcion="Shampoo SEDAL rizos definidos. 340 ml", oferta_precio="$210")
    idx, _ = c.emparejar(p["producto"], mailing["productos"])
    assert idx is None
    cands = c.candidatos(p["producto"], mailing["productos"])
    assert c.motivo_sin_pareja(cands) == "ninguna"
    assert len(cands) == 3 and all(k["puntaje"] < c.UMBRAL_EMPAREJAR for k in cands)


def test_las_filas_que_ninguna_placa_reclamo_se_dicen_cuales(mailing):
    """Se venían contando y no se decía nunca cuáles. Con una planilla es un
    hallazgo de primera: falta la placa de ese producto."""
    imgs = [{
        "id": 1, "nombre_archivo": "a.jpg", "formato": "1:1", "match_indice": 0,
        "lectura": placa_de(descripcion="Cerveza PATRICIA lata. 473 ml"),
    }]
    r = c.chequeos_del_lote(imgs, mailing["productos"])
    sin_placa = r["productos_sin_placa"]
    assert r["productos_con_placa"] == 1
    assert len(sin_placa) == 6
    assert {"indice": 6, "descripcion": "Pan negro de centeno. 500 g", "fila": 10} in sin_placa


def test_el_umbral_deja_afuera_al_peor_falso_positivo_medido():
    """Sobre las 67 placas reales de producción, el mejor puntaje contra un
    producto EQUIVOCADO fue 81,0 ("Suprema de pollo TIENDA INGLESA. Kg" contra
    "Chorizo TIENDA INGLESA. Kg") y el peor acierto 103,5. El umbral era 72, o
    sea por debajo del falso."""
    suprema = dict(descripcion="Suprema de pollo TIENDA INGLESA. Kg", precio_anterior="", oferta_precio="")
    chorizo = dict(descripcion="Chorizo TIENDA INGLESA. Kg", precio_anterior="", oferta_precio="")
    assert c.puntaje_producto(suprema, chorizo) < c.UMBRAL_EMPAREJAR


# ---------------------------------------------------------------- la evidencia

def test_cada_fila_puede_dar_su_pedazo_de_planilla_dibujado(mailing):
    """Es lo que ocupa el lugar del recorte del mailing: en la columna B del
    Excel y en el panel de la derecha de la pantalla, sin cambiar nada de los
    dos. Se dibuja CUANDO SE PIDE (ver el bloque del tope de trabajo)."""
    for i in range(len(mailing["productos"])):
        assert planilla.evidencia(mailing, i).startswith("data:image/jpeg;base64,")


def test_la_tira_muestra_la_fila_con_sus_vecinas(mailing):
    """Los vecinos son los que hacen que el diseñador CREA el emparejamiento."""
    con_vecinas = planilla.tira(mailing, 3)
    sola = planilla.tira(mailing, 3, "oferta_precio", vecinas=False)
    assert con_vecinas.height > sola.height
    assert con_vecinas.width == sola.width


def test_la_cita_permite_ir_a_mirarlo_a_mano(mailing):
    assert planilla.cita(mailing, 0) == "planilla_rrss.xlsx · hoja «Alemania 2026» · fila 4"


def test_el_recorte_de_un_campo_marca_ese_campo(mailing):
    assert planilla.recorte_de_campo(mailing, 0, "descripcion").startswith("data:image/jpeg;base64,")
    assert planilla.recorte_de_campo(mailing, 99, "descripcion") is None


# ---------------------------------------------------------------- el Excel

def _validacion(mailing: dict) -> dict:
    """Una validación con una placa mal y otra que no está en la planilla."""
    mala = placa_de(descripcion="Cerveza PATRICIA lata. 473ml.", precio_anterior="$99", oferta_precio="$74,50")
    idx, _, filas = c.comparar_placa(mala, mailing, CONFIG)
    suelta = placa_de(descripcion="Shampoo SEDAL rizos definidos. 340 ml", oferta_precio="$210")
    cands = c.candidatos(suelta["producto"], mailing["productos"])
    return {
        "mailing": mailing,
        "imagenes": [
            # `recorte_fuente` es la tira de la fila, dibujada al emparejar: es lo
            # que deja validador.validar_placa en el resultado de la placa.
            {"nombre_archivo": "campana_01.jpg", "formato": "1:1", "estado": "diferencias", "orden": 0,
             "match": {"indice": idx, "puntaje": 110.0}, "candidatos": [], "sin_pareja": None,
             "recorte_fuente": planilla.evidencia(mailing, idx),
             "filas": filas, "lectura": mala, "error": None, "vista": None},
            {"nombre_archivo": "campana_02.jpg", "formato": "1:1", "estado": "sin_match", "orden": 1,
             "match": None, "candidatos": cands, "sin_pareja": c.motivo_sin_pareja(cands),
             "filas": [], "lectura": suelta, "error": None, "vista": None},
        ],
    }


def test_el_excel_habla_de_la_planilla_y_no_de_un_mailing_que_no_existe(mailing, tmp_path):
    ruta = tmp_path / "correcciones.xlsx"
    ruta.write_bytes(excel.construir(_validacion(mailing)))
    wb = load_workbook(ruta, rich_text=True)
    ws = wb["Recomendada"]
    assert ws["B3"].value == "LA PLANILLA"
    textos = [str(ws[f"B{r}"].value or "") for r in range(4, 12)]
    assert any("Así está en la planilla" in t and "fila 4" in t for t in textos)
    assert any("No está en la planilla" in t for t in textos)
    # Lo que la planilla no dicta se dice en el encabezado, no se esconde.
    assert "no se revisó: Mecánica" in str(ws["A2"].value)


def _cabecera_de_la_planilla(ws) -> int:
    """En qué fila arranca la tabla. No es fija: arriba van los avisos de la
    lectura, que son justo lo que hay que ver primero."""
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "FILA")


def test_el_excel_trae_la_planilla_entera_con_las_filas_sin_placa_sin_pintar(mailing, tmp_path):
    ruta = tmp_path / "correcciones.xlsx"
    ruta.write_bytes(excel.construir(_validacion(mailing)))
    ws = load_workbook(ruta)["La planilla"]
    cab = _cabecera_de_la_planilla(ws)
    assert ws.cell(row=cab, column=3).value == "DESCRIPCION"
    filas = {ws.cell(row=r, column=1).value: r for r in range(cab + 1, cab + 8)}
    assert set(filas) == {4, 5, 6, 7, 8, 9, 10}
    # la fila 4 la reclamó una placa con diferencias; la 10 no la reclamó nadie
    assert ws.cell(row=filas[4], column=1).fill.fgColor.rgb.endswith("FDECEA")
    assert ws.cell(row=filas[10], column=1).fill.fgColor.rgb.endswith("FFFFFF")
    assert ws.cell(row=filas[10], column=ws.max_column).value == "Sin placa"


def test_el_sin_pareja_del_excel_muestra_contra_que_se_parecio(mailing, tmp_path):
    """Sin esto, "no encontré este producto" es una acusación sin pruebas."""
    ruta = tmp_path / "correcciones.xlsx"
    ruta.write_bytes(excel.construir(_validacion(mailing)))
    ws = load_workbook(ruta, rich_text=True)["Recomendada"]
    todo = " ".join(str(ws.cell(row=r, column=col).value or "") for r in range(4, 12) for col in (3, 4, 5))
    assert "No está en la planilla" in todo
    assert "Lo que más se le pareció" in todo


# ---------------------------------------------------------------- tipos de archivo

@pytest.mark.parametrize("nombre, tipo, origen", [
    ("LISTADO.xlsx", "application/vnd.ms-excel", "planilla"),
    ("LISTADO.XLSX", "", "planilla"),
    ("listado.csv", "", "planilla"),
    ("mailing.pdf", "application/pdf", "mailing"),
    ("mailing.jpg", "image/jpeg", "mailing"),
])
def test_el_archivo_decide_el_camino(nombre, tipo, origen):
    """La EXTENSIÓN manda: Chrome en Windows manda application/vnd.ms-excel para
    un .xlsx y un .csv puede llegar sin content_type."""
    assert archivos.origen_de(nombre, tipo) == origen


def test_los_tipos_de_archivo_salen_de_un_solo_lugar():
    """El backend y la pantalla leen el MISMO archivo: si esto se duplica,
    vuelve el bug de aceptar en un lado lo que el otro rechaza."""
    assert archivos.acepta_placa("placa.webp", "")
    assert archivos.acepta_placa("placa.sin.extension", "image/png")
    assert not archivos.acepta_placa("placa.gif", "image/gif")
    assert ".xlsx" in archivos.TIPOS["fuentes"]["planilla"]["extensiones"]


# ==========================================================================
# Lo que rompía, y no puede volver a romper
# ==========================================================================
#
# Cada bloque de acá abajo es un caso que la verificación adversarial corrió de
# verdad contra el servicio y encontró mal. No son casos inventados: el número
# de fila, el texto y el archivo son los del hallazgo.


def _xlsx(filas: list[tuple], antes: dict[str, list[tuple]] | None = None) -> bytes:
    """Un .xlsx en memoria. `antes` son las hojas que van ANTES de la de datos,
    que siempre se llama "Campaña"."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, contenido in (antes or {}).items():
        ws = wb.create_sheet(nombre)
        for fila in contenido:
            ws.append(list(fila))
    ws = wb.create_sheet("Campaña")
    for fila in filas:
        ws.append(list(fila))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- (1) el salto de línea

_CON_SALTO = (
    "CODIGO;DESCRIPCION;MONEDA;PRECIOANT;PRECIO\r\n"
    "501233;Cerveza PATRICIA lata. 473 ml;$;99;74,5\r\n"
)


def test_un_salto_de_linea_adentro_de_una_celda_no_tira_la_validacion():
    """Alt+Enter en una celda: PIL no puede medir texto multilínea y salía un
    HTTP 500 sin ninguna pista. Y alcanzaba con que el salto estuviera en una
    fila VECINA de la que se estaba dibujando, porque la tira muestra la de
    arriba y la de abajo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Campaña"
    ws.append(["CODIGO", "DESCRIPCION", "MONEDA", "PRECIOANT", "PRECIO"])
    ws.append(["501233", "Cerveza PATRICIA\nlata. 473 ml", "$", 99, 74.5])
    ws.append(["501235", "Mostaza Dijon MAILLE. 215 g", "$", 245, 199])
    ws.append(["501236", "Pan negro de centeno. 500 g", "$", 165, "1\n29"])
    buf = io.BytesIO()
    wb.save(buf)

    m = planilla.leer(buf.getvalue(), "listado.xlsx").mailing
    # el salto se va al leer: es dato sucio, no contenido
    assert m["productos"][0]["descripcion"] == "Cerveza PATRICIA lata. 473 ml"
    assert "\n" not in m["productos"][2]["oferta_precio"]
    # y la evidencia de la fila del MEDIO, que dibuja a sus dos vecinas, sale igual
    for i in range(len(m["productos"])):
        assert (planilla.evidencia(m, i) or "").startswith("data:image/jpeg;base64,")


def test_el_dibujo_de_la_evidencia_no_puede_tumbar_una_validacion(monkeypatch):
    """Segundo cinturón: aunque el dibujo reviente por lo que sea, la validación
    sigue y la diferencia se muestra igual con el texto de los dos lados. Una
    ilustración no puede voltear una request."""
    def explota(*_a, **_k):
        raise ValueError("can't measure length of multiline text")

    monkeypatch.setattr(planilla, "tira", explota)
    m = planilla.leer(_CON_SALTO.encode("utf-8"), "listado.csv").mailing
    assert planilla.evidencia(m, 0) is None
    assert planilla.recorte_de_campo(m, 0, "descripcion") is None
    # que falte no se esconde: lo dicen el Excel (test_el_excel_dice_cuando_no_pudo
    # _dibujar_la_fila) y la pantalla (test_rrss_pantalla).


# ---------------------------------------------------------------- (2) dólares contra pesos

def test_dolares_contra_pesos_no_pasa_como_correcto(mailing):
    """La fila 6 del archivo de prueba está en U$S desde el primer día y no
    había un solo test: la placa decía $149, la planilla U$S149 y el estado
    quedaba en ok, porque lo único que se comparaba era el número."""
    jarras = mailing["productos"][2]
    assert jarras["oferta_precio"] == "U$S39"
    p = placa_de(descripcion=jarras["descripcion"], precio_anterior="$49", oferta_precio="$39")
    idx, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert idx == 2
    fila = por_campo(filas)["oferta_precio"]
    assert fila["severidad"] == "error"
    assert "dólares" in fila["nota"] and "pesos" in fila["nota"]
    assert c.estado_de_la_placa(filas, idx) == "diferencias"


def test_la_misma_moneda_escrita_distinto_sigue_estando_bien(mailing):
    """"U$S", "US$" y "u$s" son la misma moneda: el diseñador la escribe como
    le sale y eso no es un error."""
    jarras = mailing["productos"][2]
    p = placa_de(descripcion=jarras["descripcion"], precio_anterior="US$49", oferta_precio="u$s 39")
    _, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert por_campo(filas)["oferta_precio"]["estado"] == "ok"
    assert por_campo(filas)["precio_anterior"]["estado"] == "ok"


def test_sin_simbolo_en_la_placa_no_se_acusa_de_moneda(mailing):
    """Si la placa no trae ninguna marca de moneda no se puede afirmar en cuál
    está: no se acusa. Lo que no se sabe no se inventa."""
    jarras = mailing["productos"][2]
    p = placa_de(descripcion=jarras["descripcion"], precio_anterior="49", oferta_precio="39")
    _, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert por_campo(filas)["oferta_precio"]["estado"] == "ok"


@pytest.mark.parametrize("texto, esperado", [
    ("$149", "UYU"), ("U$S149", "USD"), ("US$ 149", "USD"), ("u$s149", "USD"),
    ("USD 149", "USD"), ("149", ""), ("Oferta", ""),
])
def test_como_se_lee_la_moneda(texto, esperado):
    assert c.moneda(texto) == esperado


# ---------------------------------------------------------------- (3) la fila de pie

_CON_PIE = (
    "CODIGO;DESCRIPCION;MONEDA;PRECIOANT;PRECIO\n"
    "501233;Cerveza PATRICIA lata. 473 ml;$;99;74,5\n"
    "501235;Mostaza Dijon MAILLE. 215 g;$;245;199\n"
    "501236;Pan negro de centeno. 500 g;$;165;129\n"
    "501237;Mostaza Dijon MAILLE. 215 g;$;245;199\n"
    "501238;Set de jarras cerveceras. 2 unidades;U$S;49;39\n"
    ";TOTAL: 5 artículos;;;\n"
    ";Generado el 18/09/2026;;;\n"
)


def test_el_pie_del_reporte_no_es_un_producto():
    """El archivo de prueba ponía "Total: 7 artículos" en la columna CODIGO, así
    que se caía solo por no tener descripción. Gestión lo escribe en la columna
    DESCRIPCION, y entonces el Excel le decía al diseñador que le faltaba la
    placa de "TOTAL: 5 artículos"."""
    m = planilla.leer(_CON_PIE.encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["filas_leidas"] == 5
    assert all("TOTAL" not in p["descripcion"] for p in m["productos"])
    assert all("Generado" not in p["descripcion"] for p in m["productos"])
    # y se DICE, no se descarta en silencio
    aviso = next(a for a in planilla.avisos(m) if "pie del reporte" in a)
    assert "TOTAL: 5 artículos" in aviso and "fila 7" in aviso


def test_un_producto_que_empieza_como_un_pie_sigue_siendo_un_producto():
    """"Total Blue" es una marca. Para caerse hay que empezar como un pie Y no
    traer ningún otro dato: un producto viene con su código y su precio."""
    crudo = (
        "CODIGO;DESCRIPCION;PRECIO\n"
        "501233;Total Blue. 900 ml;74,5\n"
        "501235;Resumen de temporada, caja. 6 unidades;199\n"
    ).encode("utf-8")
    m = planilla.leer(crudo, "listado.csv").mailing
    assert [p["descripcion"] for p in m["productos"]] == [
        "Total Blue. 900 ml", "Resumen de temporada, caja. 6 unidades",
    ]


# ---------------------------------------------------------------- (4) el vocabulario de gestión

_DE_GESTION = (
    "CODIGO;DESCRIPCION;PRECIO\n"
    "1;ACEITE GIRASOL URUGUAY 900ML;149\n"
    "2;CERVEZA PATRICIA LATA 473ML;74,5\n"
    "3;MOSTAZA DIJON MAILLE 215G;199\n"
    "4;PAN NEGRO CENTENO 500G;129\n"
).encode("utf-8")


def test_una_planilla_en_mayusculas_no_acusa_a_todas_las_placas():
    """El caso que le quema la confianza al motor el primer día: 12 placas
    perfectas acusadas de tener la descripción mal escrita, y el Excel
    ordenándole al diseñador reescribir la campaña entera en MAYÚSCULAS con
    abreviaturas. Se detecta y se avisa; no se acusa."""
    m = planilla.leer(_DE_GESTION, "listado.csv").mailing
    assert m["planilla"]["descripcion_de_gestion"] is True
    # la descripción deja de dictar: pasa a la lista de lo que no se validó
    assert "descripcion" not in m["campos"]
    assert "Descripción" in planilla.campos_que_no_dicta(m)
    aviso = next(a for a in planilla.avisos(m) if "MAYÚSCULAS" in a)
    assert "ACEITE GIRASOL URUGUAY 900ML" in aviso

    p = placa_de(descripcion="Aceite de girasol URUGUAY. 900 ml", oferta_precio="$149")
    idx, _, filas = c.comparar_placa(p, m, CONFIG)
    # sigue emparejando (el parecido ignora mayúsculas) y NO marca la descripción
    assert idx == 0
    assert "descripcion" not in por_campo(filas)
    assert c.estado_de_la_placa(filas, idx) == "ok"


def test_una_planilla_escrita_como_va_impreso_si_dicta_la_descripcion(mailing):
    """El chequeo no puede aflojar la regla en las planillas buenas: si las
    descripciones están escritas como van impresas, la descripción se compara
    como siempre."""
    assert mailing["planilla"]["descripcion_de_gestion"] is False
    assert "descripcion" in mailing["campos"]


# ---------------------------------------------------------------- (5) emparejar con duda

_HERMANOS = (
    "CODIGO;DESCRIPCION;PRECIO\n"
    "1;Cerveza PATRICIA lata. 473 ml;74,5\n"
    "2;Cerveza PATRICIA botella. 1 l;99\n"
).encode("utf-8")


def test_el_caso_que_pregunto_ivan_vuelve_a_dar_el_diagnostico():
    """La placa es la BOTELLA y escribieron "lata": la descripción es JUSTO lo
    que está mal. Con el margen único de 8 esto se quedaba sin pareja (98,37
    contra 93,33: margen 5,04) y se perdía el diagnóstico. Ahora empareja y
    dice qué tiene que decir, con la duda a la vista."""
    m = planilla.leer(_HERMANOS, "listado.csv").mailing
    p = placa_de(descripcion="Cerveza PATRICIA lata. 1 l", oferta_precio="$99")
    par = c.emparejamiento(p["producto"], m["productos"])
    assert par["indice"] == 1
    assert par["duda"] is True
    assert round(par["puntaje"], 2) == 98.37 and round(par["segundo"], 2) == 93.33

    _, _, filas = c.comparar_placa(p, m, CONFIG, par)
    campos = por_campo(filas)
    assert campos["descripcion"]["mailing"] == "Cerveza PATRICIA botella. 1 l"
    # y la duda se ve, en su propia fila, como aviso y no como error
    assert campos["emparejamiento"]["severidad"] == "aviso"
    assert "otra fila que se le parece casi igual" in campos["emparejamiento"]["nota"]


def test_un_empate_de_verdad_sigue_sin_emparejarse(mailing):
    """WARSTEINER botella y WARSTEINER lata sacan 110,00 las dos: ahí no hay una
    mejor y elegir es tirar una moneda. Sin pareja, "ambiguo"."""
    p = placa_de(descripcion="Cerveza WARSTEINER. 500 ml", oferta_precio="$99,50")
    par = c.emparejamiento(p["producto"], mailing["productos"])
    assert par["indice"] is None
    assert par["motivo_sin_pareja"] == "ambiguo"
    assert par["margen"] == 0.0


def test_un_emparejamiento_que_calza_del_todo_no_molesta_con_dudas(mailing):
    """La duda tiene que ser rara: si aparece en todas, deja de leerse."""
    p = placa_de(descripcion="Mostaza Dijon MAILLE. 215 g", precio_anterior="$245", oferta_precio="$199")
    par = c.emparejamiento(p["producto"], mailing["productos"])
    assert par["indice"] == 1 and par["duda"] is False
    _, _, filas = c.comparar_placa(p, mailing, CONFIG, par)
    assert "emparejamiento" not in por_campo(filas)


def test_una_pareja_floja_se_empareja_pero_avisa():
    """El falso de Galletitas: la fila correcta NO está en la planilla y la
    equivocada saca 94,44 con 55 puntos de ventaja. Ningún umbral ni margen la
    para; lo único honesto es emparejar y decir que no está seguro, en vez de
    mandar a reescribir la placa con la confianza de un acierto."""
    crudo = (
        "CODIGO;DESCRIPCION;PRECIO\n"
        "1;Galletitas VAINILLA. 200 g;45\n"
        "2;Pan negro de centeno. 500 g;129\n"
    ).encode("utf-8")
    m = planilla.leer(crudo, "listado.csv").mailing
    p = placa_de(descripcion="Galletitas MARIA. 200 g", oferta_precio="$45")
    par = c.emparejamiento(p["producto"], m["productos"])
    assert par["indice"] == 0 and par["duda"] is True
    assert par["motivo_duda"] == "no_la_calza"
    assert par["parecido"] < c.PARECIDO_SEGURO


def test_lo_que_no_se_parece_a_nada_sigue_sin_pareja(mailing):
    p = placa_de(descripcion="Shampoo SEDAL rizos definidos. 340 ml", oferta_precio="$210")
    par = c.emparejamiento(p["producto"], mailing["productos"])
    assert par["indice"] is None and par["motivo_sin_pareja"] == "ninguna"


# ---------------------------------------------------------------- (6) el catálogo

def test_un_catalogo_se_rechaza():
    """Un catálogo no es una campaña, venga en .xlsx o en .csv.

    El contador se corta en el tope que se le pide y no lee el archivo entero:
    eso es lo que permite rechazar sin materializar nada (1.000.000 de filas
    eran 148,8 s y 264 MB adentro del único hilo de CPU del servicio, con todas
    las validaciones de todo el mundo frenadas). Con qué tope se lo llama ahora
    --y por qué-- está en el bloque (J)."""
    filas = [("CODIGO", "DESCRIPCION", "PRECIO")]
    filas += [(str(n), f"Producto numero {n}. 500 g", 100 + n) for n in range(5000)]
    datos = _xlsx(filas)

    # se cuenta hasta el tope y se corta: no se leen las 5001 filas
    assert planilla._cuantas_filas(datos, "catalogo.xlsx", None, planilla.MAX_FILAS) == planilla.MAX_FILAS + 1
    with pytest.raises(planilla.PlanillaInvalida) as exc:
        planilla.leer(datos, "catalogo.xlsx")
    assert "catálogo entero" in str(exc.value)


def test_un_csv_grande_tambien_se_rechaza():
    crudo = b"CODIGO;DESCRIPCION;PRECIO\n" + b"".join(
        f"{n};Producto numero {n}. 500 g;{100 + n}\n".encode("utf-8") for n in range(2000)
    )
    with pytest.raises(planilla.PlanillaInvalida):
        planilla.leer(crudo, "catalogo.csv")


def test_una_planilla_no_puede_pesar_lo_que_un_mailing():
    """El endpoint aceptaba 30 MB para cualquier fuente, y un catálogo de
    17,2 MB entraba. Un listado de campaña de 1.000 filas pesa menos de 1 MB."""
    assert archivos.max_mb("planilla") < archivos.max_mb("mailing")
    assert archivos.max_bytes("planilla") == archivos.max_mb("planilla") * 1024 * 1024


# ---------------------------------------------------------------- (9) los avisos llegan

def test_los_avisos_de_la_lectura_viajan_adentro_del_mailing():
    """Se armaban en un campo del dataclass y `preparar_planilla` devolvía solo
    `pl.mailing`: no llegaban a la pantalla ni al Excel. Ahora viven adentro del
    dict que se guarda en la base."""
    pl = planilla.leer(_CON_PIE.encode("utf-8"), "listado.csv")
    assert pl.mailing["planilla"]["avisos"] == pl.avisos
    assert pl.avisos and planilla.avisos(pl.mailing) == pl.avisos


def test_los_contadores_de_la_lectura_estan_a_la_vista():
    """filas_leidas, filas_ignoradas y fila_encabezado: el dato que delata que
    el motor contó de más."""
    info = planilla.leer(_CON_PIE.encode("utf-8"), "listado.csv").mailing["planilla"]
    assert info["filas_leidas"] == 5
    assert info["filas_ignoradas"] == 2       # las dos del pie
    assert info["fila_encabezado"] == 1


# ---------------------------------------------------------------- (12) la hoja

_DATOS = [
    ("CODIGO", "DESCRIPCION", "PRECIO"),
    ("1", "Cerveza PATRICIA lata. 473 ml", 74.5),
    ("2", "Mostaza Dijon MAILLE. 215 g", 199),
]
_PORTADA = [("CAMPAÑA ALEMANIA 2026",), ("Preparado por Marketing",)]


def test_una_portada_adelante_no_rechaza_el_archivo():
    """Se mandaba siempre la primera hoja: con una portada adelante, el archivo
    se rechazaba con "No encontré la fila de encabezados" —que apunta al
    problema equivocado— y ni siquiera nombraba las hojas."""
    m = planilla.leer(_xlsx(_DATOS, {"Portada": _PORTADA}), "listado.xlsx").mailing
    assert m["planilla"]["hoja"] == "Campaña"
    assert len(m["productos"]) == 2
    assert any("Portada" in a for a in planilla.avisos(m))


def test_con_dos_hojas_de_datos_se_dice_cual_se_leyo_y_cuales_son_las_otras():
    """Si la primera es un borrador viejo, se validaba contra el borrador sin
    chistar. Se sigue leyendo la primera —elegir por la persona es lo que no hay
    que hacer— pero se avisa, nombrando las hojas."""
    m = planilla.leer(_xlsx(_DATOS, {"Borrador": _DATOS}), "listado.xlsx").mailing
    assert m["planilla"]["hoja"] == "Borrador"
    aviso = next(a for a in planilla.avisos(m) if "más de una hoja" in a)
    assert "«Borrador»" in aviso and "«Campaña»" in aviso


def test_si_ninguna_hoja_tiene_datos_el_error_nombra_las_hojas():
    datos = _xlsx([("otra cosa",)], {"Portada": _PORTADA})
    with pytest.raises(planilla.PlanillaInvalida) as exc:
        planilla.leer(datos, "listado.xlsx")
    assert "Portada" in str(exc.value) and "Campaña" in str(exc.value)


# ---------------------------------------------------------------- (13) sin CODIGO

def test_una_planilla_con_descripcion_y_precio_alcanza():
    """Se exigía una columna CODIGO que después no se usa nunca: el código sale
    de `detectar_fila_headers`, que es del Convertidor y la necesita para armar
    una cenefa. Acá lo que amarra la placa con su fila es la DESCRIPCIÓN."""
    crudo = "DESCRIPCION;PRECIO\nCerveza PATRICIA lata. 473 ml;74,5\nMostaza Dijon MAILLE. 215 g;199\n"
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert [p["descripcion"] for p in m["productos"]] == [
        "Cerveza PATRICIA lata. 473 ml", "Mostaza Dijon MAILLE. 215 g",
    ]
    # sin columna MONEDA el precio va sin símbolo: ver el bloque (F)
    assert m["productos"][0]["oferta_precio"] == "74,50"
    assert "codigo" not in m["planilla"]["titulos"]


def test_los_encabezados_se_encuentran_aunque_haya_un_titulo_arriba_y_no_haya_codigo():
    crudo = (
        "LISTADO DE CAMPAÑA - ALEMANIA 2026;;\n"
        ";;\n"
        "DESCRIPCION;PRECIOANT;PRECIO\n"
        "Cerveza PATRICIA lata. 473 ml;99;74,5\n"
    )
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["fila_encabezado"] == 3
    assert m["productos"][0]["fila"]["numero"] == 4


# ==========================================================================
# Segunda vuelta: lo que la primera dejó abierto
# ==========================================================================

# ---------------------------------------------------------------- (A) el pie, en su forma común

_CAB = "CODIGO;DESCRIPCION;MONEDA;PRECIOANT;PRECIO\n"
_CUERPO = (
    "501233;Cerveza PATRICIA lata. 473 ml;$;99;74,5\n"
    "501235;Mostaza Dijon MAILLE. 215 g;$;245;199\n"
    "501236;Pan negro de centeno. 500 g;$;165;129\n"
)


@pytest.mark.parametrize("pie, que_trae", [
    (";TOTAL: 3 artículos;;;", "nada más"),
    (";TOTAL: 3 artículos;;;402,50", "el total de plata en la columna del precio"),
    ("3;TOTAL: 3 artículos;;;", "el contador de registros en CODIGO"),
    (";Generado el 18/09/2026;;;18/09/2026", "la fecha de emisión en otra columna"),
])
def test_el_pie_de_gestion_no_es_un_producto_aunque_traiga_su_total(pie, que_trae):
    """Se pedía que la fila no trajera NINGÚN otro dato, y un pie de gestión de
    verdad trae el total de plata, o el contador de registros, o la fecha: los
    tres entraban como un producto más y el Excel le decía al diseñador que le
    faltaba la placa de "TOTAL: 3 artículos", sin un solo aviso.

    La señal no es "no trae nada más": es estar al final del bloque, sin código
    de artículo, con un texto de pie y con un AGREGADO en las otras columnas
    (ver `_pies_del_final`)."""
    m = planilla.leer((_CAB + _CUERPO + pie + "\n").encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["filas_leidas"] == 3, que_trae
    assert all("TOTAL" not in p["descripcion"] and "Generado" not in p["descripcion"]
               for p in m["productos"])
    assert any("pie del reporte" in a for a in planilla.avisos(m))


def test_un_producto_que_se_llama_total_sigue_siendo_un_producto_aunque_este_ultimo():
    """El control que no se puede romper. "Total Blue" es una marca: al final
    del archivo, con su código y su precio, sigue siendo un producto —su precio
    no es el total de nada— y no se avisa nada."""
    crudo = (
        "CODIGO;DESCRIPCION;PRECIO\n"
        "501235;Mostaza Dijon MAILLE. 215 g;199\n"
        "501233;Total Blue agua saborizada. 900 ml;74,5\n"
    )
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert [p["descripcion"] for p in m["productos"]][-1] == "Total Blue agua saborizada. 900 ml"
    assert planilla.avisos(m) == []


def test_sin_columna_codigo_el_producto_que_se_llama_total_tampoco_se_cae():
    """Sin CODIGO la señal que queda es el agregado: 74,5 no es la suma de
    ninguna columna ni la cantidad de filas, así que es un precio y la fila es
    un producto."""
    crudo = (
        "DESCRIPCION;PRECIO\n"
        "Mostaza Dijon MAILLE. 215 g;199\n"
        "Total Blue agua saborizada. 900 ml;74,5\n"
    )
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert len(m["productos"]) == 2


def test_una_fila_que_parece_el_pie_pero_trae_datos_de_producto_se_dice():
    """El motor no la puede decidir: la cuenta como producto —que es lo menos
    destructivo— pero lo DICE. Resolverlo en silencio para cualquiera de los dos
    lados es lo que no se hace."""
    crudo = _CAB + _CUERPO + ";Resumen de temporada, caja;$;90;79\n"
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["filas_leidas"] == 4
    aviso = next(a for a in planilla.avisos(m) if "Parece el pie del reporte" in a)
    assert "fila 5" in aviso and "Resumen de temporada, caja" in aviso


def test_el_pie_del_medio_no_existe():
    """Un pie está al final. Una fila que empieza como un pie con productos
    abajo es un producto: por eso se camina de abajo para arriba."""
    crudo = _CAB + ";TOTAL: 1 artículo;;;\n" + _CUERPO
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["filas_leidas"] == 4


# ---------------------------------------------------------------- (B) el listado en formato largo

def _listado_largo(productos: int = 30, sucursales=("Central", "Punta Carretas", "Portones")) -> bytes:
    """El export de gestión que la gente tiene más a mano: una fila por producto
    Y POR SUCURSAL (ver convertidor.juntar_filas_por_producto)."""
    filas = ["ID_PRODUCTO;DESCRIPCION;SUCURSAL;STOCK;PRECIO\n"]
    for n in range(productos):
        for s, sucursal in enumerate(sucursales):
            filas.append(f"{500000 + n};Producto número {n}. 500 g;{sucursal};{s + 1};{100 + n}\n")
    return "".join(filas).encode("utf-8")


def test_el_listado_largo_entra_como_un_producto_por_descripcion():
    """90 filas repetidas: todas las placas quedaban sin pareja (dos filas
    idénticas empatan y el emparejamiento no adivina) y no se decía una palabra.
    Se junta con el juntador del Convertidor, que ya resolvía exactamente esto."""
    m = planilla.leer(_listado_largo(), "listado.csv").mailing
    assert m["planilla"]["filas_leidas"] == 30
    assert len({p["descripcion"] for p in m["productos"]}) == 30
    assert m["planilla"]["filas_juntadas"] == 90


def test_se_dice_que_se_juntaron_filas_y_cuantas_quedaron():
    m = planilla.leer(_listado_largo(), "listado.csv").mailing
    aviso = next(a for a in planilla.avisos(m) if "formato largo" in a)
    assert "90 filas" in aviso and "30 productos" in aviso and "3 sucursales" in aviso


def test_la_placa_empareja_contra_el_listado_largo():
    """Lo que importa de verdad: con las filas juntadas la placa vuelve a tener
    pareja, y la cita es la fila donde el producto aparece por PRIMERA vez."""
    m = planilla.leer(_listado_largo(productos=3), "listado.csv").mailing
    p = placa_de(descripcion="Producto número 1. 500 g", oferta_precio="101")
    idx, _, _ = c.comparar_placa(p, m, CONFIG)
    assert idx == 1
    assert m["productos"][1]["fila"]["numero"] == 5   # encabezado + las 3 filas del producto 0


def test_la_sucursal_y_el_stock_no_quedan_como_columnas_del_producto():
    """Eran datos de UNA de las filas que se juntaron: dejarlos invita a leerlos
    como si fueran del producto ("este auricular es de Central")."""
    m = planilla.leer(_listado_largo(productos=2), "listado.csv").mailing
    assert "sucursal" not in m["planilla"]["titulos"]
    assert "stock" not in m["planilla"]["titulos"]


def test_una_planilla_normal_no_pasa_por_el_juntador():
    """Sin columna SUCURSAL no se junta nada y las filas quedan TAL CUAL: los
    listados de siempre, que son la enorme mayoría, se leen igual que antes."""
    m = planilla.leer((_CAB + _CUERPO).encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["filas_juntadas"] == 0
    assert [p["fila"]["numero"] for p in m["productos"]] == [2, 3, 4]
    assert not any("formato largo" in a for a in planilla.avisos(m))


def test_si_las_filas_de_un_producto_se_contradicen_se_avisa():
    """El juntador se queda con el primer valor no vacío. Que dos filas del
    mismo producto traigan precios distintos es justo lo que hay que decir."""
    crudo = (
        "ID_PRODUCTO;DESCRIPCION;SUCURSAL;PRECIO\n"
        "1;Cerveza PATRICIA lata. 473 ml;Central;74,5\n"
        "1;Cerveza PATRICIA lata. 473 ml;Portones;99\n"
    )
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert m["productos"][0]["oferta_precio"] == "74,50"
    assert any("datos distintos entre sus filas" in a for a in planilla.avisos(m))


# ---------------------------------------------------------------- (C) el tope, sobre lo que cuesta

def test_la_tira_se_dibuja_cuando_se_pide_y_no_al_leer(monkeypatch):
    """El trabajo caro de este módulo era dibujar una tira por fila, y el 99 %
    se tiraba: solo se muestran las filas que alguna placa reclamó. 900 filas
    eran 5,92 s de CPU y 19,5 MB guardados, adentro del único hilo de CPU del
    servicio, que deja pasar UN trabajo a la vez."""
    dibujadas = []
    original = planilla.tira

    def espiar(mailing, indice, *a, **k):
        dibujadas.append(indice)
        return original(mailing, indice, *a, **k)

    monkeypatch.setattr(planilla, "tira", espiar)
    crudo = "CODIGO;DESCRIPCION;PRECIO\n" + "".join(
        f"{n};Producto número {n}. 500 g;{100 + n}\n" for n in range(900)
    )
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing

    assert len(m["productos"]) == 900
    assert dibujadas == [], "se dibujaron tiras al leer la planilla"
    assert all(p["recorte"] is None for p in m["productos"])
    # y cuando se pide, se dibuja esa y nada más
    assert planilla.evidencia(m, 7).startswith("data:image/jpeg;base64,")
    assert dibujadas == [7]


def test_el_tope_de_filas_sale_del_archivo_de_reglas():
    """El número estaba escrito en planilla.py, o sea en ningún lado que el
    navegador pueda leer. Vive en el mismo JSON que los tipos y los MB, que es
    lo que la pantalla recibe entero en GET /rrss/config."""
    assert planilla.MAX_FILAS == archivos.max_filas("planilla")
    assert archivos.TIPOS["fuentes"]["planilla"]["max_filas"] == planilla.MAX_FILAS


def test_el_excel_dice_cuando_no_pudo_dibujar_la_fila(mailing, tmp_path):
    """Que la tira falte no deja un hueco mudo en la columna B: se dice, y la
    cita (archivo · hoja · fila) alcanza para ir a mirarla a mano."""
    v = _validacion(mailing)
    v["imagenes"][0]["recorte_fuente"] = None
    ruta = tmp_path / "correcciones.xlsx"
    ruta.write_bytes(excel.construir(v))
    ws = load_workbook(ruta, rich_text=True)["Recomendada"]
    textos = [str(ws[f"B{r}"].value or "") for r in range(4, 12)]
    assert any("No pude dibujar la fila" in t for t in textos)


# ---------------------------------------------------------------- (E) el tope global nombra la fuente

def test_el_limite_de_50_mb_nombra_lo_que_se_subio():
    """"El mailing supera el límite de 50 MB" para alguien que subió una
    planilla lo manda a buscar un archivo que no existe. El nombre sale de
    `nombre_de_la_fuente`, el mismo que usan la pantalla y el Excel."""
    import asyncio

    from fastapi import HTTPException

    from app.core.uploads import read_limited
    from app.services.rrss import correccion

    class _Subido:
        async def read(self):
            return b"x" * (51 * 1024 * 1024)

    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_limited(_Subido(), correccion.nombre_de_la_fuente("planilla").capitalize()))
    assert "La planilla supera el límite de 50 MB" in exc.value.detail
    assert "mailing" not in exc.value.detail.lower()

    # y que la RUTA le pase eso, y no la palabra "mailing" escrita a mano: ahí
    # es donde se decide cómo se llama lo que se subió.
    from app.api.routes import rrss as ruta_rrss

    linea = next(l for l in pathlib.Path(ruta_rrss.__file__).read_text(encoding="utf-8").splitlines()
                 if "read_limited(mailing," in l)
    assert "nombre_de_la_fuente(origen)" in linea, linea


# ---------------------------------------------------------------- (F) sin columna MONEDA no hay símbolo

def test_sin_columna_moneda_el_precio_no_se_inventa_en_pesos():
    """Se le ponía "$" igual, y ahí el motor afirmaba en pesos algo que la
    planilla nunca dijo. Mismo criterio que la placa sin símbolo: lo que no se
    sabe no se afirma."""
    crudo = "DESCRIPCION;PRECIOANT;PRECIO\nSet de jarras cerveceras. 2 unidades;49;39\n"
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert m["productos"][0]["oferta_precio"] == "39"
    assert m["productos"][0]["precio_anterior"] == "49"
    assert c.moneda(m["productos"][0]["oferta_precio"]) == ""


def test_una_placa_en_dolares_contra_una_planilla_sin_moneda_no_se_acusa():
    """Es la consecuencia que importa: sin columna MONEDA la planilla no dice en
    qué moneda está, así que no se marca ningún error de moneda —ni al revés,
    como pasaba cuando el motor le ponía "$" de prepo."""
    crudo = "DESCRIPCION;PRECIO\nSet de jarras cerveceras. 2 unidades;39\n"
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    p = placa_de(descripcion="Set de jarras cerveceras. 2 unidades", oferta_precio="U$S39")
    _, _, filas = c.comparar_placa(p, m, CONFIG)
    assert por_campo(filas)["oferta_precio"]["estado"] == "ok"


def test_con_columna_moneda_el_simbolo_sigue_siendo_el_de_la_planilla():
    """Lo que ya andaba: si la planilla LO DICE, se usa."""
    m = planilla.leer((_CAB + _CUERPO).encode("utf-8"), "listado.csv").mailing
    assert m["productos"][0]["oferta_precio"] == "$74,50"


# ==========================================================================
# Tercera vuelta: los avisos que sobraban y los topes que contaban mal
# ==========================================================================

# ---------------------------------------------------------------- (G) el aviso que no correspondía

def _con_sucursal(productos: int = 12, sucursal: str = "Central") -> bytes:
    """Una planilla NORMAL que encima trae una columna SUCURSAL: un producto por
    fila, ninguno repetido. No hay nada que juntar."""
    filas = ["CODIGO;DESCRIPCION;SUCURSAL;PRECIO\n"]
    filas += [f"{501000 + n};Producto número {n}. 500 g;{sucursal};{100 + n}\n" for n in range(productos)]
    return "".join(filas).encode("utf-8")


def test_una_columna_sucursal_sin_nada_que_juntar_no_avisa_nada():
    """EL CASO NORMAL TIENE QUE SER CERO AVISOS. El aviso de formato largo salía
    por la SOLA existencia de la columna SUCURSAL: 12 productos distintos, una
    fila cada uno, todos en «Central», y el motor avisaba "La planilla viene en
    formato largo (una fila por producto y por sucursal...)" sin haber juntado
    una sola fila. Un aviso que salta cuando no corresponde es tan malo como el
    que falta."""
    m = planilla.leer(_con_sucursal(), "listado.csv").mailing
    assert len(m["productos"]) == 12
    assert planilla.avisos(m) == []
    assert m["planilla"]["filas_juntadas"] == 0


def test_sin_juntar_nada_la_sucursal_sigue_siendo_una_columna_de_la_fila():
    """Y es la consecuencia de lo de arriba: si no se juntó, el dato NO es de
    "una de las filas que se juntaron" —es de la fila, que es el producto— así
    que la columna se queda donde estaba."""
    m = planilla.leer(_con_sucursal(productos=3), "listado.csv").mailing
    assert m["planilla"]["titulos"]["sucursal"] == "SUCURSAL"
    assert [p["fila"]["numero"] for p in m["productos"]] == [2, 3, 4]


def test_el_listado_largo_de_verdad_sigue_avisando():
    """El otro lado del mismo control: cuando SÍ se juntó, el aviso está."""
    m = planilla.leer(_listado_largo(productos=4), "listado.csv").mailing
    assert m["planilla"]["filas_juntadas"] == 12
    assert any("formato largo" in a for a in planilla.avisos(m))


# ---------------------------------------------------------------- (H) las alternativas muertas del pie

def test_todas_las_alternativas_del_pie_matchean_su_propio_ejemplo():
    """Dos alternativas no podían matchear NADA y nadie se enteró: una
    alternativa muerta no rompe, solo deja pasar el pie como si fuera un
    producto. A las dos las mataba la frontera de palabra del final
    ("usuario:" + espacio no tiene frontera; "emisi" + "ón" tampoco).

    Este test recorre la tabla ENTERA, así la que agreguen mañana tampoco se
    puede morir en silencio."""
    muertas = [
        patron for patron, ejemplo in planilla.INICIOS_DE_PIE
        if not re.match(r"^\s*(?:%s)" % patron, ejemplo, re.IGNORECASE)
    ]
    assert muertas == [], "estas alternativas no pueden matchear ni su propio ejemplo: %s" % muertas
    # y el regex armado con la tabla las encuentra a todas
    for patron, ejemplo in planilla.INICIOS_DE_PIE:
        assert planilla._INICIO_DE_PIE.match(ejemplo), patron


@pytest.mark.parametrize("pie, que_es", [
    (";usuario: ADMIN;;;", "el usuario que exportó"),
    (";Fecha de emisión: 18/09/2026;;;18/09/2026", "la fecha de emisión, con tilde"),
    (";Fecha de emision: 18/09/2026;;;18/09/2026", "la misma, sin tilde"),
])
def test_las_dos_formas_de_pie_que_el_regex_no_podia_ver(pie, que_es):
    """La consecuencia de las alternativas muertas, en un archivo de verdad: el
    Excel le decía al diseñador que le faltaba la placa de «usuario: ADMIN»."""
    m = planilla.leer((_CAB + _CUERPO + pie + "\n").encode("utf-8"), "listado.csv").mailing
    assert m["planilla"]["filas_leidas"] == 3, que_es
    assert any("pie del reporte" in a for a in planilla.avisos(m))


# ---------------------------------------------------------------- (I) el pie que está arriba de una dudosa

def test_un_pie_limpio_arriba_de_una_fila_dudosa_se_evalua_igual():
    """`_pies_del_final` caminaba de abajo para arriba y CORTABA en la primera
    fila dudosa, así que lo que estaba más arriba entraba como producto sin un
    aviso propio: con dos pies, el "TOTAL: 3 artículos" se colaba entero.

    Una dudosa tiene cara de pie —lo único que el motor no puede decidir es qué
    trae en las otras columnas— y eso no dice nada de la fila de más arriba."""
    crudo = _CAB + _CUERPO + ";TOTAL: 3 artículos;;;402,50\n;Generado el 18/09/2026;;;por USUARIO\n"
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing

    # el TOTAL es el pie que era, y la dudosa se cuenta como producto y se dice
    assert m["planilla"]["filas_leidas"] == 4
    assert all("TOTAL" not in p["descripcion"] for p in m["productos"])
    assert m["productos"][-1]["descripcion"] == "Generado el 18/09/2026"
    limpio = next(a for a in planilla.avisos(m) if a.startswith("No lo conté como producto"))
    assert "TOTAL: 3 artículos" in limpio
    duda = next(a for a in planilla.avisos(m) if "Parece el pie del reporte" in a)
    assert "Generado el 18/09/2026" in duda


def test_el_producto_de_arriba_sigue_cortando_el_camino():
    """Lo que NO cambió: se sube hasta la primera fila que no tiene cara de pie.
    "Total Blue" al final es un producto y ahí se corta —no se sigue subiendo a
    acusar al resto del listado."""
    crudo = (
        "DESCRIPCION;PRECIO\n"
        "Mostaza Dijon MAILLE. 215 g;199\n"
        "Total Blue agua saborizada. 900 ml;74,5\n"
    )
    m = planilla.leer(crudo.encode("utf-8"), "listado.csv").mailing
    assert [p["descripcion"] for p in m["productos"]] == [
        "Mostaza Dijon MAILLE. 215 g", "Total Blue agua saborizada. 900 ml",
    ]


# ---------------------------------------------------------------- (J) el tope cuenta productos

def test_un_listado_largo_de_350_productos_no_es_un_catalogo():
    """350 productos por 3 sucursales son 1.050 filas, y el tope las miraba
    CRUDAS: el motor rechazaba con "subí el listado de esta campaña, no el
    catálogo entero" justo a quien había subido el listado de la campaña. Ahora
    que el formato largo es una entrada de primera clase, el tope cuenta
    PRODUCTOS."""
    m = planilla.leer(_listado_largo(productos=350), "listado.csv").mailing
    assert len(m["productos"]) == 350
    assert m["planilla"]["filas_juntadas"] == 1050
    assert planilla.MAX_FILAS == 1000 < 1050


def test_mas_productos_que_el_tope_sigue_siendo_el_catalogo():
    """El tope no se aflojó: lo que cambió es QUÉ cuenta. 1.001 productos son
    1.001 productos vengan en las filas que vengan."""
    crudo = "CODIGO;DESCRIPCION;PRECIO\n" + "".join(
        f"{n};Producto número {n}. 500 g;{100 + n}\n" for n in range(planilla.MAX_FILAS + 1)
    )
    with pytest.raises(planilla.PlanillaInvalida) as exc:
        planilla.leer(crudo.encode("utf-8"), "listado.csv")
    assert "más de %d productos" % planilla.MAX_FILAS in str(exc.value)
    assert "catálogo entero" in str(exc.value)


def test_el_catalogo_enorme_se_sigue_rechazando_sin_materializarlo(monkeypatch):
    """EL FRENO BARATO NO SE PUEDE PERDER. Materializar el archivo para poder
    contar los productos es exactamente lo que este freno evita: 1.000.000 de
    filas eran 148,8 s y 264 MB adentro del único hilo de CPU del servicio, con
    todas las validaciones de todo el mundo frenadas mientras tanto.

    Lo que se fija no es el reloj sino que NO SE LLAME AL LECTOR: el tiempo va
    igual porque es lo que se siente del otro lado. Medido con este mismo CSV de
    200.000 filas: 0,006 s y 0,00 MB de pico con el freno puesto, 1,12 s y 86 MB
    sin él."""
    crudo = b"CODIGO;DESCRIPCION;PRECIO\n" + b"".join(
        b"%d;Producto numero %d. 500 g;100\n" % (n, n) for n in range(200_000)
    )
    materializadas = []
    original = planilla.leer_filas
    monkeypatch.setattr(planilla, "leer_filas",
                        lambda *a, **k: materializadas.append(True) or original(*a, **k))

    arranque = time.perf_counter()
    with pytest.raises(planilla.PlanillaInvalida) as exc:
        planilla.leer(crudo, "catalogo.csv")
    tardo = time.perf_counter() - arranque

    assert "catálogo entero" in str(exc.value)
    assert materializadas == [], "se leyó el catálogo entero antes de rechazarlo"
    assert tardo < 0.5, "tardó %.2f s: se está materializando el catálogo antes de rechazarlo" % tardo


def test_el_freno_barato_se_mide_en_el_peor_formato_largo():
    """El número del freno no es un invento suelto: es el tope de productos por
    el techo de sucursales de la cadena (18 reales, 30 de techo). Un archivo más
    grande que eso es un catálogo con cualquier cuenta que se haga."""
    assert planilla.MAX_FILAS_CRUDAS == planilla.MAX_FILAS * planilla._SUCURSALES_TECHO
    assert planilla.MAX_FILAS_CRUDAS > planilla.MAX_FILAS


# ---------------------------------------------------------------- (22/09) filas vacías, columnas del informe, la cola del precio

def test_las_filas_vacias_no_son_filas_que_quedaron_afuera():
    """Un renglón en blanco entre el título y los datos, o los que Excel deja al
    final, hacían decir "Afuera: 3 filas sin descripción" en una planilla
    perfecta. Una fila con código o precio y sin descripción SÍ se cuenta."""
    limpia = planilla.leer(_xlsx([
        ("CODIGO", "DESCRIPCION", "PRECIO"),
        ("501233", "Cerveza PATRICIA lata. 473 ml", 74.5),
        (None, None, None),
        ("501235", "Mostaza Dijon MAILLE. 215 g", 199),
        (None, None, None),
        (None, None, None),
    ]), "limpia.xlsx")
    assert not any("sin descripción" in a for a in limpia.mailing["planilla"]["avisos"]), limpia.mailing["planilla"]["avisos"]
    assert limpia.mailing["planilla"]["filas_ignoradas"] == 0
    assert len(limpia.mailing["productos"]) == 2

    con_hueco = planilla.leer(_xlsx([
        ("CODIGO", "DESCRIPCION", "PRECIO"),
        ("501233", "Cerveza PATRICIA lata. 473 ml", 74.5),
        ("501299", None, 120),   # un producto al que le falta la descripción: eso sí se dice
    ]), "hueco.xlsx")
    assert "Afuera: 1 fila sin descripción" in con_hueco.mailing["planilla"]["avisos"]
    assert con_hueco.mailing["planilla"]["filas_ignoradas"] == 1


def test_los_titulos_del_informe_entran_como_columnas():
    """"ARRIBA DEL PRECIO" y "ABAJO DEL PRECIO" son los títulos que el propio
    Excel usa: si alguien los copia a su planilla tienen que entrar. Y FECHA /
    LEYENDA DE ALCOHOL también."""
    pl = planilla.leer(_xlsx([
        ("DESCRIPCIÓN", "PRECIO", "MECÁNICA", "ARRIBA DEL PRECIO", "ABAJO DEL PRECIO", "FECHA", "LEYENDA DE ALCOHOL"),
        ("Arvejas TIENDA INGLESA. 300 g", 37.5, "2x$75", "Comprando 2", "unidad",
         "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE", "Beber con moderación."),
    ]), "informe.xlsx")
    m = pl.mailing
    assert set(m["campos"]) >= {"mecanica", "oferta_encabezado", "oferta_pie", "oferta_precio", "descripcion"}
    assert m["fecha"] == "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE"
    assert m["legal_alcohol"] == "Beber con moderación."
    assert planilla.campos_que_no_dicta(m, CONFIG) == ["Precio anterior", "Si el precio anterior va tachado"]
    # y los nombres que la pantalla muestra son estos mismos
    assert [n for _, n in planilla.COLUMNAS_OPCIONALES] == [
        "MECÁNICA", "ARRIBA DEL PRECIO", "ABAJO DEL PRECIO", "VIGENCIA", "LEYENDA ALCOHOL",
    ]


def test_campos_que_no_dicta_incluye_la_fecha_y_la_leyenda_solo_si_nadie_las_dio(mailing):
    con_fecha = planilla.campos_que_no_dicta(mailing, {**CONFIG, "fecha": "DEL JUEVES 17 AL DOMINGO 20"})
    assert "Fecha de la campaña" not in con_fecha and "Leyenda de alcohol" in con_fecha
    con_todo = planilla.campos_que_no_dicta(mailing, {**CONFIG, "fecha": "x", "legal_alcohol": "Beber con moderación."})
    assert con_todo == ["Mecánica", "Texto arriba del precio", "Texto abajo del precio", "Si el precio anterior va tachado"]
    assert planilla.campos_que_no_dicta({"origen": "mailing"}, {}) == []


def test_dicta_todo(mailing):
    assert planilla.dicta_todo({"origen": "mailing"}) is True
    assert planilla.dicta_todo(mailing) is False
    assert planilla.dicta_todo({"origen": "planilla", "campos": list(planilla._DE_LA_COLUMNA)}) is True


def test_el_precio_conserva_lo_que_la_celda_trae_alrededor_del_numero():
    """"$340 unidad" quedaba en "$340": se perdía justo la palabra que la
    persona había escrito. Y sin columna MONEDA, el símbolo que la celda misma
    trae ("U$S 149") tampoco se inventa ni se tira."""
    assert planilla._precio("$340 unidad", "$") == "$340 unidad"
    assert planilla._precio("U$S 149", "") == "U$S149"
    assert planilla._precio(74.5, "$") == "$74,50"
    assert planilla._precio(1090, "") == "1.090"
    assert planilla._precio("Oferta", "$") == "Oferta"


def _planilla_con_colas() -> dict:
    return planilla.leer(_xlsx([
        ("DESCRIPCION", "MONEDA", "PRECIO ANTERIOR", "PRECIO"),
        ("Cerveza STELLA ARTOIS. Lata 710 ml", "$", "$171", "$125"),
        ("Vino BRISAS DEL ESTE. 750 ml", "$", "$340 unidad", "$199"),
    ]), "colas.xlsx").mailing


def test_si_la_planilla_escribe_lo_que_acompana_al_precio_lo_dicta():
    """La misma campaña validada contra el PDF y contra su planilla tiene que
    dar lo mismo. En el PDF, "$171" contra la placa "$171 unidad" es un error
    (sobra «unidad»). Con la planilla el precio se compara por importe y ese
    error se perdía; ahora, si alguna fila de la columna escribió la cola
    ("$340 unidad"), la columna dicta la cola en todas sus filas."""
    m = _planilla_con_colas()
    assert m["precios_con_cola"] == ["precio_anterior"]
    stella = placa_de(descripcion="Cerveza STELLA ARTOIS. Lata 710 ml", precio_anterior="$171 unidad", oferta_precio="$125")
    _, _, filas = c.comparar_placa(stella, m, CONFIG)
    fila = por_campo(filas)["precio_anterior"]
    assert fila["severidad"] == "error" and fila["estado"] == "diferente"
    from app.services.rrss import correccion
    assert correccion.instruccion(fila["estado"], fila["placa"], fila["mailing"], "planilla") == "Sobra «unidad»: sacalo"
    vino = placa_de(descripcion="Vino BRISAS DEL ESTE. 750 ml", precio_anterior="$340", oferta_precio="$199")
    _, _, filas = c.comparar_placa(vino, m, CONFIG)
    fila = por_campo(filas)["precio_anterior"]
    assert fila["severidad"] == "error"
    assert correccion.instruccion(fila["estado"], fila["placa"], fila["mailing"], "planilla") == "Falta «unidad»: agregalo"
    # el precio de oferta no escribió cola en ninguna fila: sigue siendo solo el importe
    stella_pie = placa_de(descripcion="Cerveza STELLA ARTOIS. Lata 710 ml", precio_anterior="$171", oferta_precio="$125 c/u")
    _, _, filas = c.comparar_placa(stella_pie, m, CONFIG)
    assert por_campo(filas)["oferta_precio"]["estado"] == "ok"


def test_una_planilla_de_numeros_sigue_comparando_solo_el_importe(mailing):
    """El export de gestión trae números: ahí nadie escribió una cola y "$99
    unidad" contra 99 sigue estando bien (es el test de arriba, de nuevo, para
    que la regla nueva no lo pise)."""
    assert mailing.get("precios_con_cola") == []
    p = placa_de(descripcion="Cerveza PATRICIA lata. 473 ml", precio_anterior="$99 unidad", oferta_precio="$74,50")
    _, _, filas = c.comparar_placa(p, mailing, CONFIG)
    assert por_campo(filas)["precio_anterior"]["estado"] == "ok"
