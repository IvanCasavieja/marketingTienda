"""Borrar un cuadro en el preview tiene que borrarlo del archivo, no solo de la pantalla.

El render parte del PPTX fuente: si el cuadro solo se sacara de la lista, su
forma original seguiría impresa en todas las hojas. Caso real (11/09/2026): un
"9" blanco olvidado en el diseño de Rompe Precios Congelados, que en el preview
tapaba al cuadro de "PRECIO REGULAR".

Se prueba de punta a punta: PPTX armado en memoria -> importer -> render.
"""
import io
import re
import zipfile

from pptx import Presentation
from pptx.util import Cm, Pt

from app.services.cenefas.component_renderer import render_template_to_pptx
from app.services.cenefas.jobs import aplicar_overrides
from app.services.cenefas.pptx_importer import import_pptx

PRODUCTOS = [{"precioOferta": "253"}, {"precioOferta": "199"}]


def _pptx():
    prs = Presentation()
    prs.slide_width = Cm(21.0)
    prs.slide_height = Cm(29.7)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for texto, y in (("<<precioOferta>>", 3), ("9", 12)):
        box = slide.shapes.add_textbox(Cm(2), Cm(y), Cm(17), Cm(4))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = texto
        run.font.size = Pt(60)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _textos(pptx_bytes):
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
        xml = "".join(
            z.read(n).decode("utf-8", "ignore")
            for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
    return re.findall(r"<a:t>([^<]*)</a:t>", xml)


def _nueve(definicion):
    return next(c for c in definicion["components"] if c.get("static_value") == "9")


def test_sin_eliminar_el_cuadro_se_imprime():
    src = _pptx()
    pptx, _ = render_template_to_pptx(import_pptx(src), PRODUCTOS, "a4", None, src)
    assert "9" in _textos(pptx)


def test_eliminado_en_el_preview_no_sale_en_ninguna_hoja():
    src = _pptx()
    d = import_pptx(src)
    d = aplicar_overrides(d, [{"id": _nueve(d)["id"], "eliminado": True}])
    pptx, _ = render_template_to_pptx(d, PRODUCTOS, "a4", None, src)
    textos = _textos(pptx)
    assert "9" not in textos
    # El resto de la cenefa sale igual, en las dos hojas.
    assert "253" in textos and "199" in textos


def test_eliminado_guardado_en_la_plantilla_tampoco_sale():
    """Lo que deja "Guardar en la plantilla": el cuadro ya no está en la
    definición y su forma queda anotada en formas_eliminadas."""
    src = _pptx()
    d = import_pptx(src)
    nueve = _nueve(d)
    d = {
        **d,
        "components": [c for c in d["components"] if c is not nueve],
        "formas_eliminadas": [nueve["_source_shape_id"]],
    }
    pptx, _ = render_template_to_pptx(d, PRODUCTOS, "a4", None, src)
    assert "9" not in _textos(pptx)


def test_una_forma_que_usa_otro_cuadro_no_se_saca():
    """Si la lista trae por error la forma de un cuadro que sigue en pie, el
    cuadro se imprime igual."""
    src = _pptx()
    d = import_pptx(src)
    precio = next(c for c in d["components"] if "precioOferta" in str(c))
    d = {**d, "formas_eliminadas": [precio["_source_shape_id"]]}
    pptx, _ = render_template_to_pptx(d, PRODUCTOS, "a4", None, src)
    assert "253" in _textos(pptx)
