"""Un cuadro de texto fijo del diseño no desaparece por culpa de un vecino.

Hasta el 17/09/2026 el motor hacía esto: un cuadro SIN variable propia buscaba
de quién era acompañante --el que más cobertura vertical tuviera-- y se apagaba
junto con él. Existía para que el "$" escrito a mano no quedara solo cuando el
precio de al lado no se dibujaba.

Ese motivo desapareció cuando los 41 cuadros con el símbolo escrito a mano
pasaron a usar <<unidadMoneda>> (scripts/simbolo_a_variable.py). Sin "$" fijos,
los únicos clientes de la heurística eran las 56 ETIQUETAS del parque, y les
hacía daño: la palabra "OFERTA" desaparecía de la Gran Bretaña A4 en cuanto
`ofertaUno` venía vacía, porque las dos cajas están a la misma altura aunque
haya 5 cm entre ellas.

Es la misma decisión que Ivan ya había tomado el 07/09/2026 para la otra
pareja automática ("quita eso de la pareja automatica", ver _dollar_parejas):
adivinar de quién es acompañante un cuadro sale mal en silencio.
"""
import io
import re
import zipfile

from pptx import Presentation
from pptx.util import Cm, Pt

from app.services.cenefas.component_renderer import render_template_to_pptx
from app.services.cenefas.pptx_importer import import_pptx


def _pptx(*cajas, size_pt=40):
    """Un A4 con los cuadros en las posiciones dadas: (texto, x, y, w, h)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(21.0), Cm(29.7)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for texto, x, y, w, h in cajas:
        caja = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        run = caja.text_frame.paragraphs[0].add_run()
        run.text = texto
        run.font.size = Pt(size_pt)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _textos(pptx_bytes):
    prs = Presentation(io.BytesIO(pptx_bytes))
    return [sh.text_frame.text.strip() for sl in prs.slides for sh in sl.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]


def test_una_etiqueta_no_desaparece_con_un_vecino_lejano_y_vacio():
    # La geometría real de Gran Bretaña A4: "OFERTA" a la izquierda y la
    # cocarda <<ofertaUno>> a la derecha, a la MISMA altura y a 5 cm.
    src = _pptx(
        ("OFERTA",          3.09, 16.83, 7.24, 1.62),
        ("<<ofertaUno>>",  15.28, 15.85, 4.05, 2.39),
        ("<<precioOferta>>", 0.71, 20.00, 12.21, 4.00),
    )
    d = import_pptx(src)
    salida = _textos(render_template_to_pptx(
        d, [{"ofertaUno": "", "precioOferta": "250"}], "a4", None, src)[0])
    assert "OFERTA" in salida, f"la etiqueta se fue con su vecino vacío: {salida!r}"
    assert "250" in salida


def test_la_etiqueta_sobrevive_aunque_el_vecino_la_cubra_entera():
    # Peor caso: cobertura vertical del 100%. Ni así: una etiqueta del diseño
    # no depende del dato de nadie.
    src = _pptx(
        ("PRECIO REGULAR:", 1.00, 10.00, 6.00, 2.00),
        ("<<precioBanco>>", 14.00, 10.00, 5.00, 2.00),
    )
    d = import_pptx(src)
    salida = _textos(render_template_to_pptx(d, [{"precioBanco": ""}], "a4", None, src)[0])
    assert "PRECIO REGULAR:" in salida, f"se fue con el vecino vacío: {salida!r}"


def test_un_cuadro_con_variable_vacia_si_desaparece():
    # Lo que NO cambió: un cuadro cuyo contenido sale solo de variables y todas
    # vinieron vacías no tiene nada que imprimir y se saca. Sin esto quedaría
    # el rectángulo de la cocarda impreso en un producto sin mecánica.
    src = _pptx(("<<ofertaUno>>", 15.28, 15.85, 4.05, 2.39),
                ("<<precioOferta>>", 0.71, 20.00, 12.21, 4.00))
    d = import_pptx(src)
    salida = _textos(render_template_to_pptx(
        d, [{"ofertaUno": "", "precioOferta": "250"}], "a4", None, src)[0])
    assert salida == ["250"], f"esperaba solo el precio: {salida!r}"


def test_el_simbolo_como_variable_se_apaga_solo_sin_heuristica():
    # El caso para el que existía la heurística, resuelto sin ella: con
    # <<unidadMoneda>> el cuadro tiene variable propia, así que se apaga por la
    # regla de "todas sus variables vinieron vacías" y no por adivinanza.
    src = _pptx(("<<unidadMoneda>>", 0.40, 20.00, 2.70, 2.80),
                ("<<precioOferta>>", 3.20, 20.00, 12.21, 4.00))
    d = import_pptx(src)
    con_dato = _textos(render_template_to_pptx(
        d, [{"unidadMoneda": "$", "precioOferta": "250"}], "a4", None, src)[0])
    assert "$" in con_dato and "250" in con_dato

    sin_dato = _textos(render_template_to_pptx(
        d, [{"unidadMoneda": "", "precioOferta": "250"}], "a4", None, src)[0])
    assert "$" not in sin_dato, f"el símbolo quedó solo: {sin_dato!r}"


def test_el_simbolo_variable_sigue_la_moneda_de_la_fila():
    # La razón de fondo para pasar el "$" del diseño a variable: un producto en
    # dólares imprimía "$" porque el símbolo estaba escrito a mano.
    src = _pptx(("<<unidadMoneda>>", 0.40, 20.00, 2.70, 2.80),
                ("<<precioOferta>>", 3.20, 20.00, 12.21, 4.00))
    d = import_pptx(src)
    for moneda in ("$", "U$S"):
        salida = _textos(render_template_to_pptx(
            d, [{"unidadMoneda": moneda, "precioOferta": "250"}], "a4", None, src)[0])
        assert moneda in salida, f"con moneda {moneda!r} salió {salida!r}"
