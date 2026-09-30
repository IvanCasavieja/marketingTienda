"""El run espaciador tiene que EXISTIR en el papel.

`line_height_pt` es el pedazo invisible con el que el diseño fija el alto del
renglón: apoya el "$" chico en la línea del precio grande, y deja al precio en
su lugar cuando una regla o una voladita le achican el cuerpo. El preview lo
dibuja. El export lo escribe como un run extra al final del párrafo.

Caso real (Ivan, 30/09/2026, "Fiesta de Alemania", precio banco "$501,84"
con el 501 volado): en pantalla el precio quedaba centrado en la franja roja
de Club Card y en el archivo salía 0,16 cm más arriba, pegado al borde. El
run espaciador se escribía con texto "" y PowerPoint IGNORA un run vacío al
armar el renglón. Medido con POWERPNT.EXE (PPTX -> PDF, leído con PyMuPDF):
el archivo con ese run y el archivo sin él daban la línea de base en el mismo
lugar al centésimo de punto. Con un espacio adentro, el renglón toma su alto.

Por qué un espacio COMÚN y AL FINAL, también medido: PowerPoint cuelga los
espacios del final del renglón, así que no corren el centrado ni la alineación
a derecha o izquierda (0,00 pt). Un espacio duro al final, o uno común al
principio, corren el texto ~2 pt.
"""
from pptx import Presentation
from pptx.util import Cm

from app.services.cenefas.component_renderer import _populate_text_frame

_PRECIO_BANCO = {
    "id": "banco", "type": "text",
    "style": {"font_size": 16.0, "line_height_pt": 24.0, "font_family": "Impact",
              "color": "#FFFFFF", "align": "center"},
    "segments": [
        {"type": "variable", "value": "unidadMoneda", "_resolved": "$",
         "style": {"font_size": 16.0, "font_family": "Impact", "baseline": 20000}},
        {"type": "variable", "value": "precioBanco", "_resolved": "501",
         "style": {"font_size": 24.0, "font_family": "Impact", "baseline": 5000}},
        {"type": "variable", "value": "decimalPrecioBanco", "_resolved": ",84",
         "style": {"font_size": 14.0, "font_family": "Impact", "baseline": 75000}},
    ],
}


def _runs(comp):
    prs = Presentation()
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    caja = sl.shapes.add_textbox(Cm(1), Cm(1), Cm(3), Cm(2))
    _populate_text_frame(caja.text_frame, comp, "")
    return caja.text_frame.paragraphs[0].runs


def test_el_espaciador_no_esta_vacio():
    espaciador = _runs(_PRECIO_BANCO)[-1]
    assert espaciador.text != "", (
        "el run espaciador volvió a salir VACÍO: PowerPoint lo ignora y el "
        "precio sale más arriba en el papel que en el preview"
    )


def test_el_espaciador_es_un_espacio_comun_al_final():
    runs = _runs(_PRECIO_BANCO)
    assert [r.text for r in runs] == ["$", "501", ",84", " "], (
        "el espaciador va AL FINAL y es un espacio común: al principio, o con "
        "un espacio duro, corre el centrado del texto"
    )


def test_el_espaciador_tiene_el_cuerpo_y_la_fuente_de_la_caja():
    espaciador = _runs(_PRECIO_BANCO)[-1]
    assert espaciador.font.size.pt == 24.0
    # Sin la fuente heredaría la del tema (Calibri) y el alto del renglón
    # dependería de una tipografía que el preview no conoce.
    assert espaciador.font.name == "Impact"


def test_el_espaciador_no_va_volado():
    # Volado, PowerPoint lo dibujaría a dos tercios y dejaría de sostener el
    # renglón.
    espaciador = _runs(_PRECIO_BANCO)[-1]
    assert espaciador._r.get_or_add_rPr().get("baseline") is None


def test_sin_espaciador_en_el_diseno_no_se_agrega_ninguno():
    comp = {**_PRECIO_BANCO, "style": {k: v for k, v in _PRECIO_BANCO["style"].items()
                                       if k != "line_height_pt"}}
    assert [r.text for r in _runs(comp)] == ["$", "501", ",84"]
