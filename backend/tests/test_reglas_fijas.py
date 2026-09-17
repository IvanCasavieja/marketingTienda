"""Las reglas que el sistema garantiza y una resubida de PPTX no puede borrar.

El caso que las motivó (Ivan, 17/09/2026): un M x N imprimía "$2x1" en el
cuadro grande, porque el diseño dibuja el símbolo de moneda pegado a
`promoOferta` y esa variable lleva un número en un combo ("160") pero un
literal en un M x N ("2x1"). El símbolo corresponde solo cuando lo que sigue
es plata.
"""
import io

from pptx import Presentation
from pptx.util import Cm, Pt

from app.services.cenefas.component_renderer import render_template_to_pptx
from app.services.cenefas.pptx_importer import import_pptx
from app.services.cenefas.reglas_fijas import (
    CLAVE_BLOQUEADA,
    asegurar_reglas_fijas,
)


# ---------------------------------------------------------------------------
# Las cuatro anatomías reales de un cuadro de promoOferta, medidas sobre las
# 11 plantillas que hoy lo tienen (17/09/2026).
# ---------------------------------------------------------------------------

def _comp(comp_id, segments=None, variable=None, style=None):
    c = {"id": comp_id, "type": "text", "base_bounds": {"x": 1, "y": 1, "width": 10, "height": 3}}
    if segments is not None:
        c["segments"] = segments
    if variable:
        c["variable"] = variable
    if style:
        c["style"] = style
    return c


def _var(nombre, size=None):
    return {"type": "variable", "value": nombre, "style": {"font_size": size} if size else {}}


def _reglas_de(defin, comp_id):
    return [r for r in defin["rules"] if r["target_component_id"] == comp_id]


def test_anatomia_a_con_la_variable_ya_puesta():
    # [<<unidadMoneda>>, <<promoOferta>>] -- Red Expres 17 A4 y Cenefas 3xA4.
    d = asegurar_reglas_fijas({"components": [
        _comp("c1", [_var("unidadMoneda", 108), _var("promoOferta")], variable="promoOferta")]})
    reglas = _reglas_de(d, "c1")
    assert len(reglas) == 2
    assert all(r["target_segment_index"] == 0 for r in reglas), "el símbolo está en el segmento 0"
    assert all(r[CLAVE_BLOQUEADA] for r in reglas)
    # no se le agrega un segundo símbolo
    assert len(d["components"][0]["segments"]) == 2


def test_anatomia_b_con_el_simbolo_fijo_de_preciazos():
    # [<<tipoOferta>>, " $ ", <<promoOferta>>] -- el "$" es texto del diseño,
    # no una variable, y para cuándo-se-muestra da exactamente lo mismo.
    d = asegurar_reglas_fijas({"components": [_comp("c1", [
        _var("tipoOferta", 9),
        {"type": "static", "value": " $ ", "style": {"font_size": 9}},
        _var("promoOferta", 9)])]})
    reglas = _reglas_de(d, "c1")
    assert len(reglas) == 2
    assert all(r["target_segment_index"] == 1 for r in reglas), "el '$' fijo es el segmento 1"
    assert len(d["components"][0]["segments"]) == 3, "no se agrega nada: el símbolo ya está"


def test_anatomia_c_sin_simbolo_se_le_inserta():
    # [<<promoOferta>>] solo -- Inglaterra Bebidas Rural, Red Expres-202608.
    d = asegurar_reglas_fijas({"components": [_comp("c1", [_var("promoOferta", 97)])]})
    segs = d["components"][0]["segments"]
    assert [s["value"] for s in segs] == ["unidadMoneda", "promoOferta"]
    assert segs[0]["style"]["font_size"] == 97, "el símbolo hereda el cuerpo del número"
    assert all(r["target_segment_index"] == 0 for r in _reglas_de(d, "c1"))


def test_anatomia_d_sin_segmentos_se_convierte():
    # `variable: "promoOferta"` suelta, sin segments -- A4/3xA4 HELVETICO.
    d = asegurar_reglas_fijas({"components": [
        _comp("c1", variable="promoOferta", style={"font_size": 120, "line_height_pt": 130})]})
    segs = d["components"][0]["segments"]
    assert [s["value"] for s in segs] == ["unidadMoneda", "promoOferta"]
    assert segs[0]["style"]["font_size"] == 120
    assert "line_height_pt" not in segs[0]["style"], "el interlineado es del cuadro, no del run"
    assert d["components"][0]["variable"] == "promoOferta", "el cuadro conserva su variable"


def test_una_definicion_sin_promo_oferta_no_se_toca():
    d = asegurar_reglas_fijas({"components": [_comp("c1", [_var("precioOferta")])], "rules": []})
    assert d["rules"] == []


# ---------------------------------------------------------------------------
# Lo que las hace FIJAS
# ---------------------------------------------------------------------------

def test_es_idempotente():
    base = {"components": [_comp("c1", [_var("promoOferta", 50)])]}
    una = asegurar_reglas_fijas(base)
    dos = asegurar_reglas_fijas(una)
    assert len(dos["rules"]) == 2, f"se duplicaron: {dos['rules']}"
    assert dos["rules"] == una["rules"]
    assert len(dos["components"][0]["segments"]) == 2, "no se acumula un segundo símbolo"


def test_borrarla_no_la_borra():
    # Es el corazón del pedido: una persona la saca en la UI y guarda; el
    # guardado la repone.
    d = asegurar_reglas_fijas({"components": [_comp("c1", [_var("promoOferta")])]})
    sin_reglas = {**d, "rules": []}
    assert len(asegurar_reglas_fijas(sin_reglas)["rules"]) == 2


def test_las_reglas_de_las_personas_sobreviven():
    propia = {"id": "mia", "name": "Mostrar si codigo contiene -", "action": {"type": "show"},
              "condition": {"field": "codigo", "operator": "contains", "value": "-"},
              "target_component_id": "c1"}
    d = asegurar_reglas_fijas({"components": [_comp("c1", [_var("promoOferta")])],
                               "rules": [propia]})
    assert propia in d["rules"]
    assert len(d["rules"]) == 3
    assert d["rules"][-1] == propia, "las fijas van primero, la de la persona queda abajo"


def test_una_fija_vieja_que_ya_no_corresponde_se_descarta():
    # Quedó apuntando a un cuadro que el diseño ya no tiene: si sobreviviera,
    # ocultaría el segmento equivocado del cuadro que ocupe ese lugar.
    vieja = {"id": "fantasma", "name": "vieja", "action": {"type": "hide"},
             "condition": {"field": "promoOferta", "operator": "contains", "value": "x"},
             "target_component_id": "borrado", "target_segment_index": 0,
             CLAVE_BLOQUEADA: True}
    d = asegurar_reglas_fijas({"components": [_comp("c1", [_var("promoOferta")])],
                               "rules": [vieja]})
    assert vieja not in d["rules"]
    assert len(d["rules"]) == 2


# ---------------------------------------------------------------------------
# Punta a punta: importar un PPTX y renderizar
# ---------------------------------------------------------------------------

def _pptx_promo(texto="<<unidadMoneda>><<promoOferta>>"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(21.0), Cm(29.7)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    caja = slide.shapes.add_textbox(Cm(2), Cm(10), Cm(17), Cm(6))
    run = caja.text_frame.paragraphs[0].add_run()
    run.text = texto
    run.font.size = Pt(100)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _runs(pptx_bytes):
    import re, zipfile
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as z:
        xml = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                      if re.match(r"ppt/slides/slide\d+\.xml$", n))
    return "".join(re.findall(r"<a:t>([^<]*)</a:t>", xml))


def test_importar_un_pptx_ya_trae_las_reglas_puestas():
    # Un reimport pisaba `rules` con [] y se perdía todo. Ahora las fijas
    # vuelven en el mismo acto de importar, sin script a mano.
    d = import_pptx(_pptx_promo())
    fijas = [r for r in d["rules"] if r.get(CLAVE_BLOQUEADA)]
    assert len(fijas) == 2, f"el import no repuso las reglas fijas: {d['rules']}"


def test_mxn_no_imprime_el_simbolo_y_el_combo_si():
    src = _pptx_promo()
    d = import_pptx(src)
    mxn = _runs(render_template_to_pptx(
        d, [{"promoOferta": "2x1", "unidadMoneda": "$"}], "a4", None, src)[0])
    assert "2x1" in mxn, f"se perdió el literal del M x N: {mxn!r}"
    assert "$" not in mxn, f"el M x N salió con símbolo de moneda: {mxn!r}"

    combo = _runs(render_template_to_pptx(
        d, [{"promoOferta": "160", "unidadMoneda": "$"}], "a4", None, src)[0])
    assert "$160" in combo, f"el combo perdió su símbolo: {combo!r}"


def test_sin_promo_no_queda_un_simbolo_huerfano():
    # La razón por la que son DOS reglas y no una: agregarle el símbolo a un
    # cuadro que no lo tenía dejaría un "$" solo en todo producto sin promo,
    # que son la mayoría de las filas de cualquier listado.
    src = _pptx_promo("<<promoOferta>>")
    d = import_pptx(src)
    salida = _runs(render_template_to_pptx(
        d, [{"promoOferta": "", "unidadMoneda": "$"}], "a4", None, src)[0])
    assert "$" not in salida, f"quedó un símbolo suelto sin precio al lado: {salida!r}"
