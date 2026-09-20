"""Fija la regla de tamaño por cantidad de caracteres.

Reemplaza al achique automático, eliminado el 14/09/2026 (decisión de Ivan).
Aquel medía el texto con las métricas de la tipografía y decidía solo, en
cuatro pasos encadenados de los cuales dos multiplicaban sus pisos (0,55 sobre
0,55 = 30% del cuerpo de diseño) y los otros dos copiaban el peor resultado a
toda la hoja. Peor que eso: dependía de métricas, geometría y ancho de papel,
o sea de tres cosas que había que replicar idénticas en el navegador para que
el preview no mintiera -- y no se replicaban.

Ahora el cuerpo lo declara una persona con una condición sobre el LARGO del
texto. `len()` da lo mismo en Python que en TypeScript, así que la divergencia
preview/export deja de ser algo que hay que sincronizar.
"""
from app.services.cenefas.component_renderer import detectar_solapes, preparar_componentes
from app.services.cenefas.font_metrics import digito_mas_ancho
from app.services.cenefas.rules_engine import evaluate_rules, evaluate_font_size_rules

# ---------------------------------------------------------------------------
# Andamiaje
# ---------------------------------------------------------------------------

def _precio(pt=180.0, segs=None):
    c = {
        "id": "precio", "type": "text",
        "base_bounds": {"x": 1.0, "y": 10.0, "width": 12.0, "height": 5.0},
        "style": {"font_size": pt, "align": "center",
                  "font_family": "Impact", "font_bold": True},
    }
    if segs is None:
        c["variable"] = "precioOferta"
    else:
        c["segments"] = segs
    return c


def _regla(pt, largo, segmento=None, campo="precioOferta", id_="r"):
    r = {
        "id": id_, "target_component_id": "precio",
        "condition": {"field": campo, "operator": "length_greater_than", "value": largo},
        "action": {"type": "set_font_size", "value": pt},
    }
    if segmento is not None:
        r["target_segment_index"] = segmento
    return r


def _pt(reglas, valor, comp=None):
    preparados = preparar_componentes([comp or _precio()], reglas, {"precioOferta": valor})
    return preparados[0]["style"]["font_size"]


# ---------------------------------------------------------------------------
# El operador nuevo
# ---------------------------------------------------------------------------

def test_cuenta_caracteres_no_valor_numerico():
    # `greater_than` compara como número y daría verdadero para cualquier
    # precio de 4 cifras contra el umbral 3. Por eso hizo falta un operador
    # propio: acá 3 significa "3 caracteres", no "mayor que tres pesos".
    regla = [_regla(90, 3)]
    assert _pt(regla, "999") == 180      # 3 caracteres: no matchea
    assert _pt(regla, "1000") == 90      # 4 caracteres: matchea
    # Y el valor numérico no influye: "1" es más chico que 999 pero mide 1.
    assert _pt(regla, "1") == 180


def test_length_less_than():
    regla = [{
        "id": "r", "target_component_id": "precio",
        "condition": {"field": "precioOferta", "operator": "length_less_than", "value": 3},
        "action": {"type": "set_font_size", "value": 220},
    }]
    assert _pt(regla, "99") == 220
    assert _pt(regla, "999") == 180


def test_un_umbral_no_numerico_no_rompe_ni_matchea():
    regla = [{
        "id": "r", "target_component_id": "precio",
        "condition": {"field": "precioOferta", "operator": "length_greater_than", "value": "ocho"},
        "action": {"type": "set_font_size", "value": 90},
    }]
    assert _pt(regla, "12.345") == 180


def test_el_campo_tolera_el_nombre_en_mayusculas():
    # El formulario de reglas pasa a MAYÚSCULAS lo que se escribe en "Columna
    # del Excel"; el mismo alias que ya toleran show/hide.
    regla = [_regla(90, 3, campo="PRECIOOFERTA")]
    assert _pt(regla, "1.599") == 90


# ---------------------------------------------------------------------------
# Precedencia: gana la más chica, el orden no importa
# ---------------------------------------------------------------------------

def test_si_matchean_varias_gana_la_mas_chica():
    reglas = [_regla(90, 3, id_="a"), _regla(70, 4, id_="b")]
    assert _pt(reglas, "999") == 180        # 3 caracteres: ninguna
    assert _pt(reglas, "1599") == 90        # 4 caracteres: solo la de >3
    # OJO con el punto de miles: "1.599" son CINCO caracteres, no cuatro, así
    # que matchean las dos y gana la más chica. Es la trampa de contar
    # caracteres en vez de dígitos, y hay que tenerla presente al escribir el
    # umbral de una regla sobre un precio.
    assert _pt(reglas, "1.599") == 70
    assert _pt(reglas, "12.345") == 70      # matchean las dos -> la más chica


def test_el_orden_de_las_reglas_no_cambia_nada():
    # Misma garantía que ya tienen show/hide: el panel las lista en algún orden
    # pero el motor no las recorre aplicándolas una encima de la otra.
    a, b = _regla(90, 3, id_="a"), _regla(70, 4, id_="b")
    assert _pt([a, b], "12.345") == _pt([b, a], "12.345") == 70


def test_una_regla_sin_tamano_valido_se_ignora():
    for malo in (None, 0, -10, "grande"):
        reglas = [{
            "id": "r", "target_component_id": "precio",
            "condition": {"field": "precioOferta", "operator": "length_greater_than", "value": 3},
            "action": {"type": "set_font_size", "value": malo},
        }]
        assert _pt(reglas, "12.345") == 180, malo


# ---------------------------------------------------------------------------
# Convivencia con las reglas de mostrar/ocultar
# ---------------------------------------------------------------------------

def test_una_regla_de_tamano_no_oculta_el_cuadro():
    # Regresión: _resolver_visibilidad lee action.type y, si una regla de
    # tamaño entrara por ahí, el cuadro quedaría con "regla show que no
    # matchea" y se apagaría entero.
    reglas = [_regla(90, 3)]
    assert evaluate_rules(reglas, {"precioOferta": "999"}) == {}
    preparados = preparar_componentes([_precio()], reglas, {"precioOferta": "999"})
    assert preparados[0].get("visible", True) is True


def test_una_regla_de_ocultar_no_cambia_ningun_tamano():
    reglas = [{
        "id": "r", "target_component_id": "precio",
        "condition": {"field": "precioOferta", "operator": "is_not_empty"},
        "action": {"type": "hide"},
    }]
    assert evaluate_font_size_rules(reglas, {"precioOferta": "999"}) == {}


# ---------------------------------------------------------------------------
# Segmentos
# ---------------------------------------------------------------------------

_SEGS = [
    {"type": "static",   "value": "$",            "style": {"font_size": 102.0}},
    {"type": "variable", "value": "precioOferta",  "style": {"font_size": 160.0}},
]


def _segs_pt(reglas, valor):
    comp = _precio(pt=160.0, segs=[dict(s, style=dict(s["style"])) for s in _SEGS])
    preparados = preparar_componentes([comp], reglas, {"precioOferta": valor})
    return [s["style"]["font_size"] for s in preparados[0]["segments"]]


def test_una_regla_sobre_un_segmento_toca_solo_ese_segmento():
    # El caso de Rompe Precios Congelados: achicar el NÚMERO en un precio de 4
    # cifras sin tocar el "$" de al lado.
    reglas = [_regla(120, 3, segmento=1)]
    assert _segs_pt(reglas, "253") == [102.0, 160.0]
    assert _segs_pt(reglas, "1.599") == [102.0, 120]


def test_una_regla_sobre_el_cuadro_pone_ese_pt_en_todos_sus_segmentos():
    # En un cuadro multi-segmento cada segmento lleva SU font_size y ese pisa
    # al del componente al dibujar (_populate_text_frame), así que una regla
    # sobre el cuadro tiene que llegar a los segmentos o no hace nada visible.
    #
    # Y llega con el número TAL CUAL: los dos pedazos quedan en 80, el "$"
    # incluido. Decisión de Ivan (20/09/2026): "si yo pongo 90 en el cuadro
    # entero, todo tiene que medir 90 y punto"; si alguien quiere tocar un solo
    # pedazo, para eso está la regla por segmento (el test de arriba). Antes
    # esto escalaba por pt / font_size de la CAJA y daba [51, 80] -- ver el
    # porqué en apply_font_sizes y en test_reglas_tamano_paridad.py.
    reglas = [_regla(80, 3)]
    assert _segs_pt(reglas, "1.599") == [80.0, 80.0]


# ---------------------------------------------------------------------------
# El detector de solapes: avisa, no toca
# ---------------------------------------------------------------------------

def _dos_cuadros_encimados():
    izq = _precio()
    # La caja de la cocarda NO se solapa con la del precio (1..13): si dos
    # cajas DECLARADAS se pisan más del 5%, el detector las toma por
    # superpuestas a propósito por el diseño y no avisa -- que es lo correcto
    # para el "$" metido adentro del cuadro del precio, pero acá haría que el
    # test no probara nada. El choque que se busca es el del TEXTO: el precio
    # a 180 pt mide 18 cm y se sale de su caja por los dos lados.
    der = {
        "id": "cocarda", "type": "text", "static_value": "CLUB CARD",
        "base_bounds": {"x": 13.2, "y": 10.2, "width": 4.0, "height": 2.0},
        "style": {"font_size": 40.0, "align": "center",
                  "font_family": "Impact", "font_bold": True},
    }
    return [izq, der]


def test_detectar_solapes_avisa_del_choque():
    comps = _dos_cuadros_encimados()
    avisos = detectar_solapes([(c, {"precioOferta": "12.345"}) for c in comps])
    assert avisos, "un precio de 5 cifras a 180 pt pisa la cocarda de al lado"
    assert avisos[0]["component_id"] == "precio"     # el invasor es el de texto más grande
    assert avisos[0]["contra_id"] == "cocarda"


def test_detectar_solapes_no_toca_ningun_tamano():
    # La diferencia con el resolver viejo, que es todo el punto del cambio.
    comps = _dos_cuadros_encimados()
    detectar_solapes([(c, {"precioOferta": "12.345"}) for c in comps])
    assert [c["style"]["font_size"] for c in comps] == [180.0, 40.0]


def test_sin_choque_no_hay_aviso():
    comps = _dos_cuadros_encimados()
    assert detectar_solapes([(c, {"precioOferta": "99"}) for c in comps]) == []


# ---------------------------------------------------------------------------
# El dígito con el que se calibra el pt de la regla
# ---------------------------------------------------------------------------

def test_el_digito_mas_ancho_de_impact_no_es_el_ocho():
    # La intuición dice 8; en Impact son el 6 y el 9 (0,5400 em), y el 8 mide
    # 0,5350. Importa porque el relleno de capacidad se arma con este dígito y
    # es la vara para elegir el pt que se escribe en la regla.
    assert digito_mas_ancho("Impact") in ("6", "9")


def test_en_una_fuente_tabular_cualquier_digito_sirve():
    # Siete de las nueve fuentes de la tabla tienen todos los dígitos del mismo
    # ancho: ahí la cuenta por cantidad de caracteres es exacta.
    from app.services.cenefas.font_metrics import ancho_texto_em
    anchos = {ancho_texto_em(d, "Arial") for d in "0123456789"}
    assert len(anchos) == 1


def test_fuente_desconocida_no_rompe():
    assert digito_mas_ancho("Tipografia Inexistente") == "8"
    assert digito_mas_ancho(None) == "8"


# ---------------------------------------------------------------------------
# El autoajuste de PowerPoint
# ---------------------------------------------------------------------------
#
# El achique que MÁS caro salió: no el del motor, el de PowerPoint.
#
# Si el PPTX del diseñador traía la caja con "Reducir el texto al desbordarse",
# el importer lo anotaba (style["auto_fit"]) y el export se lo volvía a activar
# al archivo final. Mientras el motor achicaba solo no se notaba --para cuando
# PowerPoint abría el archivo el texto ya entraba-- pero al eliminar el achique
# automático el de PowerPoint quedó como único actor y encogió todo hasta
# meterlo en la caja. Como el canvas del preview no hace autofit, en pantalla se
# veía el cuerpo declarado y en el archivo salía otro: el mismo síntoma de
# siempre por una causa nueva. Reportado por Ivan con el export de Congelados.

import io
import zipfile

from pptx import Presentation
from pptx.util import Cm, Pt

from app.services.cenefas.component_renderer import render_template_to_pptx
from app.services.cenefas.pptx_importer import import_pptx


def _pptx_con_autoajuste():
    """Un A4 cuyo cuadro de precio tiene 'Reducir el texto al desbordarse'."""
    from pptx.oxml.ns import qn
    from lxml import etree
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(21.0), Cm(29.7)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    caja = sl.shapes.add_textbox(Cm(13.2), Cm(10.2), Cm(4.0), Cm(2.0))
    run = caja.text_frame.paragraphs[0].add_run()
    run.text = "$<<precioOferta>>"
    run.font.size = Pt(60)
    body_pr = caja.text_frame._txBody.find(qn("a:bodyPr"))
    for tag in (qn("a:noAutofit"), qn("a:spAutoFit")):
        for el in body_pr.findall(tag):
            body_pr.remove(el)
    # Como lo deja PowerPoint cuando ya calculó el encogimiento.
    af = etree.SubElement(body_pr, qn("a:normAutofit"))
    af.set("fontScale", "45000")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _bodypr_del_export(src, fila):
    d = import_pptx(src)
    pptx, _ = render_template_to_pptx(d, [fila], "a4", None, src)
    with zipfile.ZipFile(io.BytesIO(pptx)) as z:
        return "".join(
            z.read(n).decode("utf-8", "ignore")
            for n in z.namelist() if n.startswith("ppt/slides/slide"))


def test_el_importer_reconoce_el_autoajuste_del_diseno():
    # Si esto falla, el test de abajo no prueba lo que dice probar.
    d = import_pptx(_pptx_con_autoajuste())
    assert any(c.get("style", {}).get("auto_fit") for c in d["components"])


def test_el_export_apaga_el_autoajuste_de_powerpoint():
    xml = _bodypr_del_export(_pptx_con_autoajuste(), {"precioOferta": "12.345"})
    assert "noAutofit" in xml
    assert "normAutofit" not in xml, (
        "quedó el autoajuste de PowerPoint: el archivo se ve más chico que el "
        "preview aunque el motor ya no achique")
    # Y el fontScale que PowerPoint había dejado calculado se va con él.
    assert "fontScale" not in xml


def test_el_cuerpo_del_pptx_es_el_declarado_aunque_desborde():
    # "$12.345" a 60 pt no entra en una caja de 4 cm. Antes PowerPoint lo
    # encogía al abrirlo; ahora desborda y el sz del archivo es el de diseño.
    xml = _bodypr_del_export(_pptx_con_autoajuste(), {"precioOferta": "12.345"})
    assert 'sz="6000"' in xml


# ---------------------------------------------------------------------------
# La negrita puesta a mano le gana a la automática
# ---------------------------------------------------------------------------
#
# Reportado por Ivan (14/09/2026): "la negrita no funciona, no hay diferencia
# entre activa o no activa". No estaba rota: estaba anulada. Un segmento con
# transform "smart_bold" pone en negrita SOLO las palabras en MAYÚSCULAS y
# hasta ahora ignoraba `font_bold` por completo, así que tildar la casilla del
# panel no cambiaba nada ni en el preview ni en el archivo.
#
# La automática existe para decidir cuando la persona no decidió; en cuanto
# decide, manda lo suyo.

def _runs_de(comp, product):
    from pptx import Presentation
    from pptx.util import Cm
    from app.services.cenefas.component_renderer import _populate_text_frame, _texto_resuelto
    prs = Presentation()
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    caja = sl.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(3))
    # _render_slide resuelve las variables antes de llamar; acá se hace igual.
    comp = {**comp, "segments": [
        {**s, "_resolved": product.get(s.get("value"), "")} if s.get("type") == "variable" else s
        for s in (comp.get("segments") or [])
    ]}
    _populate_text_frame(caja.text_frame, comp, _texto_resuelto(comp, product))
    return [(r.text, r.font.bold) for p in caja.text_frame.paragraphs for r in p.runs if r.text]


_DESC = {
    "id": "desc", "type": "text",
    "style": {"font_size": 27.3, "font_family": "Impact"},
    "segments": [{"type": "variable", "value": "descripcion", "transform": "smart_bold"}],
}
_FILA = {"descripcion": "Aros de Cebolla LEDUC 450 gr"}


def test_sin_negrita_a_mano_solo_se_resaltan_las_mayusculas():
    runs = _runs_de(_DESC, _FILA)
    assert len(runs) > 1, "smart_bold tiene que partir el texto en varios runs"
    assert any(b for _, b in runs), "LEDUC tiene que salir en negrita"
    assert any(not b for _, b in runs), "el resto NO tiene que salir en negrita"


def test_con_negrita_a_mano_va_todo_en_negrita():
    comp = {**_DESC, "style": {**_DESC["style"], "font_bold": True}}
    runs = _runs_de(comp, _FILA)
    assert all(b for _, b in runs), f"quedó texto sin negrita: {runs}"
    assert "".join(t for t, _ in runs) == _FILA["descripcion"]


def test_la_casilla_cambia_el_resultado():
    # La regresión concreta que reportó Ivan: antes las dos daban lo mismo.
    sin = _runs_de(_DESC, _FILA)
    con = _runs_de({**_DESC, "style": {**_DESC["style"], "font_bold": True}}, _FILA)
    assert sin != con, "tildar negrita no cambia nada"
