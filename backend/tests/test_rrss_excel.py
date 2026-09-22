"""El Excel de correcciones para el diseñador (rrss/excel.py y rrss/correccion.py).

Los casos de las instrucciones son los de la campaña real del 17 al 20 de
setiembre de 2026, sobre la que Ivan aprobó el formato.
"""
import base64
import io

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.services.rrss import correccion as c
from app.services.rrss import excel, imagenes


# ---------------------------------------------------------------- instrucciones

@pytest.mark.parametrize("placa, mailing, esperado", [
    ("Carré de cerdo. Kg", "Carré de cerdo. kg", "Va en minúscula: «kg»"),
    ("Cerveza STELLA ARTOIS. Lata 710 ml.", "Cerveza STELLA ARTOIS. Lata 710 ml", "Sobra el punto: sacalo"),
    ("$171 unidad", "$171", "Sobra «unidad»: sacalo"),
    ("$340", "$340 unidad", "Falta «unidad»: agregalo"),
    ("Arvejas TIENDA INGLESA. 300g", "Arvejas TIENDA INGLESA. 300 g", "Falta el espacio: «300 g»"),
    ("Jarra eléctrica INHAUS negra 1.7L", "Jarra eléctrica INHAUS negra. 1.7 L", "Falta el punto y el espacio: agregalo"),
    ("Cerveza STELLA ARTOIS. Lata 710ml.", "Cerveza STELLA ARTOIS. Lata 710 ml", "Sacá el punto y agregá el espacio"),
    ("Vino BRISAS DEL ESTE. Tinto Tannat, Blanco Sauvignon Blanc 750 ml",
     "Vino BRISAS DEL ESTE Tinto Tannat, Blanco Sauvignon Blanc. 750 ml",
     "El punto va después de «Blanc», no de «ESTE»"),
])
def test_instrucciones_de_la_campana_real(placa, mailing, esperado):
    assert c.instruccion("diferente", placa, mailing) == esperado


def test_la_instruccion_no_baja_a_minuscula_lo_que_va_en_mayuscula():
    """str.capitalize() baja todo el resto: decía «colonia» va en mayúscula,
    justo lo contrario de lo que hay que hacer."""
    frase = c.instruccion("diferente", "Queso colonia IL CAPITANO. 100g", "Queso Colonia IL CAPITANO. 100 g")
    assert frase == "Falta el espacio: «100 g», y «Colonia» va en mayúscula"


def test_faltante_y_sobrante_se_dicen_igual_que_en_el_resto():
    assert c.instruccion("falta_en_placa", "", "unidad") == "Falta «unidad»: agregalo"
    assert c.instruccion("sobra_en_placa", "unidad", "") == "Sobra «unidad»: sacalo"


def test_si_cambia_mucho_se_da_el_texto_final():
    assert c.instruccion("diferente", "Pollo entero", "Muslo 1/4 GRANJA TRES ARROYOS. Kg") == \
        "Escribilo así: «Muslo 1/4 GRANJA TRES ARROYOS. Kg»"


# ---------------------------------------------------------------- resaltado por palabra

def _marcado(seg):
    return "".join(f"[{t}]" if d else t for t, d in seg)


def test_un_espacio_que_falta_se_ve_como_palabra_entera():
    p, m = c.diferenciar_palabras("Arvejas TIENDA INGLESA. 300g", "Arvejas TIENDA INGLESA. 300 g")
    assert _marcado(p) == "Arvejas TIENDA INGLESA. [300g]"
    assert _marcado(m) == "Arvejas TIENDA INGLESA. [300 g]"


def test_un_punto_que_se_corrio_marca_los_dos_lugares():
    p, m = c.diferenciar_palabras("DEL ESTE. Tinto Blanc 750", "DEL ESTE Tinto Blanc. 750")
    assert _marcado(p) == "DEL [ESTE.] Tinto [Blanc] 750"
    assert _marcado(m) == "DEL [ESTE] Tinto [Blanc.] 750"


def test_textos_iguales_no_marcan_nada():
    p, m = c.diferenciar_palabras("Bola de lomo. Kg", "Bola de lomo. Kg")
    assert not any(d for _, d in p + m)


# ---------------------------------------------------------------- nombres cortos

def test_etiquetas_sacan_el_prefijo_comun_sin_comerse_un_digito():
    nombres = [f"TI-Los Rompes del Finde-RRSS 17-20_{n}.jpg" for n in ("10", "10 copia", "11", "12")]
    assert list(c.etiquetas(nombres).values()) == ["10", "10 copia", "11", "12"]
    # todos empiezan con "1": el corte va en el separador, no se lleva el 1
    nombres = [f"campaña_{n}.png" for n in ("10", "11", "12")]
    assert list(c.etiquetas(nombres).values()) == ["10", "11", "12"]


def test_etiquetas_sin_prefijo_comun_o_una_sola_placa():
    assert c.etiquetas(["verano.jpg", "invierno.jpg"]) == {"verano.jpg": "verano", "invierno.jpg": "invierno"}
    assert c.etiquetas(["TI-RRSS_10 copia.jpg"]) == {"TI-RRSS_10 copia.jpg": "TI-RRSS_10 copia"}


# ---------------------------------------------------------------- el Excel

def _jpg(color, tam=(90, 90)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", tam, color).save(buf, format="JPEG")
    return buf.getvalue()


def _uri(color) -> str:
    return "data:image/jpeg;base64," + base64.b64encode(_jpg(color, (60, 40))).decode()


def _fila(campo, estado, placa, mailing):
    return {"campo": campo, "estado": estado, "placa": placa, "mailing": mailing, "severidad": "error"}


def _placa(nombre, formato, estado, indice, filas, orden, vista=True, lectura=True):
    return {
        "nombre_archivo": f"TI-RRSS 17-20_{nombre}.jpg", "formato": formato, "estado": estado, "orden": orden,
        "match": {"indice": indice, "puntaje": 100} if indice is not None else None,
        "filas": filas, "error": None,
        "lectura": {"producto": {"descripcion": "algo"}} if lectura else None,
        "vista": _jpg((30, 60, 150), (90, 90) if formato == "1:1" else (90, 160)) if vista else None,
    }


KG = _fila("descripcion", "diferente", "Carré de cerdo. Kg", "Carré de cerdo. kg")
PUNTO = _fila("descripcion", "diferente", "Lata 710 ml.", "Lata 710 ml")
PUNTO_Y_ESPACIO = _fila("descripcion", "diferente", "Lata 710ml.", "Lata 710 ml")

VALIDACION = {
    "mailing": {
        "fecha": "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE",
        "productos": [
            {"descripcion": "Carré de cerdo. kg", "recorte": _uri((200, 40, 40))},
            {"descripcion": "Cerveza STELLA ARTOIS. Lata 710 ml", "recorte": _uri((200, 200, 40))},
            {"descripcion": "Bola de lomo. Kg", "recorte": None},
        ],
    },
    "imagenes": [
        _placa("13", "1:1", "diferencias", 0, [KG], 1),
        _placa("14", "4:5", "diferencias", 0, [KG], 2),
        _placa("15", "9:16", "diferencias", 0, [KG], 3),
        _placa("16", "1:1", "diferencias", 1, [PUNTO], 4),
        _placa("17", "4:5", "diferencias", 1, [PUNTO_Y_ESPACIO], 5),
        _placa("18", "9:16", "diferencias", 1, [PUNTO], 6),
        _placa("10", "1:1", "ok", 2, [], 7),
    ],
}


def _abrir(validacion):
    return load_workbook(io.BytesIO(excel.construir(validacion)), rich_text=True)


def _bloques(ws):
    """Las filas de la columna A que arrancan un bloque (celdas combinadas o sueltas)."""
    return [f for f in range(4, ws.max_row + 1) if ws[f"A{f}"].value]


def test_un_bloque_por_arreglo_no_por_placa():
    ws = _abrir(VALIDACION)["Recomendada"]
    # 6 placas con error -> 3 arreglos: el carré (3 placas), la Stella 16+18 y la Stella 17 sola
    assert len(_bloques(ws)) == 3
    assert "3 arreglos" not in str(ws["A2"].value)  # el texto dice "cosas para corregir"
    assert "3 cosas para corregir, en 6 de las 7 placas" in str(ws["A2"].value)


def test_las_placas_que_estan_bien_no_salen():
    ws = _abrir(VALIDACION)["Recomendada"]
    textos = " ".join(str(ws[f"A{f}"].value) for f in _bloques(ws))
    assert "Bola de lomo" not in textos


def test_la_adaptacion_distinta_de_sus_hermanas_lo_avisa():
    ws = _abrir(VALIDACION)["Recomendada"]
    avisos = [str(ws[f"A{f}"].value) for f in _bloques(ws)]
    assert sum("distinta de sus otras adaptaciones" in a for a in avisos) == 1


def test_lo_que_cambia_va_resaltado_y_la_instruccion_al_costado():
    ws = _abrir(VALIDACION)["Recomendada"]
    fila = _bloques(ws)[0]
    dice, tiene = ws[f"D{fila}"].value, ws[f"E{fila}"].value
    resaltado = lambda rico: [t.text for t in rico if not isinstance(t, str) and t.font.b]  # noqa: E731
    assert resaltado(dice) == ["Kg"] and resaltado(tiene) == ["kg"]
    assert ws[f"F{fila}"].value == "Va en minúscula: «kg»"


def test_lleva_la_hoja_con_todas_las_placas():
    wb = _abrir(VALIDACION)
    # Las tres hojas del molde, en su orden; «La planilla» solo con planilla.
    assert wb.sheetnames == ["Recomendada", "Como la pediste", "Todas las placas"]
    ws = wb["Todas las placas"]
    assert ws.max_row - 2 == len(VALIDACION["imagenes"])  # una fila por placa, también las que están bien
    # en orden natural de etiqueta (la 10 está bien y va primera igual), no en el orden de subida
    assert [ws[f"A{f}"].value for f in range(3, 6)] == ["10", "13", "14"]


def test_las_imagenes_van_adentro():
    ws = _abrir(VALIDACION)["Recomendada"]
    # por bloque: la tira de placas + el recorte del mailing
    assert len(ws._images) == 6


def test_todas_bien_y_nada_mas_que_decir_no_hay_excel():
    """Con un mailing y todas las placas bien no hay nada que mandar. Ojo: con
    una planilla puede haber (ver los dos tests de más abajo)."""
    todas_bien = {**VALIDACION, "imagenes": [VALIDACION["imagenes"][-1]]}
    with pytest.raises(excel.NadaParaCorregir):
        excel.construir(todas_bien)


# ---------------------------------------------------------------- (10) una palabra que se pueda sostener

def _con_planilla(imagenes, campos=("descripcion", "oferta_precio"), avisos=(), **info):
    """Una validación contra una planilla de 3 filas.

    `info` pisa lo que haga falta del bloque "planilla" (los contadores de la
    lectura, por ejemplo)."""
    productos = [
        {"descripcion": "Carré de cerdo. kg", "recorte": None, "fila": {"numero": 4, "valores": {"descripcion": "Carré de cerdo. kg"}}},
        {"descripcion": "Cerveza STELLA ARTOIS. Lata 710 ml", "recorte": None, "fila": {"numero": 5, "valores": {"descripcion": "Cerveza STELLA ARTOIS. Lata 710 ml"}}},
        {"descripcion": "Bola de lomo. Kg", "recorte": None, "fila": {"numero": 6, "valores": {"descripcion": "Bola de lomo. Kg"}}},
    ]
    return {
        "mailing": {
            "origen": "planilla", "fecha": "", "campos": list(campos), "productos": productos,
            "planilla": {
                "archivo": "LISTADO.xlsx", "hoja": "Campaña", "titulos": {"descripcion": "DESCRIPCION"},
                "fila_encabezado": 1, "filas_leidas": 3, "filas_ignoradas": 0,
                "filas_juntadas": 0,
                "descripcion_de_gestion": False, "avisos": list(avisos),
                **info,
            },
        },
        "imagenes": imagenes,
    }


def test_no_dice_esta_bien_cuando_no_miro_ni_la_mitad():
    """Una placa con mecánica, encabezado y pie inventados contra una planilla
    que no trae esas columnas daba "Está bien", en verde, con 2 de 7 campos
    revisados. El encabezado lo aclaraba; la hoja "Todas las placas" no."""
    bien = _placa("10", "1:1", "ok", 0, [], 1)
    mala = _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)
    ws = _abrir(_con_planilla([bien, mala]))["Todas las placas"]
    estados = [ws[f"D{f}"].value for f in range(3, 5)]
    assert "Está bien" not in estados
    assert estados[0] == "Sin diferencias"
    # y en la misma fila dice CUÁNTO se miró, no solo en el encabezado
    assert ws["E3"].value == "Coincide en lo que dicta la planilla (2 de 7 campos)"


def test_contra_un_mailing_no_se_habla_de_campos_que_quedaron_afuera():
    """Un mailing dicta los siete campos: ahí no hay nada que aclarar, y la placa
    sana dice «Está bien» con un guion en QUÉ TIENE, como en el molde."""
    ws = _abrir(VALIDACION)["Todas las placas"]
    ok = next(f for f in range(3, ws.max_row + 1) if ws[f"D{f}"].value == "Está bien")
    assert ws[f"E{ok}"].value == "—"


# ---------------------------------------------------------------- (11) el Excel que faltaba

def test_hay_excel_aunque_todas_las_placas_esten_bien_si_faltan_placas():
    """Una campaña de 3 productos con UNA sola placa, y esa placa perfecta: no
    hay ninguna placa mal, así que el Excel cortaba antes de armarse y no había
    NADA para mandarle al diseñador. Y es justo la que más lo necesita: la hoja
    "La planilla" muestra las dos filas que ninguna placa reclamó."""
    sola = _placa("10", "1:1", "ok", 0, [], 1)
    wb = _abrir(_con_planilla([sola]))
    assert "La planilla" in wb.sheetnames
    ws = wb["La planilla"]
    estados = [ws.cell(row=r, column=ws.max_column).value for r in range(4, ws.max_row + 1)]
    assert estados.count("Sin placa") == 2
    # y la hoja de correcciones lo dice con todas las letras, no con una tabla vacía
    correcciones = wb["Recomendada"]
    assert "La placa está bien" in str(correcciones["A2"].value)
    assert "La planilla" in str(correcciones["A4"].value)


def test_sin_placas_que_falten_y_sin_avisos_no_hay_excel_tampoco_con_planilla():
    """El corte no se fue: si están todas las filas cubiertas y todas las placas
    bien, no hay nada que decir."""
    placas = [_placa(str(10 + n), "1:1", "ok", n, [], n) for n in range(3)]
    with pytest.raises(excel.NadaParaCorregir):
        excel.construir(_con_planilla(placas))


# ---------------------------------------------------------------- (9) los avisos llegan al Excel

def test_los_avisos_de_la_planilla_salen_en_el_excel():
    """Se armaban y morían en un campo que nadie leía. Son los que delatan que
    el motor entendió mal el archivo, así que van en las dos hojas."""
    aviso = "No lo conté como producto, parece el pie del reporte: fila 7: «TOTAL: 5 artículos»"
    mala = _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)
    wb = _abrir(_con_planilla([mala], avisos=[aviso]))
    assert aviso in str(wb["Recomendada"]["A2"].value)
    planilla = wb["La planilla"]
    textos = [planilla.cell(row=r, column=1).value for r in range(1, 6)]
    assert aviso in textos
    # y los contadores de la lectura, que estaban tipados y no se mostraban
    assert "leídas a partir de la fila 2" in str(planilla["A2"].value)


def test_una_placa_que_no_se_pudo_leer_no_rompe_el_excel():
    rota = {**_placa("20", "", "error", None, [], 1, vista=False, lectura=False), "error": "No pude abrir el archivo"}
    ws = _abrir({**VALIDACION, "imagenes": [rota]})["Recomendada"]
    assert ws["C4"].value == "No se pudo leer"


def test_una_placa_sin_producto_en_el_mailing_va_sola():
    a = _placa("21", "1:1", "sin_match", None, [], 1)
    b = _placa("22", "1:1", "sin_match", None, [], 2)
    ws = _abrir({**VALIDACION, "imagenes": [a, b]})["Recomendada"]
    assert len(_bloques(ws)) == 2  # no se juntan: no hay nada que diga que son el mismo producto
    assert ws["B4"].value == "No está en el mailing"


def test_anda_en_un_servidor_sin_fuentes(monkeypatch):
    """Render es Linux sin fuentes instaladas: el pie de las miniaturas tiene que
    salir igual, con la fuente que trae Pillow.

    Las rutas viven en imagenes.py desde que también las usa la tira de la
    planilla (rrss/planilla.py): eran las mismas dos listas en dos módulos."""
    monkeypatch.setattr(imagenes, "_FUENTES", ("/no/existe.ttf",))
    monkeypatch.setattr(imagenes, "_FUENTES_NEGRITA", ("/no/existe-bold.ttf",))
    ws = _abrir(VALIDACION)["Recomendada"]
    assert len(ws._images) == 6


# ---------------------------------------------------------------- los contadores de la lectura

def _encabezado_de_la_planilla(**info) -> str:
    mala = _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)
    return str(_abrir(_con_planilla([mala], **info))["La planilla"]["A2"].value)


def test_un_solo_pie_no_dice_1_filas_quedaron_afuera():
    """Se leía "(1 filas quedaron afuera)", y se ve en cuanto una planilla trae
    UN solo pie de reporte, que desde que el pie se reconoce es el caso normal.
    Es una línea que el diseñador lee en cada Excel."""
    texto = _encabezado_de_la_planilla(filas_ignoradas=1)
    assert "(1 fila quedó afuera)" in texto
    assert "1 filas" not in texto


def test_con_mas_de_una_fila_afuera_sigue_en_plural():
    assert "(2 filas quedaron afuera)" in _encabezado_de_la_planilla(filas_ignoradas=2)


def test_las_filas_juntadas_tambien_se_imprimen():
    """El tercer contador de la lectura: viajaba en la API (filas_juntadas) y no
    se imprimía ni en la pantalla ni acá. Es el que delata que el motor entendió
    el archivo de otra manera que la persona: 3 productos salidos de 90 filas."""
    assert "juntando 90 filas del archivo" in _encabezado_de_la_planilla(filas_juntadas=90)


def test_sin_formato_largo_no_se_habla_de_juntar():
    """El caso normal no junta nada y no tiene nada que decir."""
    texto = _encabezado_de_la_planilla()
    assert "juntando" not in texto
    assert "quedaron afuera" not in texto


# ---------------------------------------------------------------- (22/09) el Excel es el del molde

def _fila_falta(campo, mailing):
    return {"campo": campo, "estado": "falta_en_placa", "placa": "", "mailing": mailing, "severidad": "error"}


LECTURA_COMPLETA = {
    "producto": {"descripcion": "Cerveza STELLA ARTOIS. Lata 710 ml.", "precio_anterior": "$171", "mecanica": "",
                 "oferta_encabezado": "Oferta", "oferta_precio": "$125", "oferta_pie": "", "es_alcohol": True},
    "fecha": "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE",
    "legal_bases": "Bases y condiciones en tiendainglesa.com.uy",
    "legal_alcohol": "Beber con moderación. Prohibida la venta a menores de 18 años.",
}


def _con_lectura(placa, lectura=LECTURA_COMPLETA):
    return {**placa, "lectura": lectura}


def test_las_hojas_son_las_del_molde_y_la_planilla_va_cuarta():
    """'Recomendada', 'Como la pediste', 'Todas las placas'; con planilla se
    agrega 'La planilla' al final. 'Como la pediste' se había perdido entera
    (la versión que la generaba nunca se commiteó)."""
    assert _abrir(VALIDACION).sheetnames == ["Recomendada", "Como la pediste", "Todas las placas"]
    mala = _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)
    assert _abrir(_con_planilla([mala])).sheetnames == [
        "Recomendada", "Como la pediste", "Todas las placas", "La planilla",
    ]


def test_las_placas_van_en_orden_natural_y_no_en_el_de_subida():
    """El orden de subida es el alfabético del sistema: '…_10 copia.jpg' antes
    que '…_10.jpg' (el espacio ordena antes que el punto). El molde ordena como
    los nombra el equipo: 10, 10 copia, 11, 11 copia…"""
    copia = _placa("10 copia", "1:1", "diferencias", 1, [PUNTO], 0)   # subida primera
    diez = _placa("10", "1:1", "diferencias", 0, [KG], 1)             # subida segunda
    once = _placa("9", "1:1", "ok", 2, [], 2)
    wb = _abrir({**VALIDACION, "imagenes": [copia, diez, once]})
    todas = wb["Todas las placas"]
    assert [todas[f"A{f}"].value for f in range(3, 6)] == ["9", "10", "10 copia"]
    ancha = wb["Como la pediste"]
    assert ancha["A4"].value == "Placa 10  ·  1:1" and ancha["A5"].value == "Placa 10 copia  ·  1:1"
    # y el primer bloque de la recomendada es el de la placa 10 (Carré), no el de la copia
    assert "Carré de cerdo" in str(wb["Recomendada"]["A4"].value)


def test_como_la_pediste_una_fila_por_placa_con_lo_que_dice_la_placa():
    """La tabla ancha del molde: una fila por placa con algo para corregir, una
    columna por cosa que CatTi lee; la celda con error va en rojo con marco y la
    última columna junta las correcciones."""
    stella = _con_lectura(_placa("16", "1:1", "diferencias", 1, [PUNTO, _fila_falta("oferta_pie", "unidad")], 1))
    bien = _placa("10", "1:1", "ok", 2, [], 2)
    ws = _abrir({**VALIDACION, "imagenes": [stella, bien]})["Como la pediste"]
    assert [ws[f"{col}3"].value for col in "ABCDEFGHIJ"] == [
        "LA PLACA", "DESCRIPCIÓN", "PRECIO ANTERIOR", "MECÁNICA", "ARRIBA DEL PRECIO", "PRECIO OFERTA",
        "ABAJO DEL PRECIO", "FECHA", "LEGALES", "CORRECCIÓN",
    ]
    assert ws.max_row == 4  # la placa que está bien no sale
    assert ws["A4"].value == "Placa 16  ·  1:1"
    assert ws.row_dimensions[4].height == 150
    # la descripción con error: rica, con lo que cambia en negrita, fondo rojo y marco rojo
    rico = ws["B4"].value
    assert [t.text for t in rico if not isinstance(t, str) and t.font.b] == ["ml."]
    assert ws["B4"].fill.fgColor.rgb.endswith("FDECEA") and ws["B4"].border.left.style == "medium"
    # lo que no tiene problema va plano y gris, y lo vacío con un guion
    assert ws["C4"].value == "$171" and ws["C4"].font.color.rgb.endswith("6B7686")
    assert ws["D4"].value == "—" and ws["D4"].border.left.style == "thin"
    # el pie que falta: "(no está)" en cursiva, también marcado
    assert [t.text for t in ws["G4"].value if not isinstance(t, str)] == ["(no está)"]
    assert ws["G4"].border.left.style == "medium"
    assert ws["I4"].value == (
        "Bases y condiciones en tiendainglesa.com.uy  /  Beber con moderación. Prohibida la venta a menores de 18 años."
    )
    assert ws["J4"].value == "• Descripción: Sobra el punto: sacalo\n• Texto abajo del precio: Falta «unidad»: agregalo"
    assert ws["J4"].font.color.rgb.endswith("9A6400") and ws["J4"].fill.fgColor.rgb.endswith("FFF4D6")
    assert len(ws._images) == 1 and ws.freeze_panes == "B4"


def test_como_la_pediste_dice_sin_revisar_en_vez_de_un_guion_con_planilla():
    """Con una planilla sin columna MECÁNICA ni VIGENCIA, el guion decía dos
    cosas distintas: "no lo tiene y no hace falta" y "esto no lo miró nadie".
    Ahora la segunda se dice con palabras."""
    mala = _con_lectura(_placa("11", "1:1", "diferencias", 1, [PUNTO], 2))
    ws = _abrir(_con_planilla([mala]))["Como la pediste"]
    assert ws["D4"].value == "sin revisar: la planilla no lo dice"
    assert ws["D4"].fill.fgColor.rgb.endswith("F2F4F7") and ws["D4"].font.i
    # la fecha tampoco: la planilla no trae VIGENCIA y la config no la escribió
    assert ws["H4"].value == "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE\n(sin revisar: la planilla no lo dice)"
    # y la leyenda de alcohol de una placa de alcohol, ídem
    assert "leyenda de alcohol sin revisar" in ws["I4"].value
    # con la fecha escrita en la carga sí se revisó
    v = _con_planilla([mala])
    v["config"] = {"fecha": "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE"}
    assert _abrir(v)["Como la pediste"]["H4"].value == "DEL JUEVES 17 AL DOMINGO 20 DE SETIEMBRE"
    # contra un mailing el guion sigue siendo el del molde
    ws = _abrir({**VALIDACION, "imagenes": [mala]})["Como la pediste"]
    assert ws["D4"].value == "—"


def test_una_planilla_que_dicta_todo_puede_decir_esta_bien():
    """Contra una planilla con las seis columnas escritas se miró lo mismo que
    contra un mailing: ahí sí vale 'Está bien' y el guion del molde."""
    bien = _placa("10", "1:1", "ok", 0, [], 1)
    mala = _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)
    todo = ("descripcion", "precio_anterior", "oferta_precio", "mecanica", "oferta_encabezado", "oferta_pie")
    ws = _abrir(_con_planilla([bien, mala], campos=todo))["Todas las placas"]
    assert ws["D3"].value == "Está bien" and ws["E3"].value == "—"


def test_el_banner_habla_en_singular_cuando_es_uno():
    """"1 cosas para corregir" y "Las otras 1 están bien" se leían en cuanto
    había un solo arreglo o una sola placa sana."""
    mala = _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)
    bien = _placa("10", "1:1", "ok", 0, [], 1)
    wb = _abrir({**VALIDACION, "imagenes": [mala, bien]})
    banner = str(wb["Recomendada"]["A2"].value)
    assert "1 cosa para corregir, en 1 de las 2 placas.  La otra está bien: no la toques." in banner
    assert "1 cosas" not in banner and "otras 1" not in banner
    assert "Hay que corregir 1 de las 2 placas.  La otra está bien: no la toques." in str(wb["Como la pediste"]["A2"].value)
    sola = _abrir({**VALIDACION, "imagenes": [mala]})
    assert "1 cosa para corregir, en la única placa." in str(sola["Recomendada"]["A2"].value)
    assert "Hay que corregir la única placa." in str(sola["Como la pediste"]["A2"].value)
    assert "otras" not in str(sola["Recomendada"]["A2"].value)


def test_nada_dice_de_el_mailing():
    """"Hay más de una fila de el mailing": el artículo no se contraía al
    interpolar el nombre de la fuente."""
    ambigua = {**_placa("21", "1:1", "sin_match", None, [], 1), "sin_pareja": "ambiguo",
               "candidatos": [{"descripcion": "Carré de cerdo. kg", "puntaje": 110.0}]}
    duda = {"campo": "emparejamiento", "estado": "revisar", "placa": "x", "mailing": "y", "severidad": "aviso",
            "nota": "Hay otra fila que se le parece casi igual (100.0 contra 99.0): confirmá que sea esta."}
    dudosa = _placa("22", "1:1", "avisos", 0, [duda], 2)
    wb = _abrir({**VALIDACION, "imagenes": [ambigua, dudosa]})
    textos = " | ".join(
        str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)
    )
    assert "de el " not in textos and "de El " not in textos
    assert "Hay más de una fila del mailing" in textos
    assert "Con qué fila del mailing la emparejé" in textos
    assert c.nombre_campo("emparejamiento", "planilla") == "Con qué fila de la planilla la emparejé"
    # y en la hoja ancha la duda lleva su nota, no solo "miralo a mano"
    ancha = wb["Como la pediste"]
    j = [str(ancha[f"J{r}"].value) for r in range(4, 6)]
    assert any("Miralo a mano contra el mailing: Hay otra fila" in t for t in j)
    assert any("No encontré este producto" in t or "Hay más de una fila del mailing" in t for t in j)


def test_que_hacer_no_manda_a_agregar_debe_estar():
    """Para un logo que falta, `instruccion` recibía placa="no está" y
    mailing="debe estar" y decía «Falta «debe estar»: agregalo»."""
    assert c.que_hacer({"campo": "logo_campana", "estado": "falta_en_placa", "placa": "no está", "mailing": "debe estar"}) \
        == "Falta: agregalo"
    assert c.que_hacer({"campo": "precio_anterior_tachado", "estado": "diferente", "placa": "sin tachar", "mailing": "tachado"}) \
        == "Tiene que ir tachado"
    assert c.que_hacer({"campo": "descripcion", "estado": "revisar", "placa": "a", "mailing": "b", "nota": "CatTi leyó distinto"}) \
        == "Miralo a mano contra el mailing: CatTi leyó distinto"
    assert c.que_hacer(KG) == "Va en minúscula: «kg»"


def test_la_hoja_de_la_planilla_tambien_habla_en_singular():
    """"1 filas con producto" y "las 1 que ninguna placa reclamó"."""
    dos = [_placa("10", "1:1", "ok", 0, [], 1), _placa("11", "1:1", "diferencias", 1, [PUNTO], 2)]
    texto = str(_abrir(_con_planilla(dos))["La planilla"]["A2"].value)
    assert "La que quedó SIN COLOR es la única fila que ninguna placa reclamó" in texto
    tres = dos + [_placa("12", "1:1", "ok", 2, [], 3)]
    texto = str(_abrir(_con_planilla(tres, avisos=["un aviso"]))["La planilla"]["A2"].value)
    assert "Todas las filas tienen su placa" in texto
