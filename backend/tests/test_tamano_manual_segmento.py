"""Fija la regla "lo que ponés a mano manda donde lo ponés" para el tamaño de letra.

Bug real (11/09/2026): en un cuadro "<<unidadMoneda>><<precioOferta>>", si la
caja tenía un tamaño puesto a mano (ej. 150), el tamaño que la persona le ponía
a UN segmento (el "$" a 60) se ignoraba y el "$" salía a 150 igual que el
número. El motor no tenía forma de distinguir ese tamaño del que el segmento
trae copiado del PPTX al importar.

Criterio de Ivan:
  - El tamaño puesto a mano en un segmento manda en ese segmento, aunque la
    caja tenga tamaño manual. Los segmentos sin tamaño propio siguen a la caja.
  - El tamaño manual es un techo: si el contenido no entra en su propia caja,
    baja (todo en la misma proporción), nunca sube.

Se prueba de punta a punta: PPTX armado en memoria -> importer -> render.
"""
import io

from pptx import Presentation
from pptx.util import Cm, Pt

from app.services.cenefas.component_renderer import (
    _INSET_CM,
    _ancho_medido_cm,
    render_template_to_pptx,
)
from app.services.cenefas.pptx_importer import import_pptx


def _pptx_precio(x=1.0, y=3.0, w=19.0, h=9.0):
    prs = Presentation()
    prs.slide_width = Cm(21.0)
    prs.slide_height = Cm(29.7)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "<<unidadMoneda>><<precioOferta>>"
    run.font.size = Pt(180)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _preparar(src, pt_simbolo, pt_numero, caja_manual=None, simbolo_a_mano=False):
    """Plantilla importada con los tamaños que dejaría el editor."""
    d = import_pptx(src)
    comp = next(c for c in d["components"] if "precioOferta" in str(c))
    for seg in comp["segments"]:
        es_simbolo = seg.get("value") == "unidadMoneda"
        seg["style"] = {**(seg.get("style") or {}),
                        "font_size": pt_simbolo if es_simbolo else pt_numero}
        if es_simbolo and simbolo_a_mano:
            seg["_manual_font_override"] = True
    if caja_manual is not None:
        comp["style"] = {**comp["style"], "font_size": caja_manual}
        comp["_manual_font_override"] = True
    return d, comp


def _tamanos(src, d, precio):
    pptx, _ = render_template_to_pptx(
        d, [{"unidadMoneda": "$", "precioOferta": precio}], "a4", None, src)
    out = {}
    for sh in Presentation(io.BytesIO(pptx)).slides[0].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() in ("$", precio) and r.font.size:
                    out[r.text.strip()] = r.font.size.pt
    return out["$"], out[precio]


# --- A-D: el precio entra de sobra, nadie se achica --------------------------

def test_a_sin_nada_a_mano_cada_segmento_con_su_tamano():
    src = _pptx_precio()
    d, _ = _preparar(src, 120, 180)
    assert _tamanos(src, d, "99") == (120, 180)


def test_b_segmento_a_mano_sin_tamano_en_la_caja():
    src = _pptx_precio()
    d, _ = _preparar(src, 60, 180, simbolo_a_mano=True)
    assert _tamanos(src, d, "99") == (60, 180)


def test_c_caja_a_mano_pisa_los_tamanos_que_vienen_del_pptx():
    # Los segmentos SIN marca traen el tamaño copiado del PPTX: siguen a la caja.
    src = _pptx_precio()
    d, _ = _preparar(src, 120, 180, caja_manual=150)
    assert _tamanos(src, d, "99") == (150, 150)


def test_d_segmento_a_mano_manda_sobre_la_caja_a_mano():
    # El bug: el "$" a 60 puesto a mano salía a 150, como la caja.
    src = _pptx_precio()
    d, _ = _preparar(src, 60, 180, caja_manual=150, simbolo_a_mano=True)
    assert _tamanos(src, d, "99") == (60, 150)


def test_d_el_numero_no_se_achica_por_medir_el_simbolo_al_tamano_de_la_caja():
    # Caja justa: "$1.234" entra con el "$" a 60 pero NO entraría con el "$" a
    # 150. Medir todo al tamaño de la caja achicaba un número que sí entra.
    d0, comp0 = _preparar(_pptx_precio(), 60, 150)
    familia = comp0["style"].get("font_family")
    bold = bool(comp0["style"].get("font_bold"))
    numero = _ancho_medido_cm("1.234", 150, familia, bold)
    con_chico = _ancho_medido_cm("$", 60, familia, bold) + numero
    con_grande = _ancho_medido_cm("$", 150, familia, bold) + numero
    ancho = _INSET_CM + (con_chico + con_grande) / 2
    assert ancho < 20, "la caja tiene que entrar en la hoja"

    src = _pptx_precio(w=ancho)
    d, _ = _preparar(src, 60, 150, caja_manual=150, simbolo_a_mano=True)
    assert _tamanos(src, d, "1.234") == (60, 150)


# --- E: el tamaño manual es un techo -------------------------------------------

def test_e_caja_a_mano_baja_si_el_numero_no_entra():
    src = _pptx_precio()
    d, _ = _preparar(src, 180, 180, caja_manual=180)
    simbolo, numero = _tamanos(src, d, "12.345")
    assert numero < 180
    assert simbolo == numero


def test_e_con_segmento_a_mano_bajan_todos_en_la_misma_proporcion():
    src = _pptx_precio()
    d, _ = _preparar(src, 60, 180, caja_manual=180, simbolo_a_mano=True)
    simbolo, numero = _tamanos(src, d, "12.345")
    assert numero < 180 and simbolo < 60
    assert abs(simbolo / numero - 60 / 180) < 0.01
