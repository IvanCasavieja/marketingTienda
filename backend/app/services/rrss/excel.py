"""El Excel de correcciones para el diseñador.

Tres hojas, en este orden, con el formato que Ivan aprobó el 2026-09-21 sobre
la campaña del 17 al 20 de setiembre (el archivo «Correcciones RRSS -
ejemplo.xlsx» es el molde, y este módulo lo reproduce celda por celda):

- «Recomendada»: SOLO las placas que tienen algo para corregir, un bloque por
  ARREGLO y no por placa. Las adaptaciones de un mismo producto (1:1, 4:5,
  9:16) casi siempre traen el mismo error: van juntas, así se corrige una vez
  y se exportan las tres. Si una tiene algo distinto a sus hermanas, va en su
  propio bloque y lo avisa. Placa y fuente lado a lado, y al costado solo lo
  que está mal: lo que dice (en rojo), lo que tiene que decir (en verde) y qué
  hacer, en una frase.
- «Como la pediste»: la tabla ancha que Ivan pidió primero, UNA FILA POR
  PLACA con algo para corregir y una columna por cada cosa que CatTi lee de la
  placa (descripción, precios, mecánica, textos del círculo, fecha, legales).
  Cada celda dice lo que la PLACA dice; la que está mal va en rojo con marco,
  y la última columna junta las correcciones. Esta hoja se había perdido: la
  versión que la generaba nunca se commiteó (22/09/2026).
- «Todas las placas»: una fila por placa, también las que están bien, con su
  estado.
- «La planilla» va CUARTA y solo cuando la fuente es una planilla: es el
  archivo entero, con las filas que ninguna placa reclamó sin pintar.

Con una planilla la columna de la fuente se llama LA PLANILLA y en vez del
recorte de la foto va la tira dibujada de la fila (ver planilla.tira).

Es CPU puro y sincrónico (decodifica y arma imágenes): el que lo llame desde
una corrutina tiene que mandarlo a un hilo (ver hilos.py).
"""
import base64
import io
import re

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw

from app.services.rrss import comparador
from app.services.rrss import correccion as c
from app.services.rrss import planilla as pl
from app.services.rrss.imagenes import fuente as _fuente

_EMU = 9525  # EMU por píxel

# ── Paleta ────────────────────────────────────────────────────────────────
AZUL, AZUL_2 = "1B3A6B", "2C5490"
GRIS_OSC, GRIS, GRIS_TXT, GRIS_SUAVE = "44546A", "F2F4F7", "32404F", "6B7686"
ROJO, ROJO_F = "C0392B", "FDECEA"
VERDE, VERDE_F = "1E8449", "E8F6EC"
AMBAR, AMBAR_F = "9A6400", "FFF4D6"
VACIO = "8792A2"  # el gris de "(no está)" / "(no va nada)"

_linea = Side(style="thin", color="D7DCE3")
_grueso = Side(style="medium", color="9AA4B2")
_marco_rojo = Side(style="medium", color=ROJO)
_caja = Border(left=_linea, right=_linea, top=_linea, bottom=_linea)
_caja_roja = Border(left=_marco_rojo, right=_marco_rojo, top=_marco_rojo, bottom=_marco_rojo)

_ORDEN_FORMATO = {"1:1": 0, "4:5": 1, "9:16": 2}


class NadaParaCorregir(ValueError):
    """Todas las placas están bien: no hay Excel que armar."""


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _px_col(ancho_chars: float) -> int:
    return round(ancho_chars * 7) + 5


def _pil(valor) -> Image.Image | None:
    """Una imagen desde bytes o desde un data URI."""
    if not valor:
        return None
    if isinstance(valor, str):
        valor = base64.b64decode(valor.split(",", 1)[1])
    return Image.open(io.BytesIO(valor)).convert("RGB")


def _jpeg(im: Image.Image) -> bytes:
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=86)
    return buf.getvalue()


# ── El orden de las placas ────────────────────────────────────────────────

def _clave_natural(etiqueta: str) -> tuple:
    """«10» antes que «10 copia», y las dos antes que «11».

    El `orden` de subida es el alfabético del sistema de archivos, donde
    "…_10 copia.jpg" va ANTES que "…_10.jpg" (el espacio ordena antes que el
    punto): la hoja arrancaba con la copia y los bloques salían desordenados
    (el Durazno de la 13 copia antes que el Carré de la 13). Se ordena como
    los nombra el equipo: el número primero, como número, y después el resto."""
    m = re.match(r"\s*(\d+)(.*)", etiqueta or "")
    if m:
        return (0, int(m.group(1)), m.group(2).strip().lower())
    return (1, 0, (etiqueta or "").lower())


def _en_orden(imagenes: list[dict], etiquetas: dict[str, str]) -> list[dict]:
    return sorted(
        imagenes,
        key=lambda im: (_clave_natural(etiquetas.get(im["nombre_archivo"], im["nombre_archivo"])), im.get("orden", 0)),
    )


# ── Imágenes ancladas a celdas ────────────────────────────────────────────

def _poner(ws, jpg: bytes, ancho: int, alto: int, col0: int, fila0: int, dx: float = 0, dy: float = 0) -> None:
    """Imagen anclada a una celda (col0/fila0 en base 0), corrida dx/dy px."""
    img = XLImage(io.BytesIO(jpg))
    img.width, img.height = ancho, alto
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col0, row=fila0, colOff=int(dx * _EMU), rowOff=int(dy * _EMU)),
        ext=XDRPositiveSize2D(int(ancho * _EMU), int(alto * _EMU)),
    )
    ws.add_image(img)


def _centrada(ws, valor, col0: int, fila0: int, ancho_celda: int, alto_celda: float,
              margen_arriba: int = 0, margen: int = 8) -> None:
    im = _pil(valor)
    if im is None:
        return
    escala = min((ancho_celda - 2 * margen) / im.width, (alto_celda - margen_arriba - 2 * margen) / im.height, 1.0)
    w, h = max(1, round(im.width * escala)), max(1, round(im.height * escala))
    im = im.resize((w, h), Image.LANCZOS)
    _poner(ws, _jpeg(im), w, h, col0, fila0, (ancho_celda - w) / 2, margen_arriba + (alto_celda - margen_arriba - h) / 2)


def _tira_de_placas(placas: list[dict], etiquetas: dict[str, str], alto: int = 130) -> tuple[bytes, int, int]:
    """Las placas de un bloque, una al lado de la otra, con su formato y su nombre
    abajo. Va como UNA imagen: en Excel no hay forma confiable de poner texto
    debajo de cada foto dentro de una misma celda."""
    f_formato = _fuente(16, negrita=True)
    f_nombre = _fuente(12)
    fotos = []
    for im in placas:
        pil = _pil(im.get("vista"))
        if pil is None:  # la placa no se pudo abrir: un hueco gris del tamaño de una cuadrada
            pil = Image.new("RGB", (alto, alto), (226, 230, 236))
        fotos.append((pil.resize((max(1, round(pil.width * alto / pil.height)), alto), Image.LANCZOS), im))
    hueco, pie = 12, 38
    lienzo = Image.new("RGB", (sum(f.width for f, _ in fotos) + hueco * (len(fotos) - 1), alto + pie), "white")
    dib = ImageDraw.Draw(lienzo)
    x = 0
    for foto, im in fotos:
        lienzo.paste(foto, (x, 0))
        lineas = (
            (im.get("formato") or "", f_formato, alto + 1, (32, 42, 54)),
            (etiquetas.get(im["nombre_archivo"], im["nombre_archivo"]), f_nombre, alto + 21, (107, 118, 134)),
        )
        for texto, fuente, y, color in lineas:
            ancho = dib.textlength(texto, font=fuente)
            dib.text((x + (foto.width - ancho) / 2, y), texto, font=fuente, fill=color)
        x += foto.width + hueco
    return _jpeg(lienzo), lienzo.width, lienzo.height


# ── Texto con lo que cambia resaltado ─────────────────────────────────────

def _rico(texto: str, contra: str, lado: str, tam: int = 14) -> CellRichText:
    """Del lado de la placa se resalta en rojo lo que está mal; del lado del
    mailing, en verde lo que tiene que ir. Por palabra entera (ver correccion)."""
    return CellRichText(_tramos(texto, contra, lado, tam))


def _tramos(texto: str, contra: str, lado: str, tam: int) -> list[TextBlock]:
    """Los tramos de `_rico`, sueltos, para poder pegarlos con otros."""
    base = InlineFont(sz=tam, color=GRIS_TXT, rFont="Calibri")
    if not texto:
        vacio = "(no está)" if lado == "placa" else "(no va nada)"
        return [TextBlock(InlineFont(sz=tam - 1, i=True, color=VACIO, rFont="Calibri"), vacio)]
    color = ROJO if lado == "placa" else VERDE
    resalte = InlineFont(sz=tam, b=True, u="single", color=color, rFont="Calibri")
    if not contra:  # el otro lado está vacío: todo este texto es la diferencia
        return [TextBlock(resalte, texto)]
    seg_p, seg_m = c.diferenciar_palabras(texto, contra) if lado == "placa" else c.diferenciar_palabras(contra, texto)
    segmentos = seg_p if lado == "placa" else seg_m
    return [TextBlock(resalte if distinto else base, t) for t, distinto in segmentos]


# ── Qué entra al Excel ────────────────────────────────────────────────────

def _problemas(im: dict, origen: str = "mailing") -> list[dict]:
    probs = [f for f in im.get("filas") or [] if f.get("severidad") in ("error", "aviso")]
    if probs:
        return probs
    if im.get("error"):
        return [dict(campo="lectura", estado="revisar", placa=im["error"], mailing="")]
    # Sin pareja: se dice POR QUÉ y contra qué se pareció más, con el puntaje.
    # "No encontré este producto" a secas es una acusación sin pruebas: el
    # diseñador no puede distinguir "esta placa no es de esta campaña" de "la
    # descripción está tan mal escrita que no la reconocí".
    cands = im.get("candidatos") or []
    if im.get("sin_pareja") == "ambiguo":
        texto = f"Hay más de una fila {c.de_la_fuente(origen)} que se le parece igual: no elijo yo cuál es"
    else:
        texto = f"No encontré este producto en {c.nombre_de_la_fuente(origen)}"
    parecidas = "; ".join(f"{k['descripcion']} ({k['puntaje']:g})" for k in cands[:3])
    return [dict(campo="sin_match", estado="revisar", placa=texto,
                 mailing=(f"Lo que más se le pareció: {parecidas}" if parecidas else ""))]


def _producto(mailing: dict, im: dict) -> dict | None:
    idx = (im.get("match") or {}).get("indice")
    productos = mailing.get("productos") or []
    return productos[idx] if idx is not None and 0 <= idx < len(productos) else None


def _agrupar(placas: list[dict], etiquetas: dict[str, str], origen: str = "mailing") -> list[dict]:
    """Las placas del MISMO producto con EXACTAMENTE los mismos errores van en un
    solo bloque. Una placa sin producto en el mailing no se junta con nada.
    Los bloques quedan en el orden en que aparece su primera placa."""
    grupos: dict[tuple, dict] = {}
    for im in placas:
        probs = _problemas(im, origen)
        idx = (im.get("match") or {}).get("indice")
        clave = (
            idx if idx is not None else f"sola-{im['nombre_archivo']}",
            tuple(sorted((f["campo"], f.get("placa") or "", f.get("mailing") or "") for f in probs)),
        )
        grupos.setdefault(clave, {"placas": [], "problemas": probs})["placas"].append(im)
    for g in grupos.values():
        g["placas"].sort(key=lambda im: (
            _ORDEN_FORMATO.get(im.get("formato"), 9),
            _clave_natural(etiquetas.get(im["nombre_archivo"], im["nombre_archivo"])),
        ))
    return list(grupos.values())


# ── Hojas ─────────────────────────────────────────────────────────────────

def _cuantas_placas(total: int, con_algo: int, arreglos: int | None) -> tuple[str, str]:
    """(qué hay para corregir, qué queda bien), con singular y plural escritos
    los dos. "1 cosas para corregir" y "Las otras 1 están bien" se leían en
    cuanto una campaña tenía un solo arreglo o una sola placa sana."""
    if not con_algo:
        # Puede pasar y tiene que decirse bien: el Excel existe igual porque la
        # campaña tiene filas sin placa o la planilla dejó avisos (ver
        # `construir`). Decir "hay que corregir 0 de 4 placas" sería ruido.
        que = ("La placa está bien" if total == 1 else f"Las {total} placas están bien") + \
            ": no hay nada para corregir en ellas"
        return que, ""
    donde = "en la única placa" if total == 1 else f"en {con_algo} de las {total} placas"
    if arreglos:
        cosas = "1 cosa para corregir" if arreglos == 1 else f"{arreglos} cosas para corregir"
        que = f"{cosas}, {donde}"
    else:
        que = "Hay que corregir la única placa" if total == 1 else f"Hay que corregir {con_algo} de las {total} placas"
    resto = total - con_algo
    if resto == 1:
        otras = "  La otra está bien: no la toques."
    elif resto > 1:
        otras = f"  Las otras {resto} están bien: no las toques."
    else:
        otras = ""
    return que, otras


def _banner(ws, validacion: dict, hasta_col: str, arreglos: int | None = None) -> None:
    imagenes = validacion["imagenes"]
    total = len(imagenes)
    con_algo = sum(1 for i in imagenes if i["estado"] != "ok")
    ws.merge_cells(f"A1:{hasta_col}1")
    ws["A1"] = "CORRECCIONES DE PLACAS  ·  REDES SOCIALES"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(AZUL)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    ws.merge_cells(f"A2:{hasta_col}2")
    que, otras = _cuantas_placas(total, con_algo, arreglos)
    mailing = validacion.get("mailing") or {}
    fecha = mailing.get("fecha") or ""
    campana = f"Campaña: {fecha}      " if fecha else ""
    # Lo que la fuente no dicta se DICE, no se esconde: si la planilla no trae
    # columna de mecánica, esa parte de la placa no la miró nadie y el diseñador
    # tiene que saberlo antes de dar la campaña por revisada.
    afuera = pl.campos_que_no_dicta(mailing, validacion.get("config"))
    sin_validar = f"      La planilla no dice esto, así que no se revisó: {', '.join(afuera)}." if afuera else ""
    # Lo que la lectura de la planilla tuvo para decir va acá y no en una fila
    # aparte: son justo los avisos que delatan que el motor entendió mal el
    # archivo (una hoja de más, una fila de pie contada como producto, una
    # columna en el vocabulario de gestión), y se armaban para morir en un campo
    # que nadie leía. La hoja "La planilla" los repite con más aire.
    avisos = pl.avisos(mailing)
    ojo = ("      OJO CON LA PLANILLA: " + "  ·  ".join(avisos)) if avisos else ""
    ws["A2"] = f"{campana}{que}.{otras}{sin_validar}{ojo}"
    ws["A2"].font = Font(size=12, color="FFFFFF")
    ws["A2"].fill = _fill(AZUL_2)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=bool(ojo), indent=1)
    ws.row_dimensions[2].height = 26 if not ojo else min(90, 26 + 15 * len(avisos))


def _encabezado(ws, fila: int, columnas: list[tuple[str, str, str]]) -> None:
    for letra, texto, color in columnas:
        cel = ws[f"{letra}{fila}"]
        cel.value = texto
        cel.font = Font(size=11, bold=True, color="FFFFFF")
        cel.fill = _fill(color)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = _caja
    ws.row_dimensions[fila].height = 30


def _nada_para_corregir(ws, hasta_col: str, es_planilla: bool) -> None:
    """Puede no haber ningún arreglo y existir el archivo igual: ver
    `hay_algo_que_decir`. Una tabla vacía debajo de un encabezado que dice
    "CORRECCIONES" se lee como un error del sistema."""
    ws.merge_cells(f"A4:{hasta_col}4")
    ws["A4"] = "Ninguna placa tiene algo para corregir." + (
        " Mirá la hoja «La planilla»: ahí están las filas de la campaña que ninguna placa reclamó."
        if es_planilla else ""
    )
    ws["A4"].font = Font(size=12, italic=True, color=GRIS_SUAVE)
    ws["A4"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[4].height = 30


def _hoja_recomendada(wb, validacion: dict, etiquetas: dict[str, str]) -> None:
    ws = wb.active
    ws.title = "Recomendada"
    mailing = validacion.get("mailing") or {}
    origen = mailing.get("origen") or "mailing"
    es_planilla = origen == "planilla"
    # La columna B lleva la evidencia de la fuente. Con una planilla eso es la
    # tira dibujada de la fila (ver planilla.tira), que es más ancha que alta y
    # a 28 caracteres quedaba ilegible después de que openpyxl la escala.
    anchos = {"A": 49, "B": 52 if es_planilla else 28, "C": 17, "D": 31, "E": 31, "F": 33}
    for col, a in anchos.items():
        ws.column_dimensions[col].width = a

    placas = [i for i in _en_orden(validacion["imagenes"], etiquetas) if i["estado"] != "ok"]
    grupos = _agrupar(placas, etiquetas, origen)
    _banner(ws, validacion, "F", arreglos=len(grupos))
    _encabezado(ws, 3, [
        ("A", "LAS PLACAS", GRIS_OSC), ("B", "LA PLANILLA" if es_planilla else "EL MAILING", GRIS_OSC),
        ("C", "QUÉ REVISAR", GRIS_OSC),
        ("D", "DICE LA PLACA", ROJO), ("E", "TIENE QUE DECIR", VERDE), ("F", "QUÉ HACER", AMBAR),
    ])

    # un producto en más de un bloque: sus adaptaciones no tienen todas el mismo error
    bloques_por_producto: dict = {}
    for g in grupos:
        idx = (g["placas"][0].get("match") or {}).get("indice")
        if idx is not None:
            bloques_por_producto[idx] = bloques_por_producto.get(idx, 0) + 1

    if not grupos:
        _nada_para_corregir(ws, "F", es_planilla)

    alto_minimo = 236  # px por bloque, para que se vean las fotos
    alto_error = 74    # px mínimos por error
    arriba = 46        # px del título del bloque, arriba de las fotos
    fila = 4
    for g in grupos:
        placas_g, probs = g["placas"], g["problemas"]
        n = len(probs)
        alto_bloque = max(alto_minimo, n * alto_error)
        alto_fila = alto_bloque / n
        ini, fin = fila, fila + n - 1
        prod = _producto(mailing, placas_g[0])

        for letra in "AB":
            ws.merge_cells(f"{letra}{ini}:{letra}{fin}")
            ws[f"{letra}{ini}"].border = _caja
            ws[f"{letra}{ini}"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

        lectura = placas_g[0].get("lectura") or {}
        nombre = (prod or {}).get("descripcion") or (lectura.get("producto") or {}).get("descripcion") \
            or etiquetas.get(placas_g[0]["nombre_archivo"], placas_g[0]["nombre_archivo"])
        if len(nombre) > 46:
            nombre = nombre[:44].rstrip(" ,.") + "…"
        cuantas = "1 placa" if len(placas_g) == 1 else f"{len(placas_g)} placas, el mismo arreglo"
        idx = (placas_g[0].get("match") or {}).get("indice")
        if len(placas_g) == 1 and bloques_por_producto.get(idx, 0) > 1:
            cuantas += " · ojo: distinta de sus otras adaptaciones"
        ws[f"A{ini}"] = CellRichText([
            TextBlock(InlineFont(sz=12, b=True, color=GRIS_TXT, rFont="Calibri"), nombre + "\n"),
            TextBlock(InlineFont(sz=10, i=True, color=GRIS_SUAVE, rFont="Calibri"), cuantas),
        ])
        jpg, w, h = _tira_de_placas(placas_g, etiquetas)
        ancho_a = _px_col(anchos["A"])
        # si la tira no entra en la columna, se muestra más chica (el archivo es el mismo)
        escala = min(1.0, (ancho_a - 16) / w, (alto_bloque - arriba - 10) / h)
        w, h = round(w * escala), round(h * escala)
        _poner(ws, jpg, w, h, 0, ini - 1, (ancho_a - w) / 2, arriba + (alto_bloque - arriba - h) / 2)

        fuente_nombre = c.nombre_de_la_fuente(origen)
        # La evidencia de la fuente: el recorte del producto en la página del
        # mailing, o la tira de la fila de la planilla. Con planilla la tira se
        # dibuja al emparejar y viaja con la placa (`recorte_fuente`); el
        # `prod["recorte"]` es el camino del mailing, y también el de las
        # validaciones viejas, que traían la tira guardada en cada fila.
        evidencia = placas_g[0].get("recorte_fuente") or (prod or {}).get("recorte")
        if prod:
            # Con planilla la cita (archivo · hoja · fila) también va en la
            # imagen, pero acá va en texto para que se pueda copiar y pegar.
            idx_prod = (placas_g[0].get("match") or {}).get("indice")
            donde = pl.cita(mailing, idx_prod) if es_planilla and idx_prod is not None else ""
            # Si la tira no se pudo dibujar no se deja el hueco mudo: se dice, y
            # la cita de arriba alcanza para ir a mirar la fila a mano.
            falta = "\nNo pude dibujar la fila: miralo en el archivo." if not evidencia else ""
            ws[f"B{ini}"] = f"Así está en {fuente_nombre}" + (f"\n{donde}" if donde else "") + falta
        else:
            ws[f"B{ini}"] = f"No está en {fuente_nombre}"
        ws[f"B{ini}"].font = Font(size=10, italic=True, color=GRIS_SUAVE)
        if prod:
            _centrada(ws, evidencia, 1, ini - 1, _px_col(anchos["B"]), alto_bloque,
                      margen_arriba=38 if es_planilla else 24)

        for k, f in enumerate(probs):
            r = fila + k
            ws.row_dimensions[r].height = alto_fila * 0.75  # px -> puntos
            placa_txt, mailing_txt = f.get("placa") or "", f.get("mailing") or ""
            ws[f"C{r}"] = c.nombre_campo(f["campo"], origen)
            ws[f"C{r}"].font = Font(size=11, bold=True, color=GRIS_TXT)
            ws[f"C{r}"].fill = _fill(GRIS)
            ws[f"D{r}"] = _rico(placa_txt, mailing_txt, "placa")
            ws[f"D{r}"].fill = _fill(ROJO_F)
            ws[f"E{r}"] = _rico(mailing_txt, placa_txt, "mailing")
            ws[f"E{r}"].fill = _fill(VERDE_F)
            ws[f"F{r}"] = c.que_hacer(f, origen)
            ws[f"F{r}"].font = Font(size=13, bold=True, color=AMBAR)
            ws[f"F{r}"].fill = _fill(AMBAR_F)
            for letra in "CDEF":
                ws[f"{letra}{r}"].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
                ws[f"{letra}{r}"].border = _caja
        # línea más gruesa entre bloque y bloque, para que se lea como una tarjeta
        for letra in "ABCDEF":
            b = ws[f"{letra}{fin}"].border
            ws[f"{letra}{fin}"].border = Border(left=b.left, right=b.right, top=b.top, bottom=_grueso)
        fila = fin + 1

    ws.freeze_panes = "C4"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "3:3"


# ── «Como la pediste»: una fila por placa, una columna por cosa ───────────

# (columna, campo de la lectura). Las seis del producto salen de
# lectura["producto"]; la fecha y las leyendas, de la lectura misma. El orden y
# los títulos son los del molde.
_COLUMNAS_ANCHAS: tuple[tuple[str, str, str, int], ...] = (
    ("B", "descripcion", "DESCRIPCIÓN", 30),
    ("C", "precio_anterior", "PRECIO ANTERIOR", 14),
    ("D", "mecanica", "MECÁNICA", 12),
    ("E", "oferta_encabezado", "ARRIBA DEL PRECIO", 13),
    ("F", "oferta_precio", "PRECIO OFERTA", 12),
    ("G", "oferta_pie", "ABAJO DEL PRECIO", 12),
    ("H", "fecha", "FECHA", 22),
    ("I", "legales", "LEGALES", 26),
)
_ALTO_FILA_ANCHA = 200  # px (150 puntos): entra la vista de la placa a 162 px de alto
_SIN_REVISAR = "sin revisar: la planilla no lo dice"
_SIN_REVISAR_ALCOHOL = "leyenda de alcohol sin revisar: la planilla no la dice"


def _campos_sin_revisar(mailing: dict, config: dict, lectura: dict) -> set[str]:
    """Qué columnas de la hoja ancha quedan SIN REVISAR para esta placa: las que
    la planilla no dicta (comparador.reglas_de_la_fuente) y la fecha cuando ni
    la planilla ni la pantalla de carga la dieron. Contra un mailing no hay
    nada sin revisar: dicta todo.

    Un guion suelto en esas celdas significaba dos cosas distintas --"la placa
    no lo tiene y la fuente tampoco lo pide" y "esto no lo miró nadie"-- y la
    segunda es justo la que el diseñador tiene que saber."""
    if mailing.get("origen") != "planilla":
        return set()
    dicta = set(mailing.get("campos") or ())
    afuera = {campo for campo in ("descripcion", "precio_anterior", "mecanica", "oferta_encabezado",
                                  "oferta_precio", "oferta_pie") if campo not in dicta}
    if not ((config.get("fecha") or "").strip() or mailing.get("fecha")):
        afuera.add("fecha")
    return afuera


def _celda_plana(cel, valor: str) -> None:
    cel.value = valor or "—"
    cel.font = Font(size=11, color=GRIS_SUAVE)
    cel.fill = PatternFill(fill_type=None)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border = _caja


def _celda_sin_revisar(cel, texto: str = _SIN_REVISAR) -> None:
    cel.value = texto
    cel.font = Font(size=11, italic=True, color=GRIS_SUAVE)
    cel.fill = _fill(GRIS)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border = _caja


def _celda_con_problema(cel, rico: CellRichText) -> None:
    cel.value = rico
    cel.fill = _fill(ROJO_F)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border = _caja_roja


def _legales(lectura: dict, probs: dict[str, dict], alcohol_sin_revisar: bool) -> CellRichText | str | None:
    """La columna LEGALES: bases y leyenda de alcohol juntas, separadas por
    «  /  ». Si alguna de las dos tiene problema, la celda entera va como
    problema y solo esa parte se resalta. None = las dos vacías y sin problema."""
    bases, alcohol = lectura.get("legal_bases") or "", lectura.get("legal_alcohol") or ""
    prob_b, prob_a = probs.get("legal_bases"), probs.get("legal_alcohol")
    if not prob_b and not prob_a:
        partes = [t for t in (bases, alcohol) if t]
        texto = "  /  ".join(partes)
        if alcohol_sin_revisar:
            texto = (texto + "  /  " if texto else "") + f"({_SIN_REVISAR_ALCOHOL})"
        return texto or None
    normal = InlineFont(sz=12, color=GRIS_TXT, rFont="Calibri")
    tramos: list[TextBlock] = []
    for prob, texto in ((prob_b, bases), (prob_a, alcohol)):
        if prob:
            parte = _tramos(prob.get("placa") or "", prob.get("mailing") or "", "placa", 12)
        elif texto:
            parte = [TextBlock(normal, texto)]
        else:
            continue
        if tramos:
            tramos.append(TextBlock(normal, "  /  "))
        tramos.extend(parte)
    return CellRichText(tramos)


def _hoja_como_la_pediste(wb, validacion: dict, etiquetas: dict[str, str]) -> None:
    ws = wb.create_sheet("Como la pediste")
    mailing = validacion.get("mailing") or {}
    config = validacion.get("config") or {}
    origen = mailing.get("origen") or "mailing"
    es_planilla = origen == "planilla"
    ws.column_dimensions["A"].width = 27
    for letra, _, _, ancho in _COLUMNAS_ANCHAS:
        ws.column_dimensions[letra].width = ancho
    ws.column_dimensions["J"].width = 40

    _banner(ws, validacion, "J")
    _encabezado(ws, 3, [("A", "LA PLACA", GRIS_OSC)]
                + [(letra, titulo, GRIS_OSC) for letra, _, titulo, _ in _COLUMNAS_ANCHAS]
                + [("J", "CORRECCIÓN", AMBAR)])

    placas = [i for i in _en_orden(validacion["imagenes"], etiquetas) if i["estado"] != "ok"]
    if not placas:
        _nada_para_corregir(ws, "J", es_planilla)

    ancho_a = _px_col(27)
    for k, im in enumerate(placas):
        r = 4 + k
        ws.row_dimensions[r].height = _ALTO_FILA_ANCHA * 0.75
        etiqueta = etiquetas.get(im["nombre_archivo"], im["nombre_archivo"])
        cel = ws[f"A{r}"]
        cel.value = f"Placa {etiqueta}  ·  {im.get('formato') or ''}".rstrip(" ·")
        cel.font = Font(size=11, bold=True, color=GRIS_TXT)
        cel.alignment = Alignment(horizontal="center", vertical="top")
        cel.border = _caja
        _centrada(ws, im.get("vista"), 0, r - 1, ancho_a, _ALTO_FILA_ANCHA, margen_arriba=22)

        lectura = im.get("lectura") or {}
        producto = lectura.get("producto") or {}
        probs_lista = _problemas(im, origen)
        # Una fila por campo: si un campo trajera dos, manda la primera.
        probs: dict[str, dict] = {}
        for f in probs_lista:
            probs.setdefault(f["campo"], f)
        sin_revisar = _campos_sin_revisar(mailing, config, lectura)
        es_alcohol = bool(producto.get("es_alcohol"))
        alcohol_sin_revisar = (
            es_planilla and es_alcohol and "legal_alcohol" not in probs
            and not ((config.get("legal_alcohol") or "").strip() or mailing.get("legal_alcohol"))
        )

        for letra, campo, _, _ in _COLUMNAS_ANCHAS:
            cel = ws[f"{letra}{r}"]
            if campo == "legales":
                valor = _legales(lectura, probs, alcohol_sin_revisar)
                if isinstance(valor, CellRichText):
                    _celda_con_problema(cel, valor)
                else:
                    _celda_plana(cel, valor or "")
                continue
            texto = lectura.get(campo) if campo == "fecha" else producto.get(campo)
            texto = texto or ""
            if campo in probs:
                f = probs[campo]
                _celda_con_problema(cel, _rico(f.get("placa") or "", f.get("mailing") or "", "placa", tam=12))
            elif campo in sin_revisar:
                # La placa puede decir algo ahí: se muestra, y se dice que no se revisó.
                _celda_sin_revisar(cel, f"{texto}\n({_SIN_REVISAR})" if texto else _SIN_REVISAR)
            else:
                _celda_plana(cel, texto)

        cel = ws[f"J{r}"]
        cel.value = "\n".join(
            f"• {c.nombre_campo(f['campo'], origen)}: {c.que_hacer(f, origen, con_detalle=True)}".rstrip(": ")
            for f in probs_lista
        )
        cel.font = Font(size=12, bold=True, color=AMBAR)
        cel.fill = _fill(AMBAR_F)
        cel.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        cel.border = _caja

    ws.freeze_panes = "B4"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"


# ── «Todas las placas» ────────────────────────────────────────────────────

# "Está bien" es una palabra que el motor solo puede sostener si miró TODO:
# contra un mailing sí (dicta los siete campos), y contra una planilla solo si
# trae las seis columnas con algo escrito (ver planilla.dicta_todo). Contra una
# planilla que no trae mecánica ni encabezado ni pie, una placa con esos tres
# campos inventados salía en verde con 2 de 7 campos revisados: ahí lo que el
# motor SÍ puede afirmar es que no encontró diferencias en lo que miró, y se
# dice cuánto miró en la misma fila.
_ETIQUETA_ESTADO = {
    "ok": "Sin diferencias", "diferencias": "Hay que corregirla", "avisos": "Para revisar",
    "sin_match": "Sin pareja", "error": "No se pudo leer",
}
_COLOR_ESTADO = {"ok": "E8F6EC", "diferencias": "FDECEA", "avisos": "FFF4D6",
                 "sin_match": "F1EEFF", "error": "F2F4F7"}
_SIN_PLACA = "FFFFFF"


def _hoja_todas(wb, validacion: dict, etiquetas: dict[str, str]) -> None:
    ws = wb.create_sheet("Todas las placas")
    for col, a in zip("ABCDE", (16, 46, 10, 22, 50)):
        ws.column_dimensions[col].width = a
    ws.merge_cells("A1:E1")
    ws["A1"] = "TODAS LAS PLACAS DE LA CAMPAÑA"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(AZUL)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    _encabezado(ws, 2, [("A", "PLACA", GRIS_OSC), ("B", "PRODUCTO", GRIS_OSC), ("C", "FORMATO", GRIS_OSC),
                        ("D", "ESTADO", GRIS_OSC), ("E", "QUÉ TIENE", GRIS_OSC)])
    mailing = validacion.get("mailing") or {}
    origen = mailing.get("origen") or "mailing"
    completa = pl.dicta_todo(mailing)
    etiqueta_estado = {
        **_ETIQUETA_ESTADO,
        "sin_match": f"No está en {c.nombre_de_la_fuente(origen)}",
        **({"ok": "Está bien"} if completa else {}),
    }
    color_estado = _COLOR_ESTADO
    # Cuántos de los campos del producto pudo mirar el motor. Con un mailing son
    # los siete; con una planilla, solo los que la planilla trae escritos. Una
    # placa "sin diferencias" con 2 de 7 campos revisados no es lo mismo que una
    # con 7 de 7, y la diferencia tiene que estar en la misma fila, no solo en el
    # encabezado.
    todos = len(comparador.CAMPOS_PRODUCTO)
    dictados = len(mailing.get("campos") or ()) if origen == "planilla" else todos
    fila = 3
    for im in _en_orden(validacion["imagenes"], etiquetas):
        prod = _producto(mailing, im)
        probs = [f for f in im.get("filas") or [] if f.get("severidad") in ("error", "aviso")]
        que_tiene = ", ".join(dict.fromkeys(c.nombre_campo(f["campo"], origen) for f in probs))
        if not que_tiene:
            if im["estado"] == "ok" and not completa:
                que_tiene = f"Coincide en lo que dicta {c.nombre_de_la_fuente(origen)} ({dictados} de {todos} campos)"
            else:
                que_tiene = "—"
        valores = [
            etiquetas.get(im["nombre_archivo"], im["nombre_archivo"]),
            prod["descripcion"] if prod else "—",
            im.get("formato") or "",
            etiqueta_estado.get(im["estado"], im["estado"]),
            que_tiene,
        ]
        for letra, v in zip("ABCDE", valores):
            cel = ws[f"{letra}{fila}"]
            cel.value = v
            cel.fill = _fill(color_estado.get(im["estado"], "FFFFFF"))
            cel.border = _caja
            cel.font = Font(size=11, bold=(letra == "D"), color=GRIS_TXT)
            centrada = letra in "ACD"
            cel.alignment = Alignment(vertical="center", wrap_text=True,
                                      horizontal="center" if centrada else "left", indent=0 if centrada else 1)
        # un nombre largo (el vino) ocupa dos renglones: con la fila baja quedaba cortado
        ws.row_dimensions[fila].height = 36 if len(valores[1]) > 44 else 22
        fila += 1
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


# ── «La planilla» (solo con planilla) ─────────────────────────────────────

def _hoja_planilla(wb, validacion: dict) -> None:
    """La planilla entera, como celdas de verdad y no como imagen.

    Resuelve dos cosas de un saque. Una: el diseñador tiene la fuente completa
    adentro del mismo archivo, así que puede chequear a mano cualquier
    emparejamiento sin que nadie le adjunte la planilla aparte. Dos --y es la
    que faltaba-- las FILAS QUE NINGUNA PLACA RECLAMÓ quedan visibles solas: son
    las que quedaron sin color. Eso se venía contando (productos_mailing contra
    productos_con_placa) y no se decía nunca cuáles.

    El número de fila es el REAL del archivo, el que muestra Excel al abrirlo.
    Si el diseñador va a la fila 24 y ahí hay otra cosa, se perdió toda la
    credibilidad de una."""
    mailing = validacion.get("mailing") or {}
    info = mailing.get("planilla") or {}
    productos = mailing.get("productos") or []
    ws = wb.create_sheet("La planilla")

    # Qué placa reclamó cada fila, y en qué estado quedó.
    estado_por_indice: dict[int, list[str]] = {}
    for im in validacion["imagenes"]:
        idx = (im.get("match") or {}).get("indice")
        if idx is not None:
            estado_por_indice.setdefault(idx, []).append(im["estado"])

    titulos = info.get("titulos") or {}
    # Orden de lectura: primero lo que identifica la fila, después los precios.
    orden = ["codigo", "descripcion", "precioAnterior", "precio", "mecanica", "ofertaEncabezado", "ofertaPie"]
    columnas = [k for k in orden if k in titulos] + [k for k in titulos if k not in orden]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas) + 3)
    archivo = info.get("archivo") or "la planilla"
    hoja = f"  ·  hoja «{info['hoja']}»" if info.get("hoja") else ""
    ws["A1"] = f"LA PLANILLA CONTRA LA QUE SE VALIDÓ  ·  {archivo}{hoja}"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(AZUL)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    sin_placa = sum(1 for i in range(len(productos)) if i not in estado_por_indice)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columnas) + 3)
    # Los contadores de la lectura salen acá y no solo en la pantalla: son el
    # dato que delata que el motor contó de más (una fila de pie tomada como
    # producto) o que leyó la hoja equivocada.
    afuera = info.get("filas_ignoradas") or 0
    # Singular y plural escritos los dos: "(1 filas quedaron afuera)" se ve en
    # cuanto una planilla trae UN solo pie de reporte, que es el caso normal.
    cuantas_afuera = "1 fila quedó afuera" if afuera == 1 else f"{afuera} filas quedaron afuera"
    # Cuántas filas del archivo se juntaron en esos productos (formato largo).
    # Viajaba en la API y no se imprimía en ningún lado: es el contador que
    # delata que el motor entendió el archivo de otra manera que la persona,
    # igual que los otros dos, así que va al lado de ellos. 0 no se escribe: el
    # caso normal no junta nada y no tiene nada que decir.
    juntadas = info.get("filas_juntadas") or 0
    con_producto = "1 fila con producto, leída" if len(productos) == 1 else f"{len(productos)} filas con producto, leídas"
    if sin_placa == 0:
        reclamadas = "Todas las filas tienen su placa: ninguna quedó sin color."
    elif sin_placa == 1:
        reclamadas = (
            "La que quedó SIN COLOR es la única fila que ninguna placa reclamó: "
            "o falta la placa, o la placa que le corresponde no se le parece lo suficiente."
        )
    else:
        reclamadas = (
            f"Las que quedaron SIN COLOR son las {sin_placa} que ninguna placa reclamó: "
            "o falta la placa, o la placa que le corresponde no se le parece lo suficiente."
        )
    ws["A2"] = (
        f"{con_producto} a partir de la fila "
        f"{info.get('fila_encabezado', 1) + 1}"
        + (f" ({cuantas_afuera})" if afuera else "")
        + (f", juntando {juntadas} filas del archivo" if juntadas else "")
        + f". {reclamadas}"
    )
    ws["A2"].font = Font(size=11, color="FFFFFF")
    ws["A2"].fill = _fill(AZUL_2)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 24

    cabecera = 3
    for aviso in pl.avisos(mailing):
        ws.merge_cells(start_row=cabecera, start_column=1, end_row=cabecera, end_column=len(columnas) + 3)
        cel = ws.cell(row=cabecera, column=1, value=aviso)
        cel.font = Font(size=11, bold=True, color=AMBAR)
        cel.fill = _fill(AMBAR_F)
        cel.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[cabecera].height = 30
        cabecera += 1

    cabeceras = [("FILA", 8)] + [(titulos[k], 52 if k == "descripcion" else 16) for k in columnas]
    cabeceras += [("PLACAS", 10), ("ESTADO", 22)]
    for n, (texto, ancho) in enumerate(cabeceras, start=1):
        cel = ws.cell(row=cabecera, column=n, value=texto)
        cel.font = Font(size=11, bold=True, color="FFFFFF")
        cel.fill = _fill(GRIS_OSC)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = _caja
        ws.column_dimensions[cel.column_letter].width = ancho
    ws.row_dimensions[cabecera].height = 26

    for n, prod in enumerate(productos):
        r = cabecera + 1 + n
        estados = estado_por_indice.get(n, [])
        # El peor estado manda: si una adaptación está bien y otra mal, la fila
        # se pinta de "hay que corregirla".
        peor = next((e for e in ("error", "diferencias", "avisos", "ok") if e in estados), None)
        color = _COLOR_ESTADO.get(peor or "", _SIN_PLACA)
        valores = (prod.get("fila") or {}).get("valores") or {}
        celdas = [(prod.get("fila") or {}).get("numero", "")]
        celdas += [valores.get(k, "") for k in columnas]
        celdas += [len(estados) or "", _ETIQUETA_ESTADO.get(peor or "", "Sin placa")]
        for k, valor in enumerate(celdas, start=1):
            cel = ws.cell(row=r, column=k, value=valor)
            cel.fill = _fill(color)
            cel.border = _caja
            cel.font = Font(size=11, color=GRIS_TXT, bold=(k == len(celdas)))
            centrada = k == 1 or k > len(columnas) + 1
            cel.alignment = Alignment(vertical="center", wrap_text=True,
                                      horizontal="center" if centrada else "left", indent=0 if centrada else 1)
        ws.row_dimensions[r].height = 30 if len(str(celdas[columnas.index("descripcion") + 1])) > 48 else 20

    ws.freeze_panes = f"B{cabecera + 1}"
    ws.sheet_view.showGridLines = False


def _filas_sin_placa(validacion: dict) -> int:
    productos = (validacion.get("mailing") or {}).get("productos") or []
    reclamadas = {
        (im.get("match") or {}).get("indice") for im in validacion["imagenes"]
    } - {None}
    return len(productos) - len(reclamadas)


def hay_algo_que_decir(validacion: dict) -> bool:
    """¿Este Excel le sirve de algo al diseñador?

    No alcanza con "hay placas para corregir". Una campaña de 4 productos a la
    que le faltan 3 placas no tiene NINGUNA placa mal --las que están, están
    bien-- y es justo la que más necesita el archivo: la hoja "La planilla"
    muestra las filas que nadie reclamó. Eso se veía en pantalla y el Excel
    cortaba antes de armarse, así que no había nada para mandarle a nadie.
    Lo mismo con los avisos de la lectura de la planilla."""
    if any(i["estado"] != "ok" for i in validacion["imagenes"]):
        return True
    mailing = validacion.get("mailing") or {}
    if mailing.get("origen") != "planilla":
        # Con un mailing no existe la hoja "La planilla": no hay qué agregar.
        return False
    return bool(_filas_sin_placa(validacion) or pl.avisos(mailing))


def construir(validacion: dict) -> bytes:
    """El .xlsx de una validación.

    `validacion` tiene la forma del detalle de la API: {"mailing": {...},
    "config": {...}, "imagenes": [{nombre_archivo, formato, estado, orden,
    filas, lectura, match, error, vista}]}. La vista puede venir en bytes o
    como data URI. NadaParaCorregir cuando no hay nada que decir (ver
    `hay_algo_que_decir`)."""
    if not hay_algo_que_decir(validacion):
        raise NadaParaCorregir(
            "No hay nada para corregir: todas las placas están bien y ninguna fila quedó sin placa"
        )
    etiquetas = c.etiquetas([i["nombre_archivo"] for i in validacion["imagenes"]])
    wb = Workbook()
    _hoja_recomendada(wb, validacion, etiquetas)
    _hoja_como_la_pediste(wb, validacion, etiquetas)
    _hoja_todas(wb, validacion, etiquetas)
    if (validacion.get("mailing") or {}).get("origen") == "planilla":
        _hoja_planilla(wb, validacion)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
