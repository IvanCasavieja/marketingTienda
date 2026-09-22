"""La fuente de verdad cuando la campaña NO tiene un mailing: una planilla.

Pedido de Ivan (18/09/2026): "Hay campañas que no tienen un mailing físico...
tendríamos que agregar también la opción de que se pueda validar directamente
con un Excel", y sobre CatTi: "ahí ya directamente ni siquiera tiene que
hacerlo visual, sino que ya tiene que LEER, porque es un Excel con datos".

Eso es exactamente lo que hace este módulo, y por eso es tan corto para lo que
resuelve: la representación intermedia de la validación ya existía y ya estaba
desacoplada del PDF. `validador.preparar_mailing` devuelve un dict
{fecha, legal_alcohol, productos:[{descripcion, precio_anterior, ...}]} que se
guarda en JSONB, y todo `comparador.py` trabaja sobre ESE dict: el PDF ni se
guarda ni se le pasa a la validación de las placas. Una planilla, entonces, no
es otro motor: es otra forma de producir el mismo dict -- sin IA, sin tokens y
en milisegundos.

QUÉ PUEDE EXIGIR UNA PLANILLA Y QUÉ NO
Solo lo que trae escrito. La descripción y los dos precios están siempre; la
mecánica, el texto de arriba y el de abajo del círculo, la vigencia y la leyenda
de alcohol solo si la planilla trae una columna para eso. Lo que no trae NO se
valida y NO se inventa: se lista en `mailing["campos"]` y la pantalla y el Excel
muestran cuáles quedaron afuera. Una validación que dice "esto no lo sé" es
mejor que una que acusa a una placa con un dato que nadie escribió.

LOS PRECIOS SE COMPARAN POR IMPORTE, NO POR TEXTO
Contra un mailing la comparación es carácter por carácter porque del otro lado
hay un texto impreso: "$1090" y "$1.090" son dos decisiones distintas y una de
las dos está mal. Una planilla no trae un texto, trae un NÚMERO (74,5): cómo se
escribe en la placa no lo dicta la planilla, así que exigir un formato sería
inventarlo. Se compara el importe (ver comparador.importe) y se muestra el
precio canónico como referencia. Ver la nota del informe: es una decisión, no
un descuido.

EL LECTOR ES EL DEL CONVERTIDOR DE CENEFAS, A PROPÓSITO
`convertidor.leer_filas` ya lee .csv, .xlsx y el .xlsx-que-en-realidad-es-un-CSV
que escupe gestión (con sniffing de separador y de encoding), y
`mapear_columnas` resuelve los nombres de columna con los cien alias que se
fueron aprendiendo de los exports reales. Escribir otro lector acá sería tener
dos herramientas leyendo distinto el mismo archivo de la misma empresa.

Lo único que NO se comparte es qué fila es la de encabezados. El Convertidor
(`detectar_fila_headers`) exige una columna CODIGO porque para armar una cenefa
la necesita; acá el código no se usa para NADA -- lo que amarra la placa con su
fila es la DESCRIPCIÓN -- y exigirlo rechazaba planillas con DESCRIPCION +
PRECIO, que es todo lo que este motor necesita. Ver `_fila_de_encabezados`.

CUANDO EL MOTOR NO ENTIENDE LA PLANILLA, LO DICE
Un salto de línea adentro de una celda, una fila de pie de reporte con el total
escrito en la columna DESCRIPCION, una hoja de portada adelante de los datos,
una columna en MAYÚSCULAS con el vocabulario de gestión, un listado en formato
largo con el mismo producto repetido una vez por sucursal: en todos esos casos
el motor CREÍA entender un archivo que no entendió, y adivinaba (o reventaba).
La salida correcta no es ninguna de las dos: es avisar con un texto que una
persona pueda usar. Esos avisos viven en `mailing["planilla"]["avisos"]`, que es
lo que se guarda en la base, y salen en la pantalla Y en el Excel.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from itertools import islice

import openpyxl
from PIL import Image, ImageDraw

from app.services.cenefas.convertidor import (
    ConvertidorParseError,
    juntar_filas_por_producto,
    leer_filas,
    mapear_columnas,
    normalizar_header,
)
from app.services.cenefas.formatters import fmt_price
from app.services.rrss import archivos, imagenes
from app.services.rrss.comparador import cola_del_precio, importe

logger = logging.getLogger(__name__)

# Tope de filas de una planilla. El número NO se escribe acá: vive en
# app/data/rrss_archivos.json, que es el único lugar donde se escriben las
# reglas de los archivos que acepta la validación, y que el navegador recibe
# entero en GET /rrss/config. Así la pantalla puede decir el mismo número que
# el backend usa para rechazar, sin que nadie lo copie a mano.
#
# LO QUE JUSTIFICA EL TOPE NO ES LEER, ES TRABAJAR. Decía que "leer 10.000 filas
# es instantáneo" y hacía rato que no: con la tira de evidencia dibujada por
# fila, 900 filas eran 5,92 s de CPU y 19,5 MB de imágenes guardadas en la base,
# y eso corre adentro de `en_hilo`, que deja pasar UN trabajo de CPU a la vez
# (hilos.py) -- seis segundos con todas las validaciones de todo el mundo
# frenadas. El arreglo no fue bajar el tope: fue dejar de hacer ese trabajo. La
# tira se dibuja SOLO CUANDO SE PIDE (ver `evidencia`), o sea una por placa
# emparejada --30 placas son 0,16 s, 5 ms cada tira-- y no una por fila del
# archivo. Medido de nuevo con el cambio puesto: leer esas mismas 900 filas son
# 0,08 s de CPU, 1,2 MB de pico y 0,4 MB de JSON, contra 5,92 s y 19,5 MB.
#
# Lo que queda es lo que el tope siempre quiso decir: una validación es una
# campaña, y una planilla de miles de filas es casi seguro el archivo
# equivocado. Mejor decirlo que emparejar 30 placas contra un catálogo entero.
#
# Se chequea ANTES de materializar el archivo (ver `_cuantas_filas`). Cuando se
# chequeaba después, un catálogo de 1.000.000 de filas tardaba 148,8 s en ser
# rechazado, con un pico de 264 MB -- y por lo mismo de arriba, durante esos dos
# minutos y medio no avanzaba NINGUNA placa de NINGUNA validación de nadie.
MAX_FILAS = archivos.max_filas("planilla")

# EL TOPE CUENTA PRODUCTOS, NO FILAS CRUDAS. Desde que el formato largo es una
# entrada de primera clase (ver `_juntar_si_viene_en_formato_largo`), un listado
# trae una fila por producto Y POR SUCURSAL: 350 productos por 3 sucursales son
# 1.050 filas, y rechazarlas era acusar de "catálogo entero" justo a quien subió
# el listado que se le pidió. Así que el tope de verdad se mira DESPUÉS de
# juntar, contando productos.
#
# Pero el freno barato no se puede perder: lo que justifica todo esto es que un
# catálogo se rechaza sin materializarlo (un CSV de 200.000 filas hoy se rechaza
# en 0,00 s contando saltos de línea sobre los bytes, y uno de 1.000.000 tardaba
# 148,8 s y 264 MB cuando se chequeaba después). Sigue estando, solo que medido
# en el peor formato largo posible: MAX_FILAS productos, cada uno repetido una
# vez por sucursal. La cadena tiene 18 sucursales en el export real; 30 es techo
# con aire para las que abran, y cualquier archivo más grande que eso es un
# catálogo con cualquier cuenta que se haga.
_SUCURSALES_TECHO = 30
MAX_FILAS_CRUDAS = MAX_FILAS * _SUCURSALES_TECHO

# Cuántas filas se miran buscando la fila de encabezados. Mismo criterio que el
# Convertidor (_HEADER_SCAN_ROWS): el export real trae una fila de título y una
# en blanco antes, no veinte.
_FILAS_A_MIRAR = 10


class PlanillaInvalida(ValueError):
    """La planilla no se puede usar como fuente -- mensaje apto para el usuario."""


# Los campos del producto que una planilla PUEDE traer, y de qué campo del
# lector salen. Los tres primeros los trae cualquier listado de campaña; los
# otros solo si alguien les puso una columna (ver _ALIAS_RRSS).
_DE_LA_COLUMNA: dict[str, str] = {
    "descripcion": "descripcion",
    "precio_anterior": "precioAnterior",
    "oferta_precio": "precio",
    "mecanica": "mecanica",
    "oferta_encabezado": "ofertaEncabezado",
    "oferta_pie": "ofertaPie",
}

# Columnas que son solo de la validación de placas y que el Convertidor no
# conoce (no le sirven para una cenefa). Se normalizan con el mismo _norm del
# Convertidor para que "Texto arriba" y "TEXTOARRIBA" sean la misma columna.
#
# Los títulos que el propio informe usa ("ARRIBA DEL PRECIO", "ABAJO DEL PRECIO",
# ver _TITULO_POR_DEFECTO y la hoja «Como la pediste» del Excel) están acá a
# propósito: si alguien copia una columna del informe a su planilla, tiene que
# entrar. Hasta el 22/09/2026 no entraban y nadie podía adivinar los nombres.
_ALIAS_RRSS: dict[str, str] = {
    "mecanica": "mecanica",
    "titularoferta": "mecanica",
    "encabezadooferta": "ofertaEncabezado",
    "textoarriba": "ofertaEncabezado",
    "textoarribadelprecio": "ofertaEncabezado",
    "arribadelprecio": "ofertaEncabezado",
    "pieoferta": "ofertaPie",
    "textoabajo": "ofertaPie",
    "textoabajodelprecio": "ofertaPie",
    "abajodelprecio": "ofertaPie",
    "vigencia": "vigencia",
    "vigenciacampana": "vigencia",
    "vigenciadelacampana": "vigencia",
    "fecha": "vigencia",
    "fechas": "vigencia",
    "fechacampana": "vigencia",
    "fechadecampana": "vigencia",
    "fechadelacampana": "vigencia",
    "legalalcohol": "legalAlcohol",
    "legaldealcohol": "legalAlcohol",
    "leyendaalcohol": "legalAlcohol",
    "leyendadealcohol": "legalAlcohol",
}

# Cómo se llama cada columna opcional para una persona: es lo que la pantalla
# de carga muestra como "columnas que entiende". Un solo nombre por columna, el
# que también usa el informe, para no enseñar dos vocabularios.
COLUMNAS_OPCIONALES: tuple[tuple[str, str], ...] = (
    ("mecanica", "MECÁNICA"),
    ("ofertaEncabezado", "ARRIBA DEL PRECIO"),
    ("ofertaPie", "ABAJO DEL PRECIO"),
    ("vigencia", "VIGENCIA"),
    ("legalAlcohol", "LEYENDA ALCOHOL"),
)

# Qué columna de la planilla mira cada campo, para poder encuadrar la celda
# exacta en la tira de evidencia. Los dos precios y la descripción son los que
# de verdad se marcan; el resto cae en su propia columna.
COLUMNA_DEL_CAMPO = dict(_DE_LA_COLUMNA)

_TITULO_POR_DEFECTO = {
    "codigo": "CÓDIGO",
    "descripcion": "DESCRIPCIÓN",
    "precioAnterior": "PRECIO ANTERIOR",
    "precio": "PRECIO",
    "mecanica": "MECÁNICA",
    "ofertaEncabezado": "ARRIBA DEL PRECIO",
    "ofertaPie": "ABAJO DEL PRECIO",
}

# LA PLANILLA PUEDE VENIR DE CUALQUIER FORMA. Pedido de Ivan (22/09/2026), el día
# que subió el listado real de la campaña y el motor lo rechazó:
#   "La idea de esta validación es no esperar un formato de Excel. Puede venir un
#    Excel de cualquier manera -- normalmente viene de la forma que está en
#    descargas, pero puede venir de cualquier forma habida y por haber. Lo único
#    que tenés que esperar es descripción, y el resto de las cosas
#    lamentablemente el agente tiene que razonarlas en el momento."
# El listado real decía CODIGO | NOMBRE ARTÍCULO | MONEDA | PRECIOANT | PRECIO |
# OFERTA: la descripción estaba en NOMBRE ARTÍCULO --una columna que el lector
# descartaba a propósito, suponiendo que ahí siempre viene el nombre corto de
# gestión-- y la mecánica ("2x$75") en OFERTA, una columna que ni se miraba. Con
# una tabla de nombres fija eso no se arregla agregando alias: el próximo listado
# va a traer otros nombres.
#
# Así que qué es cada columna lo decide CatTi leyendo la planilla (ver
# catti.interpretar_planilla), igual que decide qué es cada cosa en un PDF. Lo que
# NO decide son los valores: esos se copian de las celdas tal cual, acá abajo.
# Una planilla es dato exacto y tiene que seguir siéndolo -- el modelo elige la
# columna, no reescribe el contenido.
#
# Estos son los datos que se le pide que ubique: la clave es el nombre con el que
# CatTi los devuelve, el valor es el nombre interno de este módulo (el mismo que
# usa el lector por nombres de columna, así el resto de `leer` no se entera de
# quién decidió). catti.py arma su herramienta con ESTAS claves y se niega a
# arrancar si alguna falta o sobra: la lista vive en un solo lugar.
CAMPOS_DE_LA_PLANILLA: dict[str, str] = {
    "descripcion": "descripcion",
    "precio_anterior": "precioAnterior",
    "precio_oferta": "precio",
    "moneda": "moneda",
    "mecanica": "mecanica",
    "arriba_del_precio": "ofertaEncabezado",
    "abajo_del_precio": "ofertaPie",
    "vigencia": "vigencia",
    "leyenda_alcohol": "legalAlcohol",
    "sucursal": "sucursal",
    "stock": "stock",
    "codigo": "codigo",
}

# Cómo se le nombra cada dato a una persona cuando se le cuenta qué entendió
# CatTi ("Descripción ← NOMBRE ARTÍCULO").
_NOMBRE_DEL_DATO = {
    "descripcion": "Descripción",
    "precioAnterior": "Precio anterior",
    "precio": "Precio de oferta",
    "moneda": "Moneda",
    "mecanica": "Mecánica",
    "ofertaEncabezado": "Arriba del precio",
    "ofertaPie": "Abajo del precio",
    "vigencia": "Vigencia",
    "legalAlcohol": "Leyenda de alcohol",
    "sucursal": "Sucursal",
    "stock": "Stock",
    "codigo": "Código",
}

# Lo que la pantalla de carga dice que CatTi busca en una planilla. Sale de la
# misma lista que usa CatTi (menos sucursal, stock y código, que no se comparan
# contra la placa): así la ayuda nunca promete algo que el motor no mira.
DATOS_QUE_BUSCA: list[str] = [
    _NOMBRE_DEL_DATO[campo].lower()
    for campo in CAMPOS_DE_LA_PLANILLA.values()
    if campo not in ("sucursal", "stock", "codigo")
]

# Cuánto de la planilla se le muestra a CatTi para que decida. No hace falta el
# archivo entero: con el principio (títulos y las primeras filas) y el final (el
# pie del reporte, si lo hay) alcanza para saber qué es cada columna, y así una
# campaña de 800 productos cuesta lo mismo que una de 8.
_MUESTRA_FILAS = 25
_MUESTRA_COLA = 5
_MUESTRA_HOJAS = 8
_MUESTRA_COLUMNAS = 30
_MUESTRA_ANCHO_CELDA = 60


@dataclass
class Planilla:
    mailing: dict
    avisos: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# Leer
# --------------------------------------------------------------------------

# Un salto de línea adentro de una celda (Alt+Enter, o una celda pegada desde
# otro lado) no es contenido: es dato sucio. Del otro lado la placa dice el
# texto en un renglón, así que comparar contra un texto con un \n adentro sería
# marcar como error algo que nadie ve. Y además reventaba: PIL no puede medir
# texto multilínea ("can't measure length of multiline text"), así que un solo
# Alt+Enter en una fila VECINA de la que se estaba dibujando tiraba la
# validación entera con un HTTP 500 sin ninguna pista.
# Los de Unicode (\u2028 y \u2029) llegan de verdad: una celda pegada desde
# una web los trae, y openpyxl los entrega tal cual.
_SALTOS = "\r\n\t\v\f\u2028\u2029"
_SALTO_DE_LINEA = re.compile(r"\s*[%s]+\s*" % re.escape(_SALTOS))


def _una_linea(texto: str) -> str:
    """El texto en un solo renglón. Los espacios dobles NO se tocan (un "300  g"
    es una diferencia de verdad y tiene que verse); lo que se va son los saltos
    de línea, que no son contenido -- ver `_SALTO_DE_LINEA`."""
    return _SALTO_DE_LINEA.sub(" ", texto or "").strip()


def _texto(valor) -> str:
    """El valor de una celda como lo ve una persona al abrir el archivo.

    Un entero que openpyxl entrega como float (501233.0) se escribe sin el
    ".0": ese número es un código de artículo, y con la cola decimal deja de
    parecerse al que está en la planilla."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, float):
        # coma decimal, que es como se ve en un Excel en español
        return f"{valor:g}".replace(".", ",")
    return _una_linea(str(valor))


def _precio(valor, moneda: str) -> str:
    """El precio canónico: el símbolo de la planilla y el formato uruguayo
    (`fmt_price`, el mismo con el que se imprimen las cenefas). Es lo que se
    MUESTRA como referencia; lo que se COMPARA es el importe.

    SIN COLUMNA MONEDA NO VA SÍMBOLO. Antes se ponía "$" igual, y ahí el motor
    afirmaba en pesos algo que la planilla nunca dijo: contra una placa en
    dólares salía un error de moneda inventado por nosotros. Es el mismo
    criterio que ya usa `comparador.misma_moneda` para la placa sin símbolo: si
    un lado no trae marca reconocible no se afirma nada. El precio queda
    "74,50" y la comparación por importe sigue igual."""
    texto = _texto(valor)
    if not texto:
        return ""
    numero = importe(texto)
    if numero is None:
        return texto  # la celda trae algo que no es un número: se muestra tal cual
    m = _RE_NUMERO.search(texto)
    # Lo que la celda trae ALREDEDOR del número no se tira: el símbolo que la
    # persona escribió ("U$S 149" sin columna MONEDA) y lo que acompaña al
    # precio ("$340 unidad"). Sin esto, "$340 unidad" quedaba en "$340" y la
    # planilla perdía justo la palabra que el diseñador había puesto ahí.
    simbolo = (moneda or "").strip() or texto[:m.start()].strip()
    cola = texto[m.end():].strip()
    return f"{simbolo}{fmt_price(numero)}" + (f" {cola}" if cola else "")


_RE_NUMERO = re.compile(r"\d[\d.,]*")


def _fila_de_datos(valores: dict) -> bool:
    """Una fila sirve si tiene descripción."""
    return bool(valores.get("descripcion"))


def _fila_vacia(valores: dict) -> bool:
    """Una fila en la que ninguna columna reconocida trae nada. Es el renglón en
    blanco que queda entre el título y los datos, o los que Excel agrega al
    final: no es una fila que "quedó afuera", y contarla como "sin descripción"
    le hacía decir al informe "Afuera: 3 filas sin descripción" en una planilla
    perfecta. Una fila con código o precio y sin descripción SÍ se cuenta: ahí
    hay un producto al que le falta lo que lo amarra con la placa."""
    return not any(v for v in valores.values())


# Cómo empieza la fila de pie de un reporte de gestión. Hasta ahora se decía que
# estas filas "se caen solas porque no tienen descripción", y era cierto SOLO
# porque en el archivo de prueba el total estaba escrito en la columna CODIGO.
# Los listados de gestión de verdad lo escriben en la columna DESCRIPCION, y
# entonces "TOTAL: 5 artículos" entraba como un producto más: el Excel le decía
# al diseñador que le faltaba la placa de "TOTAL: 5 artículos".
#
# CADA ALTERNATIVA VIENE CON UN EJEMPLO AL LADO, Y EL REGEX SE ARMA CON ESTA
# MISMA TABLA. No es decoración: dos de las alternativas no podían matchear
# NADA y nadie se enteró, porque una alternativa muerta no rompe nada -- solo
# deja pasar el pie como si fuera un producto, que es justo lo que esto venía a
# arreglar. A las dos las mataba la frontera de palabra que estaba al final del
# grupo entero:
#   - "usuario:" seguido de un espacio: ":" y " " son los dos no-palabra, así
#     que ahí NO hay frontera y "usuario: ADMIN" nunca matcheó;
#   - "fecha de emisi" seguido de "ón": "i" y "ó" son los dos caracteres de
#     palabra, misma historia al revés.
# Por eso la frontera ahora va ALTERNATIVA POR ALTERNATIVA, donde de verdad
# corresponde, y las dos que no terminan en palabra escriben su propio final.
# El test recorre esta tabla entera y falla si alguna alternativa no puede
# matchear ni su propio ejemplo: así la que agreguen mañana no se muere en
# silencio.
INICIOS_DE_PIE: tuple[tuple[str, str], ...] = (
    (r"total(es)?\b", "TOTAL: 3 artículos"),
    (r"sub\s*total\b", "Sub total general"),
    (r"cantidad\s+(de|total)\b", "Cantidad de artículos: 12"),
    (r"nro\.?\s*de\s+registros\b", "Nro. de registros: 12"),
    (r"registros\b", "Registros: 12"),
    (r"generado\b", "Generado el 18/09/2026"),
    (r"emitido\b", "Emitido por gestión"),
    (r"impreso\b", "Impreso el 18/09/2026"),
    (r"exportado\b", "Exportado el 18/09/2026"),
    (r"fin\s+del?\s+(reporte|listado|informe)\b", "Fin del reporte"),
    (r"resumen\b", "Resumen de la campaña"),
    # el ":" ES el final de la palabra: pedir una frontera después lo mataba
    (r"usuario\s*:", "usuario: ADMIN"),
    # con tilde y sin tilde: los dos salen de gestión según quién exportó
    (r"fecha\s+de\s+(emisi|impresi)[óo]n", "Fecha de emisión: 18/09/2026"),
)

_INICIO_DE_PIE = re.compile(
    r"^\s*(?:%s)" % "|".join(patron for patron, _ in INICIOS_DE_PIE),
    re.IGNORECASE,
)


# Un código de artículo de gestión tiene cuatro dígitos o más (501233). Un "3"
# escrito en la columna CODIGO no es un código: es el contador de un pie de
# reporte, que es justo lo que hay que poder distinguir.
_RE_CODIGO = re.compile(r"\d{4}")
_RE_FECHA = re.compile(r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}$")


def _tiene_codigo(valores: dict) -> bool:
    """¿La fila trae un código de artículo de verdad?"""
    return bool(_RE_CODIGO.search(valores.get("codigo", "")))


def _agregados(filas: list[dict]) -> set[float]:
    """Los números que un PIE puede traer y una fila de producto no: el total de
    cada columna numérica de las filas de arriba, y cuántas filas son.

    Es la diferencia entre "TOTAL: 3 artículos | 402,50" y un producto: el 402,50
    no es el precio de nada, es la suma de la columna. Un producto que empieza
    como un pie ("Total Blue agua saborizada. 900 ml") trae su propio precio, y
    su precio no es el total de nadie."""
    totales: dict[str, float] = {}
    for v in filas:
        for campo, valor in v.items():
            if campo == "descripcion" or _RE_FECHA.search(valor or ""):
                continue
            n = importe(valor)
            if n is not None:
                totales[campo] = totales.get(campo, 0.0) + n
    return set(totales.values()) | {float(len(filas))}


def _solo_trae_agregados(valores: dict, arriba: list[dict]) -> bool:
    """Lo que esta fila trae en las OTRAS columnas, ¿es un agregado del reporte
    o un dato de producto?

    Agregado: una fecha ("Generado el 18/09/2026" con la fecha al lado), el
    total de una columna, la cantidad de filas. Cualquier otra cosa --un precio
    que no es la suma de nada, un texto suelto-- es un dato de producto, y
    entonces esta fila no se descarta: se cuenta como producto y se avisa."""
    posibles = _agregados(arriba)
    for campo, valor in valores.items():
        if campo == "descripcion" or not valor:
            continue
        if _RE_FECHA.search(valor):
            continue
        n = importe(valor)
        if n is None or not any(abs(n - a) < 0.01 for a in posibles):
            return False
    return True


def _pies_del_final(crudas: list[tuple[int, dict]]) -> tuple[set[int], set[int]]:
    """Cuáles de las filas leídas son el pie del reporte y cuáles lo PARECEN.

    Devuelve (posiciones que son pie, posiciones dudosas).

    Antes se pedía que la fila no trajera NINGÚN otro dato, y por eso el pie
    entraba como producto en su forma más común: un pie de gestión de verdad
    trae el total de plata, o el contador de registros, o la fecha de emisión.
    Con "TOTAL: 3 artículos | 402,50" al pie, el Excel le decía al diseñador que
    le faltaba la placa de "TOTAL: 3 artículos", sin un solo aviso.

    La señal no es "no trae nada más", son cuatro cosas juntas:
      1. el texto empieza como un pie ("Total: 5 artículos", "Generado el...");
      2. la fila está AL FINAL del bloque de datos -- por eso se camina de abajo
         para arriba y se corta en la primera fila que es un producto: un pie en
         el medio del listado no existe, un producto que se llama "Total" sí;
      3. no trae un código de artículo (ver `_tiene_codigo`);
      4. lo que trae en las otras columnas es un AGREGADO y no un dato de
         producto (ver `_solo_trae_agregados`).

    El control que no se puede romper es que "Total Blue agua saborizada. 900 ml"
    siga siendo un producto: cae por el 3 (trae su código) y, en una planilla sin
    columna CODIGO, por el 4 (su precio no es el total de nada).

    Una fila que cumple 1, 2 y 3 pero NO el 4 queda como dudosa: se cuenta como
    producto --que es lo que era hasta ahora-- pero se avisa, porque el motor no
    puede decidirla y quien mira el archivo sí."""
    pies: set[int] = set()
    dudosas: set[int] = set()
    for pos in range(len(crudas) - 1, -1, -1):
        valores = crudas[pos][1]
        if not _INICIO_DE_PIE.match(valores.get("descripcion", "")) or _tiene_codigo(valores):
            break
        if _solo_trae_agregados(valores, [v for _, v in crudas[:pos]]):
            pies.add(pos)
            continue
        # Y se SIGUE subiendo, no se corta. Una dudosa cortaba el camino, así
        # que un pie limpio que estuviera ARRIBA de una dudosa no se evaluaba
        # nunca y entraba como producto sin un aviso propio: con
        # "TOTAL: 3 artículos | 402,50" y abajo "Generado el 18/09/2026 | por
        # USUARIO", el TOTAL se colaba entero. Lo que corta sigue siendo la
        # primera fila que NO tiene cara de pie (el criterio 2: un pie está al
        # final); una dudosa tiene cara de pie, solo que el motor no la puede
        # decidir, y eso no dice nada de la que está más arriba.
        dudosas.add(pos)
    return pies, dudosas


def _juntar_si_viene_en_formato_largo(filas: list[dict]) -> tuple[list[dict], dict]:
    """Un listado en FORMATO LARGO --una fila por producto y por sucursal-- con
    una fila por producto, usando el juntador del Convertidor.

    Es el archivo que la gente tiene más a mano, y entraba como 90 productos
    repetidos: todas las placas quedaban sin pareja (dos filas idénticas empatan
    y el emparejamiento no adivina) y no se decía ni una palabra.

    NO se escribe un segundo juntador: `convertidor.juntar_filas_por_producto`
    ya resuelve exactamente esto --agrupa respetando el orden de aparición, se
    queda con el primer valor no vacío de cada campo y cuenta los productos que
    traen datos distintos entre sus filas-- y tenerlo dos veces es la forma
    segura de que un día las dos herramientas junten distinto el mismo archivo.

    Se agrupa por DESCRIPCIÓN y no por código porque acá la identidad del
    producto es la descripción: es lo que amarra la placa con su fila (decisión
    de Ivan), y una planilla de campaña puede no traer columna CODIGO.

    El stock se manda vacío: el juntador SUMA las unidades de cada sucursal y acá
    los valores son texto (sumar "12" + "3" reventaría), y este motor no usa el
    stock para nada. Si la planilla no trae columna SUCURSAL --o la trae y no hay
    nada que juntar-- `resumen["filas"]` viene en cero y se devuelven las filas
    TAL CUAL, sin tocar una sola clave."""
    juntadas, resumen = juntar_filas_por_producto(
        [dict(v, stock=None) for v in filas], clave="descripcion",
    )
    # HUBO JUNTADO SOLO SI ENTRARON MÁS FILAS QUE LOS PRODUCTOS QUE SALIERON.
    # El juntador arranca por la SOLA existencia de la columna SUCURSAL y
    # devuelve `filas` = todas las que le pasaron, haya juntado o no: una
    # planilla normal que trae una columna SUCURSAL --12 productos distintos,
    # todos "Central", una fila cada uno-- salía con `filas` = 12 y disparaba el
    # aviso de "viene en formato largo" sin que se hubiera juntado una sola
    # fila. Un aviso que salta cuando no corresponde es tan malo como el que
    # falta: el caso normal tiene que ser CERO avisos. Así que cuando no se
    # juntó nada se devuelve lo mismo que cuando no hay columna SUCURSAL --las
    # filas TAL CUAL y el resumen en cero-- y de ahí para abajo no cambia nada:
    # ni el aviso, ni `filas_juntadas`, ni el renumerado de las filas.
    if resumen["filas"] <= resumen["productos"]:
        return filas, {"filas": 0, "productos": 0, "sucursales": [], "con_diferencias": 0}
    for v in juntadas:
        # "_stock" es la carpeta por sucursal que arma el Convertidor para
        # partir la descarga de cenefas; acá no se muestra ni se usa.
        v.pop("_stock", None)
    return juntadas, resumen


def _columna_de_descripcion(col_map: dict[int, str]) -> str | None:
    """Cuál de las columnas del archivo es la descripción que va en la placa,
    cuando la planilla se lee por los NOMBRES de sus columnas -- que es el
    camino de repuesto: el de siempre es que lo decida CatTi (ver
    CAMPOS_DE_LA_PLANILLA).

    DESCRIPCION gana sobre DESCRIPCIONWEB, y las dos sobre NOMBREARTICULO. Hasta
    el 22/09/2026 NOMBREARTICULO no contaba nunca, suponiendo que ahí siempre
    viene el nombre corto de gestión en mayúsculas ("CERVEZA PATRICIA LATA
    473ML"). El listado real de la campaña de los rompeprecios trae la
    descripción impresa justo ahí ("Arvejas TIENDA INGLESA. 300 g") y ninguna
    otra columna de texto, y el motor lo rechazaba entero. Ahora cuenta si es la
    única; y si de verdad viene en mayúsculas de gestión, lo agarra
    `_parece_vocabulario_de_gestion` y lo avisa, en vez de acusar placas."""
    campos = set(col_map.values())
    for candidata in ("descripcionExcel", "descripcionWeb", "nombreArticulo"):
        if candidata in campos:
            return candidata
    return None


def _columnas_de(encabezados) -> dict[int, str]:
    """{índice de columna -> campo} de una fila de encabezados: los alias del
    Convertidor más los que solo le importan a la validación de placas."""
    col_map = dict(mapear_columnas(encabezados))
    for idx, celda in enumerate(encabezados):
        if celda is None or idx in col_map:
            continue
        campo = _ALIAS_RRSS.get(normalizar_header(celda))
        if campo:
            col_map[idx] = campo
    return col_map


def _fila_de_encabezados(filas: list[tuple]) -> int | None:
    """Índice (0-based) de la fila de encabezados, o None.

    La señal es una columna de DESCRIPCIÓN, no una de CÓDIGO. El Convertidor
    busca CODIGO porque para armar una cenefa lo necesita; acá el código no se
    usa nunca -- lo que amarra la placa con su fila es la descripción -- y
    exigirlo rechazaba una planilla con DESCRIPCION + PRECIO, que es todo lo
    que este motor necesita.

    Entre dos filas candidatas gana la que reconoce más columnas: en un export
    con título arriba, la fila del título puede tener por casualidad una celda
    que normaliza a un alias conocido."""
    mejor, mejor_cuantas = None, 0
    for i, fila in enumerate(filas[:_FILAS_A_MIRAR]):
        col_map = _columnas_de(fila)
        if _columna_de_descripcion(col_map) is None:
            continue
        if len(col_map) > mejor_cuantas:
            mejor, mejor_cuantas = i, len(col_map)
    return mejor


def _hojas(datos: bytes, filename: str) -> list[str]:
    if (filename or "").lower().endswith(".csv"):
        return []
    try:
        from app.services.cenefas.convertidor import listar_hojas

        return listar_hojas(datos, filename)
    except Exception:
        return []


def _cuantas_filas(datos: bytes, filename: str, hoja: str | int | None, tope: int) -> int:
    """Cuántas filas tiene el archivo, mirando a lo sumo `tope`+1.

    Existe para poder rechazar un catálogo ANTES de materializarlo. `leer_filas`
    devuelve una lista: con 1.000.000 de filas eso son 149 segundos y 264 MB
    adentro del único hilo de CPU que tiene el servicio (hilos.py), o sea con
    todas las validaciones de todo el mundo frenadas mientras tanto. Acá se leen
    MAX_FILAS_CRUDAS+1 y se corta.

    Si no se puede contar (un .xlsx raro, un .xlsx que en realidad es un CSV) se
    devuelve 0: el chequeo de verdad lo vuelve a hacer `leer` con la lista ya
    armada, así que no se deja pasar nada -- solo se pierde el atajo."""
    if (filename or "").lower().endswith(".csv"):
        # Contar saltos de línea sobre los bytes es un barrido de memoria: 17 MB
        # en milisegundos, contra parsear un millón de filas de CSV.
        return datos.count(b"\n") + (0 if datos.endswith(b"\n") else 1)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        try:
            nombres = wb.sheetnames
            if isinstance(hoja, str):
                candidatas = [hoja] if hoja in nombres else nombres[:1]
            elif hoja is not None:
                candidatas = nombres[hoja:hoja + 1] or nombres[:1]
            else:
                # Sin hoja elegida se van a mirar TODAS (ver `_elegir_hoja`), así
                # que el tope tiene que mirarlas todas también: si no, un
                # catálogo en la segunda hoja entraba igual.
                candidatas = nombres
            # read_only entrega las filas de a una, perezosamente: con islice
            # solo se parsean las primeras `tope`+1 de cada hoja.
            return max(
                (sum(1 for _ in islice(wb[n].iter_rows(values_only=True), tope + 1)) for n in candidatas),
                default=0,
            )
        finally:
            wb.close()
    except Exception:
        logger.warning("rrss: no pude contar las filas de %s antes de leerla", filename, exc_info=True)
        return 0


def _demasiados_productos() -> PlanillaInvalida:
    """El tope de verdad, el que se mira después de juntar: son PRODUCTOS."""
    return PlanillaInvalida(
        f"La planilla tiene más de {MAX_FILAS} productos. Subí el listado de esta "
        f"campaña, no el catálogo entero."
    )


def _demasiadas_filas_crudas() -> PlanillaInvalida:
    """El freno barato, el de antes de materializar el archivo. Sin el número
    exacto de filas a propósito: se corta de contar en MAX_FILAS_CRUDAS+1, justo
    para no tener que leer el archivo entero.

    Dice las dos cuentas porque el que sube el archivo no tiene por qué saber
    cuál de los dos números lo frenó: ni juntando una fila por sucursal eso
    entra en una campaña."""
    return PlanillaInvalida(
        f"La planilla tiene más de {MAX_FILAS_CRUDAS} filas. Aun si viniera en formato "
        f"largo (una fila por producto y por sucursal) son muchos más de {MAX_FILAS} "
        f"productos: subí el listado de esta campaña, no el catálogo entero."
    )


def _parece_vocabulario_de_gestion(descripciones: list[str]) -> bool:
    """¿La columna elegida trae el nombre CORTO de gestión en vez del texto que
    va impreso en la placa?

    Una descripción impresa lleva minúsculas ("Cerveza PATRICIA lata. 473 ml");
    el nombre de gestión no tiene ni una ("CERVEZA PATRICIA LATA 473ML"). Con
    una planilla así, TODAS las placas quedaban acusadas de tener la descripción
    mal escrita y el Excel le ordenaba al diseñador reescribir la campaña entera
    en MAYÚSCULAS con abreviaturas. Eso quema la confianza en el motor el primer
    día, y encima el equivocado era el motor.

    Se piden al menos 3 filas para no disparar con una planilla de dos líneas
    donde el diseñador escribió dos títulos en mayúscula a propósito."""
    utiles = [d for d in descripciones if any(c.isalpha() for c in d)]
    if len(utiles) < 3:
        return False
    sin_minusculas = sum(1 for d in utiles if not any(c.islower() for c in d))
    return sin_minusculas / len(utiles) >= 0.8


def _elegir_hoja(datos: bytes, filename: str, hoja: str | int | None) -> tuple[list[tuple], str, list[str]]:
    """(filas, nombre de la hoja leída, avisos) de la hoja que de verdad tiene
    los datos.

    Hasta ahora se mandaba siempre la PRIMERA hoja. Con una portada adelante eso
    rechazaba el archivo con "No encontré la fila de encabezados" --que apunta al
    problema equivocado y ni siquiera nombra las hojas que hay-- y con un
    borrador viejo adelante validaba contra el borrador sin chistar.

    Ahora se elige la primera hoja que tenga una fila de encabezados de verdad, y
    si hay más de una candidata se AVISA cuál se leyó y cuáles son las otras:
    elegir por la persona es exactamente lo que no hay que hacer, callarse
    tampoco."""
    avisos: list[str] = []
    nombres = _hojas(datos, filename)
    if hoja is not None or len(nombres) <= 1:
        return _filas_de(datos, filename, hoja), _nombre_de_hoja(nombres, hoja), avisos

    con_datos: list[tuple[str, list[tuple]]] = []
    for n, nombre in enumerate(nombres):
        filas = _filas_de(datos, filename, n)
        if _fila_de_encabezados(filas) is not None:
            con_datos.append((nombre, filas))
    if not con_datos:
        raise PlanillaInvalida(
            "No encontré la fila de encabezados en ninguna hoja del archivo: tiene que "
            "haber una fila con una columna DESCRIPCION. Las hojas del archivo son: "
            + ", ".join(nombres)
        )
    nombre, filas = con_datos[0]
    if len(con_datos) > 1:
        otras = ", ".join("«%s»" % n for n, _ in con_datos[1:])
        avisos.append(
            "El archivo tiene más de una hoja con datos: leí «%s», la primera. Las otras "
            "son %s — si la campaña está en alguna de esas, poné esa hoja primera y "
            "volvé a subir el archivo." % (nombre, otras)
        )
    elif nombre != nombres[0]:
        saltadas = ", ".join("«%s»" % n for n in nombres[:nombres.index(nombre)])
        avisos.append("Leí la hoja «%s»: las de adelante (%s) no tienen datos." % (nombre, saltadas))
    return filas, nombre, avisos


def _filas_de(datos: bytes, filename: str, hoja: str | int | None) -> list[tuple]:
    try:
        return leer_filas(datos, filename, hoja)
    except ConvertidorParseError as exc:
        raise PlanillaInvalida(str(exc)) from exc
    except Exception as exc:
        raise PlanillaInvalida("No pude abrir la planilla; ¿es un .xlsx o un .csv?") from exc


def letra(indice: int) -> str:
    """La letra de columna como la muestra Excel: 0 -> A, 25 -> Z, 26 -> AA."""
    s = ""
    n = indice + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def indice_de_letra(texto: str) -> int | None:
    """La inversa de `letra`: 'A' -> 0, 'AA' -> 26. None si no es una letra de
    columna (CatTi devuelve cadena vacía para "esta planilla no trae ese dato")."""
    t = (texto or "").strip().upper()
    if not t or not t.isalpha() or not t.isascii():
        return None
    n = 0
    for ch in t:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _celda_para_catti(valor) -> str:
    t = _texto(valor)
    return t if len(t) <= _MUESTRA_ANCHO_CELDA else t[:_MUESTRA_ANCHO_CELDA - 1] + "…"


def _tabla_para_catti(nombre: str, filas: list[tuple]) -> str:
    """Una hoja como la ve una persona al abrirla: número de fila de Excel y
    letra de columna. CatTi contesta con esas mismas coordenadas, así lo que
    decide se puede verificar contra el archivo y mostrarle a la persona."""
    ancho = 0
    for f in filas:
        for i, v in enumerate(f[:_MUESTRA_COLUMNAS]):
            if _texto(v):
                ancho = max(ancho, i + 1)
    total = len(filas)
    titulo = "Hoja «%s» — %d filas" % (nombre, total) if nombre else "Archivo CSV — %d filas" % total
    if not ancho:
        return titulo + " (vacía)"
    lineas = [titulo, "fila | " + " | ".join(letra(i) for i in range(ancho))]
    mostrar = list(range(min(total, _MUESTRA_FILAS)))
    if total > _MUESTRA_FILAS:
        cola = list(range(max(_MUESTRA_FILAS, total - _MUESTRA_COLA), total))
        if cola and cola[0] > _MUESTRA_FILAS:
            lineas_cortadas = (cola[0] - _MUESTRA_FILAS)
            mostrar.append(-lineas_cortadas)  # marca de "no se muestran N filas"
        mostrar.extend(cola)
    for i in mostrar:
        if i < 0:
            lineas.append("  (… %d filas más con el mismo formato, no se muestran …)" % -i)
            continue
        celdas = [_celda_para_catti(v) for v in filas[i][:ancho]]
        if not any(celdas):
            continue  # un renglón vacío no aporta, y su número de fila igual se ve en el salto
        lineas.append("%4d | " % (i + 1) + " | ".join(celdas + [""] * (ancho - len(celdas))))
    return "\n".join(lineas)


def muestra(datos: bytes, filename: str = "") -> dict:
    """SINCRÓNICO (va en un hilo). Lo que se le muestra a CatTi para que decida
    qué es cada columna: cada hoja (hasta ocho) con sus primeras filas y las
    últimas, en texto. Devuelve {"texto", "hojas": [nombres]}.

    El freno barato del catálogo corre ANTES, igual que en `leer`: mandarle un
    catálogo de 200.000 filas al modelo para que después lo rechace el tope
    sería gastar tokens y tiempo en decir que no."""
    if _cuantas_filas(datos, filename, None, MAX_FILAS_CRUDAS) > MAX_FILAS_CRUDAS:
        raise _demasiadas_filas_crudas()
    nombres = _hojas(datos, filename)
    partes: list[str] = []
    if nombres:
        for n, nombre in enumerate(nombres[:_MUESTRA_HOJAS]):
            partes.append(_tabla_para_catti(nombre, _filas_de(datos, filename, n)))
        if len(nombres) > _MUESTRA_HOJAS:
            partes.append("(El archivo tiene %d hojas más que no se muestran: %s)" % (
                len(nombres) - _MUESTRA_HOJAS, ", ".join(nombres[_MUESTRA_HOJAS:])))
    else:
        partes.append(_tabla_para_catti("", _filas_de(datos, filename, None)))
    return {"texto": "\n\n".join(partes), "hojas": nombres}


def _segun_catti(
    datos: bytes, filename: str, hoja: str | int | None, mapeo: dict,
) -> tuple[list[tuple], str, list[str], int, dict[int, str], dict]:
    """Lo que decidió CatTi (catti.interpretar_planilla), verificado contra el
    archivo antes de creerle: la hoja tiene que existir, las columnas tienen que
    estar dentro de la tabla, y la de la descripción tiene que traer texto. Lo
    que no cierra se descarta y se dice; lo único sin lo cual no se puede
    seguir es la descripción, porque es lo que amarra cada placa con su fila.

    Devuelve (filas, hoja, avisos, índice de la fila de títulos o -1 si no hay,
    {índice de columna: campo}, interpretación para mostrarle a la persona)."""
    avisos: list[str] = []
    nombres = _hojas(datos, filename)

    if hoja is not None:
        # La hoja que eligió la persona manda sobre la que eligió CatTi.
        nombre_hoja = _nombre_de_hoja(nombres, hoja)
        filas = _filas_de(datos, filename, hoja)
    elif nombres:
        pedida = (mapeo.get("hoja") or "").strip()
        if pedida in nombres:
            nombre_hoja = pedida
        else:
            nombre_hoja = nombres[0]
            if pedida:
                avisos.append("CatTi nombró una hoja que no está en el archivo («%s»): leí «%s»." % (pedida, nombre_hoja))
        filas = _filas_de(datos, filename, nombres.index(nombre_hoja))
        if nombre_hoja != nombres[0]:
            saltadas = ", ".join("«%s»" % n for n in nombres[:nombres.index(nombre_hoja)])
            avisos.append("Leí la hoja «%s»: CatTi vio que ahí están los productos (las de adelante son %s)." % (nombre_hoja, saltadas))
    else:
        nombre_hoja = ""
        filas = _filas_de(datos, filename, None)

    fila = mapeo.get("fila_encabezados") or 0
    idx_headers = fila - 1 if isinstance(fila, int) and 1 <= fila <= len(filas) else -1
    ancho = max((len(f) for f in filas), default=0)

    col_map: dict[int, str] = {}
    usados: dict[int, str] = {}
    for clave, campo in CAMPOS_DE_LA_PLANILLA.items():
        texto = (mapeo.get("columnas") or {}).get(clave) or ""
        i = indice_de_letra(texto)
        if i is None:
            continue
        if i >= ancho:
            avisos.append("CatTi ubicó %s en la columna %s, que no existe en la planilla: no lo tuve en cuenta." % (
                _NOMBRE_DEL_DATO[campo].lower(), texto))
            continue
        if i in usados:
            # Una columna es UN dato. Si CatTi pone dos en la misma, se queda el
            # primero (el orden de CAMPOS_DE_LA_PLANILLA va de lo que importa a
            # lo accesorio) y se dice.
            avisos.append("CatTi puso %s y %s en la misma columna (%s): me quedé con %s." % (
                _NOMBRE_DEL_DATO[usados[i]].lower(), _NOMBRE_DEL_DATO[campo].lower(), letra(i),
                _NOMBRE_DEL_DATO[usados[i]].lower()))
            continue
        col_map[i] = campo
        usados[i] = campo

    encabezados = filas[idx_headers] if idx_headers >= 0 else ()
    desc = next((i for i, c in col_map.items() if c == "descripcion"), None)
    con_texto = desc is not None and any(
        _texto(f[desc]) for f in filas[idx_headers + 1:] if desc < len(f)
    )
    if not con_texto:
        titulos_vistos = ", ".join(_texto(c) for c in encabezados if _texto(c))[:200]
        raise PlanillaInvalida(
            "CatTi no encontró en la planilla ninguna columna con la descripción del producto, "
            "que es lo que amarra cada placa con su fila."
            + (" Los títulos que vio son: %s." % titulos_vistos if titulos_vistos else "")
            + ((" " + mapeo["dudas"]) if mapeo.get("dudas") else "")
        )

    interpretacion = {
        "explicacion": mapeo.get("explicacion") or "",
        "dudas": mapeo.get("dudas") or "",
        "columnas": [
            {
                "dato": _NOMBRE_DEL_DATO[campo],
                "letra": letra(i),
                "titulo": _titulo_de_columna(encabezados, i),
            }
            for i, campo in sorted(col_map.items(), key=lambda kv: list(CAMPOS_DE_LA_PLANILLA.values()).index(kv[1]))
        ],
    }
    if interpretacion["dudas"]:
        avisos.append("CatTi tuvo una duda al leer la planilla: " + interpretacion["dudas"])
    return filas, nombre_hoja, avisos, idx_headers, col_map, interpretacion


def _titulo_de_columna(encabezados, i: int) -> str:
    """El título de la columna como está en el archivo, o "columna F" si la
    planilla no tiene fila de títulos."""
    t = _texto(encabezados[i]) if i < len(encabezados) else ""
    return t or "columna %s" % letra(i)


def leer(
    datos: bytes, filename: str = "", hoja: str | int | None = None, mapeo: dict | None = None,
) -> Planilla:
    """SINCRÓNICO (va en un hilo, ver hilos.py): la planilla entera convertida
    en el mismo dict que devuelve la lectura de un mailing.

    `mapeo` es lo que decidió CatTi sobre qué es cada columna (ver
    CAMPOS_DE_LA_PLANILLA y catti.interpretar_planilla); es el camino de
    siempre. Sin `mapeo` las columnas se reconocen por su NOMBRE, que es el
    repuesto para cuando no se puede consultar al modelo.

    openpyxl es CPU puro y bloqueante. Lo que acá NO se hace es dibujar: la tira
    de evidencia de cada fila se dibuja cuando se pide (ver `evidencia`)."""
    # El freno barato: se mira ANTES de materializar nada -- ver `_cuantas_filas`
    # y el comentario de MAX_FILAS_CRUDAS. El tope de verdad (los PRODUCTOS) se
    # mira más abajo, recién cuando se sabe cuántos son.
    if _cuantas_filas(datos, filename, hoja, MAX_FILAS_CRUDAS) > MAX_FILAS_CRUDAS:
        raise _demasiadas_filas_crudas()

    interpretacion = None
    if mapeo is not None:
        filas, nombre_hoja, avisos, idx_headers, col_map, interpretacion = _segun_catti(
            datos, filename, hoja, mapeo,
        )
        if len(filas) > MAX_FILAS_CRUDAS:  # el atajo no pudo contar: se chequea igual
            raise _demasiadas_filas_crudas()
    else:
        filas, nombre_hoja, avisos = _elegir_hoja(datos, filename, hoja)
        if len(filas) > MAX_FILAS_CRUDAS:  # el atajo no pudo contar: se chequea igual
            raise _demasiadas_filas_crudas()

        idx_headers = _fila_de_encabezados(filas)
        if idx_headers is None:
            primeras = ", ".join(str(c) for c in (filas[0] if filas else ()) if c is not None)[:200]
            raise PlanillaInvalida(
                "No encontré la fila de encabezados de la planilla: tiene que tener una "
                "columna con la descripción del producto (DESCRIPCION o NOMBRE ARTÍCULO)."
                + (" La primera fila dice: %s" % primeras if primeras else "")
            )
        col_map = _columnas_de(filas[idx_headers])

        col_desc = _columna_de_descripcion(col_map)
        if col_desc is None:  # pragma: no cover - _fila_de_encabezados ya lo garantiza
            raise PlanillaInvalida("La planilla no trae una columna con la descripción del producto")
        # De acá en adelante la descripción elegida se llama "descripcion" y punto:
        # así el resto del módulo no tiene que preguntarse cuál de las dos era.
        col_map = {i: ("descripcion" if c == col_desc else c) for i, c in col_map.items()}
        # La otra columna de descripción no se usa ni se muestra: dos descripciones
        # en la evidencia es una invitación a corregir la placa contra la que no era.
        col_map = {
            i: c for i, c in col_map.items() if c not in ("descripcionExcel", "descripcionWeb", "nombreArticulo")
        }

    encabezados = filas[idx_headers] if idx_headers >= 0 else ()
    titulos = {campo: _titulo_de_columna(encabezados, i) for i, campo in col_map.items()}

    crudas: list[tuple[int, dict]] = []
    sin_descripcion = 0
    for n, cruda in enumerate(filas[idx_headers + 1:], start=idx_headers + 2):
        valores = {campo: _texto(cruda[i]) if i < len(cruda) else "" for i, campo in col_map.items()}
        if _fila_vacia(valores):
            continue  # un renglón en blanco no es una fila que quedó afuera
        if not _fila_de_datos(valores):
            sin_descripcion += 1
            continue
        crudas.append((n, valores))

    # El pie del reporte se decide mirando el bloque entero, no fila por fila:
    # hace falta saber qué hay ABAJO (nada) y qué hay ARRIBA (los totales).
    posiciones_pie, dudosas = _pies_del_final(crudas)
    pies = ["fila %s: «%s»" % (crudas[p][0], crudas[p][1]["descripcion"]) for p in sorted(posiciones_pie)]
    dudas_de_pie = ["fila %s: «%s»" % (crudas[p][0], crudas[p][1]["descripcion"]) for p in sorted(dudosas)]
    crudas = [(n, v) for p, (n, v) in enumerate(crudas) if p not in posiciones_pie]

    # Y recién después se junta: un listado en formato largo trae el mismo
    # producto repetido una vez por sucursal (ver `_juntar_si_viene_en_formato_largo`).
    juntadas, largo = _juntar_si_viene_en_formato_largo([v for _, v in crudas])
    if largo["filas"]:
        # El número de fila del producto es el de su PRIMERA aparición: es donde
        # el diseñador lo va a encontrar al abrir el archivo.
        primera = {}
        for n, v in crudas:
            primera.setdefault(v["descripcion"], n)
        crudas = [(primera[v["descripcion"]], v) for v in juntadas]
        # La sucursal y el stock eran datos de UNA de las filas que se juntaron,
        # no del producto: dejarlos como columna invita a leerlos como si lo
        # fueran ("este auricular es de Central").
        titulos = {c: t for c, t in titulos.items() if c not in ("sucursal", "stock")}

    # Y ACÁ SE MIRA EL TOPE DE VERDAD, que cuenta PRODUCTOS: recién después de
    # juntar se sabe cuántos son. Antes se miraba contra las filas crudas y un
    # listado de 350 productos por 3 sucursales --1.050 filas, el listado
    # correcto-- se rechazaba con "subí el listado de esta campaña, no el
    # catálogo entero", que es acusar de lo contrario de lo que pasó.
    if len(crudas) > MAX_FILAS:
        raise _demasiados_productos()

    productos: list[dict] = []
    vigencias: list[str] = []
    alcoholes: list[str] = []
    for n, valores in crudas:
        moneda = valores.get("moneda", "")
        if valores.get("vigencia"):
            vigencias.append(valores["vigencia"])
        if valores.get("legalAlcohol"):
            alcoholes.append(valores["legalAlcohol"])
        producto = {
            "descripcion": valores["descripcion"],
            "precio_anterior": _precio(valores.get("precioAnterior"), moneda),
            "precio_anterior_tachado": False,  # la planilla no lo dice: no se compara (ver "campos")
            "mecanica": valores.get("mecanica", ""),
            "oferta_encabezado": valores.get("ofertaEncabezado", ""),
            "oferta_precio": _precio(valores.get("precio"), moneda),
            "oferta_pie": valores.get("ofertaPie", ""),
            "es_alcohol": False,  # tampoco: la leyenda se decide por lo que lee la placa
            "pagina": None,
            "cajas": {},
            "recorte": None,
            # `numero` es el número de fila REAL del archivo, el que muestra
            # Excel al abrirlo. Si el diseñador va a la fila 24 y ahí hay otra
            # cosa, se perdió toda la credibilidad de una.
            "fila": {"numero": n, "valores": valores},
        }
        productos.append(producto)

    if not productos:
        raise PlanillaInvalida("La planilla no tiene ninguna fila con descripción")

    # Qué puede exigir esta planilla: los campos cuya columna existe Y tiene
    # algún valor. Una columna MECANICA entera vacía no habilita a marcar como
    # "sobra" la mecánica de una placa.
    campos = [
        campo for campo, columna in _DE_LA_COLUMNA.items()
        if columna in titulos and any(p[campo] for p in productos)
    ]

    # Las columnas de precio en las que alguna fila escribió lo que acompaña al
    # número ("$340 unidad"): ahí la columna dicta también la cola, y una fila
    # que dice "$171" a secas está diciendo que NO va "unidad" (ver
    # comparador._comparar_precio). Un export de gestión trae números y no
    # entra acá: sus precios se siguen comparando solo por importe.
    precios_con_cola = [
        campo for campo, columna in (("precio_anterior", "precioAnterior"), ("oferta_precio", "precio"))
        if columna in titulos and any(cola_del_precio(p[campo]) for p in productos)
    ]

    # La columna elegida trae el nombre corto de gestión y no el texto impreso:
    # se sigue usando para EMPAREJAR (el parecido ignora mayúsculas y tildes)
    # pero deja de dictar la descripción, igual que cualquier otra columna que
    # la planilla no trae. Acusar a las 12 placas de una campaña perfecta porque
    # la planilla está en MAYÚSCULAS es el motor equivocándose en voz alta.
    de_gestion = _parece_vocabulario_de_gestion([p["descripcion"] for p in productos])
    if de_gestion:
        campos = [c for c in campos if c != "descripcion"]
        ejemplo = next(p["descripcion"] for p in productos if not any(c.islower() for c in p["descripcion"]))
        avisos.append(
            "La columna %s viene en MAYÚSCULAS y abreviada («%s»): así se llama el "
            "artículo en gestión, no como va impreso en la placa. No comparo la "
            "descripción contra eso —marcaría mal a todas las placas—, pero la sigo "
            "usando para emparejar cada placa con su fila."
            % (titulos.get("descripcion", "DESCRIPCIÓN"), ejemplo)
        )

    if sin_descripcion:
        cuantas = "1 fila" if sin_descripcion == 1 else "%d filas" % sin_descripcion
        avisos.append("Afuera: %s sin descripción" % cuantas)
    if pies:
        avisos.append("No lo conté como producto, parece el pie del reporte: " + "; ".join(pies))
    if dudas_de_pie:
        # El motor no la pudo decidir: la contó como producto (que es lo menos
        # destructivo) y lo dice, en vez de resolverlo en silencio para cualquier
        # lado. Quien mira el archivo sabe en dos segundos cuál de las dos es.
        avisos.append(
            "Parece el pie del reporte pero trae datos en otras columnas, así que la conté "
            "como producto: " + "; ".join(dudas_de_pie) + ". Si no es un producto, sacala del archivo."
        )
    if largo["filas"]:
        sucursales = len(largo["sucursales"])
        avisos.append(
            "La planilla viene en formato largo (una fila por producto y por sucursal, %d "
            "sucursales): junté las %d filas en %d productos. Cada placa se empareja con un "
            "producto, no con una fila."
            % (sucursales, largo["filas"], largo["productos"])
        )
        if largo["con_diferencias"]:
            avisos.append(
                "Ojo: %d de esos productos traen datos distintos entre sus filas (precios o "
                "textos que no coinciden). Me quedé con el primer valor que no estaba vacío."
                % largo["con_diferencias"]
            )

    mailing = {
        "origen": "planilla",
        "fecha": vigencias[0] if vigencias else "",
        "fecha_pagina": None, "fecha_caja": None,
        "legal_alcohol": alcoholes[0] if alcoholes else "",
        "legal_alcohol_pagina": None, "legal_alcohol_caja": None,
        "campos": campos,
        "precios_con_cola": precios_con_cola,
        "planilla": {
            "archivo": filename or "planilla",
            "hoja": nombre_hoja,
            "titulos": titulos,
            "fila_encabezado": idx_headers + 1,
            "filas_leidas": len(productos),
            "filas_ignoradas": sin_descripcion + len(pies),
            # Cuántas filas del archivo se juntaron en esos productos, 0 si la
            # planilla no venía en formato largo. Está a la vista por lo mismo
            # que los otros contadores: es el dato que delata que el motor
            # entendió el archivo de otra manera que la persona.
            "filas_juntadas": largo["filas"],
            "descripcion_de_gestion": de_gestion,
            # Qué entendió CatTi de la planilla: qué columna es cada dato, en
            # sus palabras. Va a la pantalla y al Excel para que la persona vea
            # CÓMO se leyó su archivo antes de creerle a una sola corrección.
            # None cuando la planilla se leyó por los nombres de las columnas.
            "interpretacion": interpretacion,
            # Los avisos viven ACÁ y no solo en el dataclass: esto es lo que se
            # guarda en la base y lo que llega a la pantalla y al Excel.
            # `Planilla.avisos` los armaba y se perdían en el camino
            # (preparar_planilla devolvía solo pl.mailing).
            "avisos": avisos,
        },
        "productos": productos,
    }
    # La tira de evidencia NO se dibuja acá: se dibuja cuando se pide (ver
    # `evidencia` y el comentario de MAX_FILAS). Dibujar una por fila era el
    # trabajo caro de este módulo y el 99 % se tiraba: solo se muestran las
    # filas que alguna placa reclamó.
    return Planilla(mailing, avisos)


def _nombre_de_hoja(nombres: list[str], hoja) -> str:
    """Cómo se llama la hoja que se leyó. Un CSV no tiene hojas: cadena vacía,
    y la cita queda "listado.csv · fila 24" sin un «hoja «CSV»» que no existe."""
    if isinstance(hoja, str):
        return hoja
    if not nombres:
        return ""
    n = hoja or 0
    return nombres[n] if 0 <= n < len(nombres) else nombres[0]


# --------------------------------------------------------------------------
# La evidencia: la fila de la planilla dibujada
# --------------------------------------------------------------------------
#
# Contra un mailing la evidencia es un RECORTE: un rectángulo de la fuente con
# lo relevante adentro, que se muestra al lado de la placa en pantalla y se pega
# en la columna B del Excel de correcciones. Con una planilla no hay foto que
# recortar, y si de este lado pasáramos a texto plano el Excel perdería la cara
# visual de un lado y se rompería la simetría placa <-> fuente, que es lo que
# hace que se entienda de un vistazo.
#
# Así que la fila se DIBUJA como lo que es: un pedazo de la planilla, con su
# encabezado real, la fila marcada con ▶, una vecina arriba y otra abajo (los
# vecinos son los que hacen que el diseñador CREA el emparejamiento: ve que
# arriba y abajo hay otros productos) y, cuando se está mostrando un campo
# puntual, la celda de ese campo encuadrada. Arriba va siempre la cita
# --archivo, hoja y número de fila-- para que lo pueda chequear él mismo: esa
# es la diferencia entre creerle al sistema y poder auditarlo.
#
# Se dibuja CHICO y con letra grande a propósito: en el Excel esta imagen entra
# en una columna de ~240 px y openpyxl la escala. Una tira de 1200 px con letra
# de 13 px llega ahí como letra de 3 px.

_ANCHO_TIRA = 600
_ALTO_FILA = 34
_PAD = 10
_TAM_TEXTO = 17
_TAM_TITULO = 14

_TINTA = (32, 42, 54)
_SUAVE = (107, 118, 134)
_LINEA = (215, 220, 227)
_FONDO_ELEGIDA = (255, 244, 214)
_MARCO = (192, 57, 43)


def _ancho(dib, texto: str, fuente) -> float:
    """Lo que mide un texto, sin poder reventar.

    `ImageDraw.textlength` tira `ValueError: can't measure length of multiline
    text` con un solo salto de línea adentro, y eso tumbaba la request entera
    con un 500 -- aunque el salto estuviera en una fila VECINA de la que se
    estaba dibujando. El texto ya viene limpio de `_texto`, así que esto es el
    segundo cinturón: lo que se dibuja NUNCA puede voltear una validación."""
    return dib.textlength(_una_linea(texto), font=fuente)


def _acortar(dib, texto: str, fuente, ancho: int) -> str:
    texto = _una_linea(texto)
    if _ancho(dib, texto, fuente) <= ancho:
        return texto
    while texto and _ancho(dib, texto + "…", fuente) > ancho:
        texto = texto[:-1]
    return texto + "…"


def _partir(dib, texto: str, fuente, ancho: int, max_lineas: int) -> list[str]:
    lineas: list[str] = []
    actual = ""
    for palabra in _una_linea(texto).split(" "):
        prueba = f"{actual} {palabra}".strip()
        if _ancho(dib, prueba, fuente) <= ancho or not actual:
            actual = prueba
        else:
            lineas.append(actual)
            actual = palabra
            if len(lineas) == max_lineas:
                break
    if len(lineas) < max_lineas and actual:
        lineas.append(actual)
    if len(lineas) == max_lineas and actual and lineas[-1] != actual:
        lineas[-1] = _acortar(dib, lineas[-1] + " …", fuente, ancho)
    return lineas or [""]


def _columnas_a_mostrar(mailing: dict, campo: str | None) -> list[str]:
    """Qué columnas entran en la tira. Siempre la descripción (es la que amarra
    la fila con la placa) y, si se está mostrando un campo puntual, la columna
    de ese campo. Más que eso no entra legible en el ancho que después le deja
    el Excel."""
    titulos = mailing["planilla"]["titulos"]
    columnas = ["descripcion"]
    objetivo = COLUMNA_DEL_CAMPO.get(campo or "")
    if objetivo and objetivo != "descripcion" and objetivo in titulos:
        columnas.append(objetivo)
    elif not objetivo and "precio" in titulos:
        columnas.append("precio")
    return columnas


def tira(mailing: dict, indice: int, campo: str | None = None, vecinas: bool = True) -> Image.Image:
    """Un pedazo de la planilla alrededor de la fila `indice`, como imagen."""
    productos = mailing["productos"]
    pl = mailing["planilla"]
    titulos = pl["titulos"]
    columnas = _columnas_a_mostrar(mailing, campo)
    marcada = COLUMNA_DEL_CAMPO.get(campo or "")

    vecinos = [indice - 1, indice, indice + 1] if vecinas else [indice]
    vecinos = [i for i in vecinos if 0 <= i < len(productos)]

    f_texto = imagenes.fuente(_TAM_TEXTO)
    f_negrita = imagenes.fuente(_TAM_TEXTO, negrita=True)
    f_titulo = imagenes.fuente(_TAM_TITULO, negrita=True)
    f_cita = imagenes.fuente(_TAM_TITULO)

    # Anchos: la columna del número es fija, la del valor se queda con lo que
    # necesita su encabezado y la descripción con todo lo que sobra.
    lienzo = Image.new("RGB", (_ANCHO_TIRA, 10), "white")
    dib = ImageDraw.Draw(lienzo)
    x_fila = 54
    ancho_valor = 0
    if len(columnas) > 1:
        ancho_valor = int(max(
            _ancho(dib, titulos.get(columnas[1], ""), f_titulo) + 18,
            *(_ancho(dib, productos[i]["fila"]["valores"].get(columnas[1], ""), f_negrita) + 18 for i in vecinos),
            90,
        ))
        ancho_valor = min(ancho_valor, 200)
    ancho_desc = _ANCHO_TIRA - 2 * _PAD - x_fila - ancho_valor

    # Alto: la descripción de la fila elegida se parte hasta en 3 renglones.
    altos = []
    textos_desc = {}
    for i in vecinos:
        lineas = _partir(dib, productos[i]["fila"]["valores"].get("descripcion", ""), f_texto,
                         ancho_desc - 12, 3 if i == indice else 1)
        textos_desc[i] = lineas
        altos.append(max(_ALTO_FILA, len(lineas) * 22 + 12))

    cita_alto = 46
    alto_cabecera = 28
    alto = cita_alto + alto_cabecera + sum(altos) + _PAD
    lienzo = Image.new("RGB", (_ANCHO_TIRA, alto), "white")
    dib = ImageDraw.Draw(lienzo)

    # La cita: archivo · hoja · fila, arriba de todo.
    hoja = f" · hoja «{pl['hoja']}»" if pl.get("hoja") else ""
    dib.text((_PAD, 6), _acortar(dib, f"{pl['archivo']}{hoja}", f_cita, _ANCHO_TIRA - 2 * _PAD), font=f_cita, fill=_SUAVE)
    titulo_col = titulos.get(marcada or "", "")
    donde = f"fila {productos[indice]['fila']['numero']}"
    if titulo_col:
        donde += f" · columna {titulo_col}"
    dib.text((_PAD, 24), _una_linea(donde), font=f_titulo, fill=_TINTA)

    y = cita_alto
    # Encabezado
    dib.rectangle([_PAD, y, _ANCHO_TIRA - _PAD, y + alto_cabecera], fill=(242, 244, 247))
    dib.text((_PAD + 6, y + 6), "fila", font=f_titulo, fill=_SUAVE)
    dib.text((_PAD + x_fila, y + 6), _una_linea(titulos.get("descripcion", "DESCRIPCIÓN")), font=f_titulo, fill=_SUAVE)
    if len(columnas) > 1:
        dib.text((_ANCHO_TIRA - _PAD - ancho_valor + 6, y + 6),
                 _acortar(dib, titulos.get(columnas[1], _TITULO_POR_DEFECTO.get(columnas[1], "")), f_titulo, ancho_valor - 12),
                 font=f_titulo, fill=_SUAVE)
    y += alto_cabecera

    for i, alto_fila in zip(vecinos, altos):
        elegida = i == indice
        if elegida:
            dib.rectangle([_PAD, y, _ANCHO_TIRA - _PAD, y + alto_fila], fill=_FONDO_ELEGIDA)
        dib.line([(_PAD, y), (_ANCHO_TIRA - _PAD, y)], fill=_LINEA)
        if elegida:
            # El triangulito se DIBUJA, no se escribe: "▶" no existe en Segoe
            # UI y en Windows salía como un cuadradito vacío. Mismo problema que
            # ya había dejado escrito excel.py con la negrita del "9:16": con
            # estas fuentes no se puede confiar en un carácter decorativo.
            cy = y + 16
            dib.polygon([(_PAD + 5, cy - 6), (_PAD + 5, cy + 6), (_PAD + 14, cy)], fill=_MARCO)
        dib.text((_PAD + 20, y + 8), str(productos[i]["fila"]["numero"]),
                 font=f_negrita if elegida else f_texto, fill=_TINTA if elegida else _SUAVE)
        for k, linea in enumerate(textos_desc[i]):
            dib.text((_PAD + x_fila, y + 8 + k * 22), linea, font=f_texto, fill=_TINTA if elegida else _SUAVE)
        if len(columnas) > 1:
            valor = productos[i]["fila"]["valores"].get(columnas[1], "")
            x0 = _ANCHO_TIRA - _PAD - ancho_valor
            dib.text((x0 + 8, y + 8), _acortar(dib, valor, f_negrita if elegida else f_texto, ancho_valor - 16),
                     font=f_negrita if elegida else f_texto, fill=_TINTA if elegida else _SUAVE)
            if elegida and marcada == columnas[1]:
                dib.rectangle([x0 + 2, y + 3, _ANCHO_TIRA - _PAD - 3, y + alto_fila - 3], outline=_MARCO, width=3)
        if elegida and marcada == "descripcion":
            dib.rectangle([_PAD + x_fila - 6, y + 3, _ANCHO_TIRA - _PAD - ancho_valor - 4, y + alto_fila - 3],
                          outline=_MARCO, width=3)
        y += alto_fila
    dib.rectangle([_PAD, cita_alto, _ANCHO_TIRA - _PAD, y], outline=_LINEA)
    return lienzo


def evidencia(mailing: dict, indice: int, campo: str | None = None, vecinas: bool = True) -> str | None:
    """La tira como data URI, lista para pegar en el Excel o mostrar en pantalla.

    ES LA PUERTA DE LA EVIDENCIA, Y SE LLAMA CUANDO SE PIDE. Se dibujaba una por
    fila al leer el archivo, y el 99 % se tiraba: la tira solo se muestra de las
    filas que alguna placa reclamó. Con 900 filas eran 5,92 s de CPU y 19,5 MB
    guardados en la base para nada, adentro del único hilo de CPU del servicio
    (ver MAX_FILAS). Ahora se dibuja una por placa emparejada, que es lo que ya
    hacía `recorte_de_campo` para el lado derecho de cada diferencia.

    None si no se pudo dibujar. Es una ILUSTRACIÓN: si falla, la validación
    sigue y la diferencia se muestra igual con el texto de los dos lados. Que
    un dibujo pueda tumbar una validación entera con un 500 --como pasaba con
    un Alt+Enter en una celda-- es lo que no puede volver a pasar. Que la tira
    falte NO se esconde: la pantalla y el Excel lo dicen (ver
    validador.validar_placa)."""
    try:
        return imagenes.data_uri(imagenes.jpeg(tira(mailing, indice, campo, vecinas), 88))
    except Exception:
        logger.warning("rrss: no pude dibujar la fila %s de la planilla", indice, exc_info=True)
        return None


def avisos(mailing: dict) -> list[str]:
    """Los avisos de la planilla, vacíos si la fuente fue un mailing."""
    return list(((mailing or {}).get("planilla") or {}).get("avisos") or [])


def recorte_de_campo(mailing: dict, indice: int, campo: str) -> str | None:
    """La fila sola, con la celda del campo encuadrada: es el lado derecho de
    cada diferencia en pantalla."""
    if not (0 <= indice < len(mailing.get("productos") or [])):
        return None
    return evidencia(mailing, indice, campo, vecinas=False)


# --------------------------------------------------------------------------
# Nombres para una persona
# --------------------------------------------------------------------------

_ETIQUETA_CAMPO = {
    "descripcion": "Descripción",
    "precio_anterior": "Precio anterior",
    "oferta_precio": "Precio de oferta",
    "mecanica": "Mecánica",
    "oferta_encabezado": "Texto arriba del precio",
    "oferta_pie": "Texto abajo del precio",
    "precio_anterior_tachado": "Si el precio anterior va tachado",
}


def campos_que_no_dicta(mailing: dict, config: dict | None = None) -> list[str]:
    """Los campos que esta planilla NO puede exigir, con el nombre con que se
    los conoce en pantalla. Se muestran siempre: "lo muestro distinto según el
    caso" solo vale si se ve POR QUÉ.

    Además de los campos del producto entran la FECHA y la LEYENDA DE ALCOHOL:
    se comparan solo si la planilla trae la columna (VIGENCIA, LEYENDA ALCOHOL)
    o la persona las escribió en la pantalla de carga (`config`). Sin ninguna de
    las dos no se comparan (comparador.comparar_elementos), y hasta el
    22/09/2026 el informe no lo decía en ningún lado."""
    if mailing.get("origen") != "planilla":
        return []
    config = config or {}
    trae = set(mailing.get("campos") or ())
    afuera = [c for c in _DE_LA_COLUMNA if c not in trae] + ["precio_anterior_tachado"]
    nombres = [_ETIQUETA_CAMPO[c] for c in afuera]
    if not ((config.get("fecha") or "").strip() or mailing.get("fecha")):
        nombres.append("Fecha de la campaña")
    if not ((config.get("legal_alcohol") or "").strip() or mailing.get("legal_alcohol")):
        nombres.append("Leyenda de alcohol")
    return nombres


def dicta_todo(mailing: dict) -> bool:
    """¿Esta fuente dicta TODOS los campos del producto que una planilla puede
    traer? Un mailing siempre; una planilla, solo si trae las seis columnas con
    algo escrito. Es lo que decide si el Excel puede decir "Está bien" de una
    placa o tiene que conformarse con "Sin diferencias (2 de 7 campos)"."""
    if mailing.get("origen") != "planilla":
        return True
    return set(_DE_LA_COLUMNA) <= set(mailing.get("campos") or ())


def cita(mailing: dict, indice: int) -> str:
    """"LISTADO.xlsx · hoja «Alemania 2026» · fila 24" -- la coordenada que
    permite ir a mirarlo a mano."""
    pl = mailing.get("planilla") or {}
    productos = mailing.get("productos") or []
    if not (0 <= indice < len(productos)):
        return ""
    partes = [pl.get("archivo") or "la planilla"]
    if pl.get("hoja"):
        partes.append(f"hoja «{pl['hoja']}»")
    partes.append(f"fila {productos[indice]['fila']['numero']}")
    return " · ".join(partes)
